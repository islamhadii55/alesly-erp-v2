import os
import io
import json
import base64
import sqlite3
import uuid
from datetime import datetime, date, timedelta
from functools import wraps
from flask import (
    Flask, g, render_template, request, redirect, url_for,
    session, flash, jsonify, send_from_directory, make_response, send_file
)
from werkzeug.middleware.proxy_fix import ProxyFix

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as PdfImage
    from reportlab.graphics.barcode import code128
    from reportlab.graphics.shapes import Drawing
    import arabic_reshaper
    from bidi.algorithm import get_display
    pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    REPORTLAB_OK = True
except Exception:
    REPORTLAB_OK = False
    code128 = None
    Drawing = None


def rtl_pdf(value):
    text = str(value or "")
    if not REPORTLAB_OK:
        return text
    return get_display(arabic_reshaper.reshape(text))

APP_NAME = "الاصلي لتجارة قطع الغيار وخدمات صيانة السيارات"


def resolve_db_path():
    explicit = (os.environ.get("DATABASE_PATH") or os.environ.get("SQLITE_PATH") or "").strip()
    if explicit:
        return explicit
    volume = (os.environ.get("RAILWAY_VOLUME_MOUNT_PATH") or "").strip()
    if volume:
        return os.path.join(volume, "original_auto.db")
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "original_auto.db")


DB_PATH = resolve_db_path()
IS_PRODUCTION = bool(
    os.environ.get("RAILWAY_ENVIRONMENT")
    or os.environ.get("RAILWAY_PUBLIC_DOMAIN")
    or os.environ.get("RAILWAY_PROJECT_ID")
)

app = Flask(__name__)
app.secret_key = os.environ.get("APP_SECRET_KEY", "original-auto-parts-local-secret")
app.permanent_session_lifetime = timedelta(days=30)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)
app.config.update(
    PREFERRED_URL_SCHEME="https" if IS_PRODUCTION else "http",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=IS_PRODUCTION,
)

# Permission modules: key -> (label, group)
PERMISSION_MODULES = (
    ("dashboard", "لوحة التحكم", "الرئيسية"),
    ("branches", "الفروع وإدارة التوزيع", "الإدارة"),
    ("pos", "كاشير قطع الغيار", "المبيعات والورشة"),
    ("sales", "فواتير البيع والمرتجعات والصيانة", "المبيعات والورشة"),
    ("quotes", "عروض الأسعار", "المبيعات والورشة"),
    ("jobs", "أوامر الشغل", "المبيعات والورشة"),
    ("purchases", "فواتير الشراء ومرتجعاتها", "المشتريات"),
    ("inventory", "المخزون والأصناف", "المخزون والحسابات"),
    ("manufacturing", "التصنيع وأوامر الإنتاج", "المخزون والحسابات"),
    ("parties", "العملاء والموردون", "المخزون والحسابات"),
    ("accounts", "حسابات العملاء والموردين", "المخزون والحسابات"),
    ("treasury", "الخزينة والورديات", "المخزون والحسابات"),
    ("finance", "شيكات وأقساط وسندات", "المخزون والحسابات"),
    ("journal", "القيود اليومية", "المخزون والحسابات"),
    ("costs", "حسابات التكلفة", "المخزون والحسابات"),
    ("pnl", "الأرباح والخسائر", "المخزون والحسابات"),
    ("expenses", "المصروفات", "المخزون والحسابات"),
    ("marketing", "حملات وكوبونات", "المخزون والحسابات"),
    ("reports", "التقارير", "المخزون والحسابات"),
    ("sync", "المزامنة بدون إنترنت", "المخزون والحسابات"),
    ("hr", "الموارد البشرية", "الموارد البشرية"),
    ("attendance", "الحضور والغياب والتأخير", "الموارد البشرية"),
    ("shift_start", "بدء الوردية", "الموارد البشرية"),
    ("shift_end", "إنهاء الوردية", "الموارد البشرية"),
    ("attendance_review", "مراجعة خصومات الحضور", "الموارد البشرية"),
)
PERMISSION_KEYS = {key for key, _, _ in PERMISSION_MODULES}
PERMISSION_LABELS = {key: label for key, label, _ in PERMISSION_MODULES}

# endpoint -> permission module. Endpoints missing here need no permission.
ENDPOINT_PERMISSIONS = {
    "dashboard": "dashboard",
    "pos": "pos",
    "shift_open": "shift_start",
    "shift_close": "shift_end",
    "attendance_mark": "attendance",
    "attendance": "attendance",
    "invoices_list": "invoices",
    "invoice_new": "invoices",
    "invoice_view": "invoices",
    "invoice_print": "invoices",
    "invoice_pdf": "invoices",
    "invoice_cancel": "invoices",
    "receipt_print": "finance",
    "day_close": "treasury",
    "day_archive": "treasury",
    "quotes": "quotes",
    "quote_to_invoice": "quotes",
    "quote_new": "quotes",
    "quote_pdf": "quotes",
    "jobs": "jobs",
    "job_new": "jobs",
    "job_view": "jobs",
    "job_to_invoice": "jobs",
    "products": "inventory",
    "delete_product": "inventory",
    "inventory": "inventory",
    "inventory_export": "inventory",
    "inventory_template": "inventory",
    "inventory_import": "inventory",
    "inventory_adjust": "inventory",
    "inventory_location": "inventory",
    "inventory_transfer": "inventory",
    "inventory_count": "inventory",
    "inventory_serial": "inventory",
    "inventory_shortages": "inventory",
    "barcode_center": "inventory",
    "product_barcode_pdf": "inventory",
    "manufacturing": "manufacturing",
    "manufacturing_material_delete": "manufacturing",
    "manufacturing_start": "manufacturing",
    "manufacturing_issue": "manufacturing",
    "manufacturing_waste": "manufacturing",
    "manufacturing_finish": "manufacturing",
    "customers": "parties",
    "delete_customer": "parties",
    "suppliers": "parties",
    "delete_supplier": "parties",
    "accounts": "accounts",
    "treasury": "treasury",
    "finance": "finance",
    "journal": "journal",
    "costs": "costs",
    "pnl": "pnl",
    "expenses": "expenses",
    "marketing": "marketing",
    "reports": "reports",
    "reports_export_pdf": "reports",
    "shift_reports": "reports",
    "payroll_deductions": "attendance_review",
    "payroll_deductions_export_pdf": "attendance_review",
    "sync_center": "sync",
    "employees": "hr",
    "delete_employee": "hr",
    "advances": "hr",
    "salaries": "hr",
    "settings": "settings",
    "branches": "branches",
    "branch_select": "branches",
    "branch_transfer": "branches",
}

# Kinds handled by the sales vs purchases permission
SALES_KINDS = {"sale", "sale_return", "maintenance"}
PURCHASE_KINDS = {"purchase", "purchase_return"}


def permission_list(user):
    if not user:
        return set()
    if user["role"] == "مدير":
        return set(PERMISSION_KEYS)
    raw = (user["permissions"] or "").strip()
    if raw == "all":
        return set(PERMISSION_KEYS)
    return {p.strip() for p in raw.split(",") if p.strip() in PERMISSION_KEYS}


def user_has_permission(module):
    if not session.get("user"):
        return False
    if session.get("role") == "مدير":
        return True
    branch_id = session.get("branch_id")
    if branch_id and session.get("user_id"):
        branch_access = query(
            "SELECT permissions, status FROM user_branch_permissions WHERE user_id=? AND branch_id=?",
            (session.get("user_id"), branch_id), one=True,
        )
        if branch_access and branch_access["status"] != "نشط":
            return False
        if branch_access:
            perms = branch_access["permissions"]
        else:
            perms = session.get("permissions")
    else:
        perms = session.get("permissions")
    if perms is None:
        return False
    if perms == "all":
        return True
    granted = {p.strip() for p in str(perms).split(",") if p.strip()}
    if module in {"shift_start", "shift_end"} and "shift" in granted:
        return True
    return module in granted


def accessible_branches_for_user():
    if session.get("role") == "مدير":
        return query("SELECT * FROM branches WHERE status!='معطل' ORDER BY id")
    return query(
        """SELECT b.* FROM branches b JOIN user_branch_permissions ubp ON ubp.branch_id=b.id
           WHERE ubp.user_id=? AND ubp.status='نشط' AND b.status!='معطل' ORDER BY b.id""",
        (session.get("user_id"),),
    )


def resolve_endpoint_permission():
    module = ENDPOINT_PERMISSIONS.get(request.endpoint)
    if module == "invoices":
        kind = (request.view_args or {}).get("kind")
        if kind in SALES_KINDS:
            return "sales"
        if kind in PURCHASE_KINDS:
            return "purchases"
        return "sales"
    return module


OPEN_ENDPOINTS = {
    "login", "logout", "static", "web_manifest", "service_worker",
    "api_offline_catalog", "api_sync_status", "api_sync_push",
    "api_sync_pull", "api_search", "api_product", "api_next_sku",
    "api_barcode", "api_coupon", "dashboard", "health",
}


@app.before_request
def enforce_permissions():
    if not session.get("user") or request.endpoint is None:
        return None
    if session.get("role") == "مدير" or request.endpoint in OPEN_ENDPOINTS:
        return None
    if request.endpoint == "branch_select":
        bid = request.form.get("branch_id")
        if query("SELECT 1 FROM user_branch_permissions WHERE user_id=? AND branch_id=? AND status='نشط'", (session.get("user_id"), bid), one=True):
            return None
    module = resolve_endpoint_permission()
    if module and not user_has_permission(module):
        if request.path.startswith("/api/"):
            return jsonify({"ok": False, "error": "لا تملك صلاحية لهذا الإجراء"}), 403
        flash("لا تملك صلاحية الوصول لهذه الشاشة", "err")
        return redirect(url_for("dashboard"))
    return None


def ensure_db_dir():
    directory = os.path.dirname(os.path.abspath(DB_PATH))
    if directory:
        os.makedirs(directory, exist_ok=True)


def connect_sqlite():
    ensure_db_dir()
    conn = sqlite3.connect(DB_PATH, timeout=30, check_same_thread=False)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA busy_timeout = 30000")
    try:
        conn.execute("PRAGMA journal_mode = WAL")
    except sqlite3.OperationalError:
        pass
    return conn


def db():
    if "db" not in g:
        g.db = connect_sqlite()
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


def query(sql, args=(), one=False):
    cur = db().execute(sql, args)
    rows = cur.fetchall()
    return (rows[0] if rows else None) if one else rows


def execute(sql, args=()):
    cur = db().execute(sql, args)
    db().commit()
    return cur.lastrowid


def money(value):
    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return "0.00"


def time12(value):
    if not value:
        return "—"
    text = str(value)
    try:
        parsed = datetime.strptime(text[:16], "%Y-%m-%d %H:%M")
        suffix = "ص" if parsed.hour < 12 else "م"
        hour = parsed.hour % 12 or 12
        return f"{parsed.strftime('%Y-%m-%d')} {hour:02d}:{parsed.minute:02d} {suffix}"
    except ValueError:
        return text


app.jinja_env.filters["money"] = money
app.jinja_env.filters["time12"] = time12


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapper


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login"))
        if session.get("role") != "مدير":
            flash("هذه الصفحة متاحة للمدير فقط", "err")
            return redirect(url_for("dashboard"))
        return fn(*args, **kwargs)
    return wrapper


CATEGORY_PREFIX = {
    "فرامل": "BRK",
    "فلاتر": "FLT",
    "زيوت": "OIL",
    "كهرباء": "ELC",
    "محرك": "ENG",
    "تبريد": "CLG",
    "تعليق": "SUS",
    "إطارات": "TIR",
    "اطارات": "TIR",
    "إنارة": "LGT",
    "انارة": "LGT",
    "ناقل حركة": "TRN",
    "هياكل": "BDY",
    "إكسسوارات": "ACC",
}


def get_setting(key, default=""):
    row = query("SELECT value FROM settings WHERE key=?", (key,), one=True)
    return row["value"] if row and row["value"] is not None else default


def all_settings():
    return {r["key"]: r["value"] for r in query("SELECT key, value FROM settings")}


def format_location(warehouse="", aisle="", shelf="", bin_code=""):
    parts = []
    if warehouse:
        parts.append(f"مخزن {warehouse}")
    if aisle:
        parts.append(f"ممر {aisle}")
    if shelf:
        parts.append(f"رف {shelf}")
    if bin_code:
        parts.append(f"موضع {bin_code}")
    return " / ".join(parts)


def next_sku(category=""):
    prefix = CATEGORY_PREFIX.get((category or "").strip())
    if not prefix:
        prefix = (get_setting("sku_prefix", "PRD") or "PRD").strip().upper()[:6] or "PRD"
    row = query(
        "SELECT sku FROM products WHERE sku LIKE ? ORDER BY id DESC LIMIT 1",
        (f"{prefix}-%",),
        one=True,
    )
    seq = 1
    if row:
        try:
            seq = int(str(row["sku"]).rsplit("-", 1)[-1]) + 1
        except ValueError:
            seq = 1
    else:
        count = query(
            "SELECT COUNT(*) c FROM products WHERE sku LIKE ?",
            (f"{prefix}-%",),
            one=True,
        )["c"]
        seq = int(count) + 1
    return f"{prefix}-{seq:04d}"


def next_number(prefix, table="invoices"):
    year = datetime.now().strftime("%Y")
    like = f"{prefix}-{year}-%"
    row = query(
        f"SELECT number FROM {table} WHERE number LIKE ? ORDER BY id DESC LIMIT 1",
        (like,),
        one=True,
    )
    seq = 1
    if row:
        try:
            seq = int(str(row["number"]).split("-")[-1]) + 1
        except ValueError:
            seq = 1
    return f"{prefix}-{year}-{seq:04d}"


def seed_demo_documents(conn):
    today = date.today().isoformat()
    year = datetime.now().strftime("%Y")

    def add_invoice(number, kind, party_type, party_id, party_name, status, items, paid=None, vehicle=None, notes=""):
        subtotal = sum(q * price for _, _, q, _, price in items)
        cost_total = sum(q * cost for _, _, q, cost, _ in items)
        total = round(subtotal, 2)
        profit = round(subtotal - cost_total, 2)
        if kind in ("sale_return",):
            profit = round(-profit, 2)
        if kind in ("purchase", "purchase_return"):
            profit = 0
        paid_val = total if paid is None else paid
        cur = conn.execute(
            """INSERT INTO invoices
               (number, kind, party_type, party_id, party_name, date, status, subtotal, discount, tax,
                total, paid, cost_total, profit, vehicle, notes, created_by)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                number, kind, party_type, party_id, party_name, today, status,
                subtotal, 0, 0, total, paid_val, cost_total, profit, vehicle, notes, "admin",
            ),
        )
        inv_id = cur.lastrowid
        for pid, desc, qty, cost, price in items:
            line_total = qty * price
            line_cost = qty * cost
            conn.execute(
                """INSERT INTO invoice_items
                   (invoice_id, product_id, description, qty, unit_cost, unit_price, line_total, line_cost, line_profit)
                   VALUES (?,?,?,?,?,?,?,?,?)""",
                (inv_id, pid, desc, qty, cost, price, line_total, line_cost, line_total - line_cost),
            )
            prod = conn.execute("SELECT qty, cost FROM products WHERE id=?", (pid,)).fetchone()
            if not prod:
                continue
            delta = qty if kind in ("purchase", "sale_return") else -qty
            new_qty = prod[0] + delta
            conn.execute("UPDATE products SET qty=? WHERE id=?", (new_qty, pid))
            conn.execute(
                """INSERT INTO stock_moves (product_id, qty, unit_cost, move_type, ref, date, notes)
                   VALUES (?,?,?,?,?,?,?)""",
                (pid, delta, cost, kind, number, today, desc),
            )
        return inv_id

    add_invoice(
        f"PUR-{year}-0001", "purchase", "supplier", 1, "الوكلاء المتحدة لقطع الغيار", "نقدي",
        [
            (1, "تيل فرامل أمامي", 10, 85, 85),
            (5, "بطارية 70 أمبير", 4, 280, 280),
            (10, "إطار 205/55 R16", 8, 260, 260),
        ],
        notes="توريد أول الشهر",
    )
    add_invoice(
        f"SAL-{year}-0001", "sale", "customer", 1, "ورشة النور للصيانة", "نقدي",
        [
            (1, "تيل فرامل أمامي", 2, 85, 140),
            (2, "فلتر زيت أصلي", 6, 12, 25),
            (4, "زيت محرك 5W30 4 لتر", 4, 55, 85),
        ],
    )
    add_invoice(
        f"SAL-{year}-0002", "sale", "customer", 4, "شركة أسطول النقل", "آجل",
        [
            (5, "بطارية 70 أمبير", 2, 280, 420),
            (10, "إطار 205/55 R16", 4, 260, 390),
        ],
        paid=0,
    )
    add_invoice(
        f"MNT-{year}-0001", "maintenance", "customer", 3, "أحمد سعيد الغامدي", "نقدي",
        [
            (6, "بواجي إيريديوم", 1, 48, 90),
            (4, "زيت محرك 5W30 4 لتر", 1, 55, 85),
            (2, "فلتر زيت أصلي", 1, 12, 25),
        ],
        vehicle="كامري 2018 - أ ب ج 4321",
        notes="صيانة دورية + تغيير زيت وبواجي",
    )
    add_invoice(
        f"SRT-{year}-0001", "sale_return", "customer", 1, "ورشة النور للصيانة", "نقدي",
        [(2, "فلتر زيت أصلي", 1, 12, 25)],
        notes="مرتجع قطعة تالفة من التغليف",
    )
    conn.execute(
        "INSERT INTO advances (employee_id, amount, remaining, date, notes) VALUES (?,?,?,?,?)",
        (2, 800, 300, today, "سلفة طارئة"),
    )
    conn.execute(
        """INSERT INTO salaries (employee_id, period, base_salary, bonus, advance_deduct, net, status, paid_at, notes)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (1, datetime.now().strftime("%Y-%m"), 7500, 500, 0, 8000, "مدفوع", today, "راتب المدير"),
    )


def init_db():
    conn = connect_sqlite()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'مدير'
        );

        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sku TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            category TEXT,
            brand TEXT,
            car_model TEXT,
            unit TEXT DEFAULT 'قطعة',
            cost REAL NOT NULL DEFAULT 0,
            price REAL NOT NULL DEFAULT 0,
            qty REAL NOT NULL DEFAULT 0,
            min_qty REAL NOT NULL DEFAULT 2,
            location TEXT,
            warehouse TEXT,
            aisle TEXT,
            shelf TEXT,
            bin TEXT,
            barcode TEXT,
            year_from TEXT,
            year_to TEXT,
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS customers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            phone TEXT,
            address TEXT,
            tax_no TEXT,
            balance REAL NOT NULL DEFAULT 0,
            points REAL NOT NULL DEFAULT 0,
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS suppliers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            phone TEXT,
            address TEXT,
            tax_no TEXT,
            balance REAL NOT NULL DEFAULT 0,
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            job_title TEXT,
            phone TEXT,
            hire_date TEXT,
            salary REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'نشط',
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS advances (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER NOT NULL REFERENCES employees(id),
            amount REAL NOT NULL,
            remaining REAL NOT NULL,
            date TEXT NOT NULL,
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS salaries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER NOT NULL REFERENCES employees(id),
            period TEXT NOT NULL,
            base_salary REAL NOT NULL,
            bonus REAL NOT NULL DEFAULT 0,
            advance_deduct REAL NOT NULL DEFAULT 0,
            net REAL NOT NULL,
            status TEXT NOT NULL DEFAULT 'مدفوع',
            paid_at TEXT NOT NULL,
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE NOT NULL,
            kind TEXT NOT NULL,
            party_type TEXT,
            party_id INTEGER,
            party_name TEXT,
            date TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'نقدي',
            subtotal REAL NOT NULL DEFAULT 0,
            discount REAL NOT NULL DEFAULT 0,
            tax REAL NOT NULL DEFAULT 0,
            total REAL NOT NULL DEFAULT 0,
            paid REAL NOT NULL DEFAULT 0,
            cost_total REAL NOT NULL DEFAULT 0,
            profit REAL NOT NULL DEFAULT 0,
            vehicle TEXT,
            related_id INTEGER,
            notes TEXT,
            created_by TEXT,
            payment_method TEXT DEFAULT 'نقدي',
            labor_total REAL NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS invoice_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_id INTEGER NOT NULL REFERENCES invoices(id) ON DELETE CASCADE,
            product_id INTEGER,
            description TEXT NOT NULL,
            qty REAL NOT NULL,
            unit_cost REAL NOT NULL DEFAULT 0,
            unit_price REAL NOT NULL DEFAULT 0,
            line_total REAL NOT NULL DEFAULT 0,
            line_cost REAL NOT NULL DEFAULT 0,
            line_profit REAL NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS stock_moves (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER NOT NULL REFERENCES products(id),
            qty REAL NOT NULL,
            unit_cost REAL NOT NULL DEFAULT 0,
            move_type TEXT NOT NULL,
            ref TEXT,
            date TEXT NOT NULL,
            notes TEXT
        );

        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT
        );

        CREATE TABLE IF NOT EXISTS job_cards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE NOT NULL,
            customer_id INTEGER,
            customer_name TEXT,
            phone TEXT,
            plate TEXT,
            vehicle_type TEXT,
            vehicle_model TEXT,
            vehicle_year TEXT,
            complaint TEXT,
            technician_id INTEGER,
            technician_name TEXT,
            status TEXT NOT NULL DEFAULT 'قيد الانتظار',
            opened_at TEXT NOT NULL,
            closed_at TEXT,
            labor_total REAL NOT NULL DEFAULT 0,
            parts_total REAL NOT NULL DEFAULT 0,
            discount REAL NOT NULL DEFAULT 0,
            tax REAL NOT NULL DEFAULT 0,
            total REAL NOT NULL DEFAULT 0,
            invoice_id INTEGER,
            notes TEXT,
            created_by TEXT
        );

        CREATE TABLE IF NOT EXISTS job_parts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER NOT NULL REFERENCES job_cards(id) ON DELETE CASCADE,
            product_id INTEGER,
            description TEXT NOT NULL,
            qty REAL NOT NULL,
            unit_cost REAL NOT NULL DEFAULT 0,
            unit_price REAL NOT NULL DEFAULT 0,
            line_total REAL NOT NULL DEFAULT 0,
            line_cost REAL NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS job_labor (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER NOT NULL REFERENCES job_cards(id) ON DELETE CASCADE,
            employee_id INTEGER,
            service_name TEXT NOT NULL,
            hours REAL NOT NULL DEFAULT 1,
            rate REAL NOT NULL DEFAULT 0,
            amount REAL NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS expenses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category TEXT NOT NULL,
            amount REAL NOT NULL,
            date TEXT NOT NULL,
            paid_to TEXT,
            notes TEXT,
            created_by TEXT
        );

        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            party_type TEXT NOT NULL,
            party_id INTEGER,
            party_name TEXT,
            amount REAL NOT NULL,
            method TEXT NOT NULL DEFAULT 'نقدي',
            date TEXT NOT NULL,
            invoice_id INTEGER,
            notes TEXT
        );
        CREATE TABLE IF NOT EXISTS warehouses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            branch TEXT,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS stock_balances (
            warehouse_id INTEGER NOT NULL,
            product_id INTEGER NOT NULL,
            qty REAL NOT NULL DEFAULT 0,
            PRIMARY KEY (warehouse_id, product_id)
        );
        CREATE TABLE IF NOT EXISTS inventory_lots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER NOT NULL,
            warehouse_id INTEGER,
            qty_remaining REAL NOT NULL DEFAULT 0,
            unit_cost REAL NOT NULL DEFAULT 0,
            received_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS item_units (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER NOT NULL,
            unit_name TEXT NOT NULL,
            conversion_factor REAL NOT NULL DEFAULT 1,
            UNIQUE (product_id, unit_name)
        );
        CREATE TABLE IF NOT EXISTS item_components (
            assembly_id INTEGER NOT NULL,
            component_id INTEGER NOT NULL,
            qty_per_unit REAL NOT NULL DEFAULT 1,
            PRIMARY KEY (assembly_id, component_id)
        );
        CREATE TABLE IF NOT EXISTS stock_transfers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE NOT NULL,
            source_id INTEGER NOT NULL,
            dest_id INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'مسودة',
            tracking_number TEXT,
            requested_by TEXT,
            created_at TEXT NOT NULL,
            shipped_at TEXT,
            received_at TEXT
        );
        CREATE TABLE IF NOT EXISTS stock_transfer_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            transfer_id INTEGER NOT NULL REFERENCES stock_transfers(id) ON DELETE CASCADE,
            product_id INTEGER NOT NULL,
            qty REAL NOT NULL
        );
        CREATE TABLE IF NOT EXISTS stock_counts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            warehouse_id INTEGER NOT NULL,
            count_type TEXT NOT NULL DEFAULT 'دوري',
            status TEXT NOT NULL DEFAULT 'مفتوح',
            counted_by TEXT,
            created_at TEXT NOT NULL,
            closed_at TEXT
        );
        CREATE TABLE IF NOT EXISTS stock_count_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            count_id INTEGER NOT NULL REFERENCES stock_counts(id) ON DELETE CASCADE,
            product_id INTEGER NOT NULL,
            expected_qty REAL NOT NULL DEFAULT 0,
            counted_qty REAL NOT NULL DEFAULT 0,
            difference REAL NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS serial_numbers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER NOT NULL,
            serial_number TEXT UNIQUE NOT NULL,
            status TEXT NOT NULL DEFAULT 'متاح',
            invoice_id INTEGER,
            notes TEXT
        );
        CREATE TABLE IF NOT EXISTS shifts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            started_at TEXT NOT NULL,
            ended_at TEXT,
            opening_cash REAL NOT NULL DEFAULT 0,
            employee_id INTEGER REFERENCES employees(id)
        );
        CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER NOT NULL REFERENCES employees(id),
            attendance_date TEXT NOT NULL,
            shift_id INTEGER REFERENCES shifts(id),
            status TEXT NOT NULL DEFAULT 'حاضر',
            started_at TEXT,
            ended_at TEXT,
            late_minutes INTEGER NOT NULL DEFAULT 0,
            deduction REAL NOT NULL DEFAULT 0,
            reviewed INTEGER NOT NULL DEFAULT 0,
            notes TEXT,
            UNIQUE(employee_id, attendance_date)
        );
        CREATE TABLE IF NOT EXISTS cash_accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            account_type TEXT NOT NULL DEFAULT 'خزينة',
            opening_balance REAL NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS cash_transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            account_id INTEGER NOT NULL,
            transaction_type TEXT NOT NULL,
            amount REAL NOT NULL,
            description TEXT NOT NULL,
            reference_type TEXT,
            reference_id INTEGER,
            transaction_date TEXT NOT NULL,
            created_by TEXT
        );
        CREATE TABLE IF NOT EXISTS journal_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL,
            description TEXT,
            account_name TEXT NOT NULL,
            debit REAL NOT NULL DEFAULT 0,
            credit REAL NOT NULL DEFAULT 0,
            reference_type TEXT,
            reference_id INTEGER,
            created_by TEXT
        );
        CREATE TABLE IF NOT EXISTS party_ledger (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            party_type TEXT NOT NULL,
            party_id INTEGER NOT NULL,
            invoice_id INTEGER,
            entry_type TEXT NOT NULL,
            description TEXT NOT NULL,
            debit REAL NOT NULL DEFAULT 0,
            credit REAL NOT NULL DEFAULT 0,
            due_date TEXT,
            payment_date TEXT,
            created_by TEXT,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS quotes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE NOT NULL,
            date TEXT NOT NULL,
            customer_id INTEGER,
            customer_name TEXT NOT NULL,
            subtotal REAL NOT NULL DEFAULT 0,
            discount REAL NOT NULL DEFAULT 0,
            tax REAL NOT NULL DEFAULT 0,
            total REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'ساري',
            notes TEXT,
            created_by TEXT,
            invoice_id INTEGER
        );
        CREATE TABLE IF NOT EXISTS quote_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            quote_id INTEGER NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,
            product_id INTEGER,
            description TEXT NOT NULL,
            qty REAL NOT NULL,
            unit_price REAL NOT NULL DEFAULT 0,
            line_total REAL NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS campaigns (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            discount_type TEXT NOT NULL DEFAULT 'نسبة',
            discount_value REAL NOT NULL DEFAULT 0,
            starts_on TEXT NOT NULL,
            ends_on TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS coupons (
            code TEXT PRIMARY KEY,
            campaign_id INTEGER,
            discount_type TEXT NOT NULL DEFAULT 'نسبة',
            discount_value REAL NOT NULL DEFAULT 0,
            max_uses INTEGER NOT NULL DEFAULT 1,
            used_count INTEGER NOT NULL DEFAULT 0,
            expires_on TEXT,
            active INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS sales_targets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            target_month TEXT NOT NULL,
            target_amount REAL NOT NULL DEFAULT 0,
            commission_percent REAL NOT NULL DEFAULT 0,
            UNIQUE (username, target_month)
        );
        CREATE TABLE IF NOT EXISTS categories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL
        );
        CREATE TABLE IF NOT EXISTS unit_defs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL
        );
        CREATE TABLE IF NOT EXISTS branches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            code TEXT UNIQUE,
            status TEXT NOT NULL DEFAULT 'نشط',
            is_default INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            address TEXT,
            phone TEXT
        );
        CREATE TABLE IF NOT EXISTS user_branch_permissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            branch_id INTEGER NOT NULL REFERENCES branches(id) ON DELETE CASCADE,
            permissions TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT 'نشط',
            UNIQUE(user_id, branch_id)
        );
        CREATE TABLE IF NOT EXISTS audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT,
            full_name TEXT,
            branch_id INTEGER,
            action TEXT NOT NULL,
            entity_type TEXT NOT NULL,
            entity_id INTEGER,
            entity_ref TEXT,
            details TEXT,
            ip_address TEXT,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS installments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_id INTEGER NOT NULL,
            invoice_id INTEGER,
            installment_no INTEGER NOT NULL,
            due_date TEXT NOT NULL,
            amount REAL NOT NULL,
            paid_amount REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'مستحق'
        );
        CREATE TABLE IF NOT EXISTS checks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_id INTEGER,
            supplier_id INTEGER,
            check_number TEXT NOT NULL,
            check_type TEXT NOT NULL,
            bank_name TEXT,
            amount REAL NOT NULL,
            due_date TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'معلق',
            notes TEXT,
            created_by TEXT,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS daily_closures (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shift_id INTEGER,
            closing_date TEXT NOT NULL,
            expected_cash REAL NOT NULL DEFAULT 0,
            actual_cash REAL NOT NULL DEFAULT 0,
            difference REAL NOT NULL DEFAULT 0,
            closed_by TEXT NOT NULL,
            closed_at TEXT NOT NULL,
            notes TEXT
        );
        CREATE TABLE IF NOT EXISTS sync_devices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            device_id TEXT UNIQUE NOT NULL,
            device_name TEXT,
            platform TEXT,
            last_seen TEXT,
            last_push TEXT,
            last_pull TEXT
        );
        CREATE TABLE IF NOT EXISTS sync_queue (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            op_uuid TEXT UNIQUE NOT NULL,
            device_id TEXT,
            op_type TEXT NOT NULL,
            payload TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            error TEXT,
            created_at TEXT NOT NULL,
            applied_at TEXT
        );
        CREATE TABLE IF NOT EXISTS sync_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            direction TEXT NOT NULL,
            status TEXT NOT NULL,
            message TEXT,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS raw_materials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            unit TEXT NOT NULL DEFAULT 'قطعة',
            cost REAL NOT NULL DEFAULT 0,
            qty REAL NOT NULL DEFAULT 0,
            min_qty REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'فعال',
            notes TEXT
        );
        CREATE TABLE IF NOT EXISTS manufacturing_boms (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER NOT NULL REFERENCES products(id),
            version TEXT NOT NULL DEFAULT '1',
            status TEXT NOT NULL DEFAULT 'فعال',
            labor_cost REAL NOT NULL DEFAULT 0,
            overhead_cost REAL NOT NULL DEFAULT 0,
            notes TEXT,
            UNIQUE(product_id, version)
        );
        CREATE TABLE IF NOT EXISTS manufacturing_bom_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            bom_id INTEGER NOT NULL REFERENCES manufacturing_boms(id) ON DELETE CASCADE,
            raw_material_id INTEGER NOT NULL REFERENCES raw_materials(id),
            qty REAL NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS production_orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            number TEXT UNIQUE NOT NULL,
            product_id INTEGER NOT NULL REFERENCES products(id),
            planned_qty REAL NOT NULL DEFAULT 0,
            completed_qty REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'مفتوح',
            current_stage TEXT,
            opened_at TEXT NOT NULL,
            started_at TEXT,
            completed_at TEXT,
            labor_cost REAL NOT NULL DEFAULT 0,
            overhead_cost REAL NOT NULL DEFAULT 0,
            waste_cost REAL NOT NULL DEFAULT 0,
            warehouse_id INTEGER,
            notes TEXT,
            created_by TEXT
        );
        CREATE TABLE IF NOT EXISTS production_stages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_id INTEGER NOT NULL REFERENCES production_orders(id) ON DELETE CASCADE,
            name TEXT NOT NULL,
            sequence INTEGER NOT NULL DEFAULT 1,
            status TEXT NOT NULL DEFAULT 'معلق',
            started_at TEXT,
            ended_at TEXT,
            notes TEXT
        );
        CREATE TABLE IF NOT EXISTS raw_material_moves (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            raw_material_id INTEGER NOT NULL REFERENCES raw_materials(id),
            qty REAL NOT NULL,
            move_type TEXT NOT NULL,
            reference_type TEXT,
            reference_id INTEGER,
            unit_cost REAL NOT NULL DEFAULT 0,
            date TEXT NOT NULL,
            created_by TEXT
        );
        """
    )
    defaults = {
        "shop_name": APP_NAME,
        "shop_subtitle": "تجارة قطع الغيار وخدمات صيانة السيارات",
        "phone": "0120000000",
        "address": "المملكة العربية السعودية",
        "tax_no": "",
        "print_footer": "شكراً لتعاملكم معنا — الأصلي لقطع الغيار",
        "paper_size": "A4",
        "show_cost": "0",
        "show_profit": "0",
        "print_copies": "1",
        "sku_prefix": "PRD",
        "vat_rate": "15",
        "points_per_100": "1",
        "default_labor_rate": "80",
        "work_start_time": "09:00",
        "attendance_grace_minutes": "15",
        "sync_enabled": "1",
        "sync_server_url": "",
        "sync_token": "alasly-sync-local",
        "device_name": "جهاز المحل",
        "device_id": "",
        "last_sync_at": "",
        "last_sync_status": "",
    }
    for key, value in defaults.items():
        exists = conn.execute("SELECT 1 FROM settings WHERE key=?", (key,)).fetchone()
        if not exists:
            conn.execute("INSERT INTO settings (key, value) VALUES (?,?)", (key, value))
    device_row = conn.execute("SELECT value FROM settings WHERE key='device_id'").fetchone()
    if not device_row or not (device_row[0] or "").strip():
        conn.execute(
            "INSERT INTO settings (key, value) VALUES ('device_id', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (str(uuid.uuid4()),),
        )
    def ensure_columns(table, cols):
        existing = {r[1] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}
        for col, definition in cols:
            if col not in existing:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} {definition}")

    ensure_columns(
        "products",
        (
            ("warehouse", "TEXT"),
            ("aisle", "TEXT"),
            ("shelf", "TEXT"),
            ("bin", "TEXT"),
            ("barcode", "TEXT"),
            ("year_from", "TEXT"),
            ("year_to", "TEXT"),
            ("item_type", "TEXT NOT NULL DEFAULT 'منتج'"),
            ("base_unit", "TEXT NOT NULL DEFAULT 'قطعة'"),
            ("serial_tracking", "INTEGER NOT NULL DEFAULT 0"),
            ("warehouse_id", "INTEGER"),
        ),
    )
    ensure_columns(
        "customers",
        (
            ("points", "REAL NOT NULL DEFAULT 0"),
            ("credit_limit", "REAL NOT NULL DEFAULT 0"),
            ("payment_term_days", "INTEGER NOT NULL DEFAULT 0"),
            ("discount_percent", "REAL NOT NULL DEFAULT 0"),
            ("is_supplier", "INTEGER NOT NULL DEFAULT 0"),
            ("linked_id", "INTEGER"),
        ),
    )
    ensure_columns(
        "suppliers",
        (
            ("is_customer", "INTEGER NOT NULL DEFAULT 0"),
            ("linked_id", "INTEGER"),
            ("credit_limit", "REAL NOT NULL DEFAULT 0"),
            ("payment_term_days", "INTEGER NOT NULL DEFAULT 0"),
            ("discount_percent", "REAL NOT NULL DEFAULT 0"),
        ),
    )
    ensure_columns(
        "users",
        (
            ("job_title", "TEXT"),
            ("permissions", "TEXT NOT NULL DEFAULT 'sales,shift'"),
        ),
    )
    # Branch migration is additive: existing installations keep all legacy columns/data.
    ensure_columns("branches", (("code", "TEXT"), ("status", "TEXT NOT NULL DEFAULT 'نشط'"), ("is_default", "INTEGER NOT NULL DEFAULT 0"), ("created_at", "TEXT")))
    ensure_columns("users", (("branch_id", "INTEGER"),))
    ensure_columns("employees", (("branch_id", "INTEGER"),))
    ensure_columns("invoices", (("branch_id", "INTEGER"),))
    ensure_columns("shifts", (("branch_id", "INTEGER"),))
    ensure_columns("production_orders", (("branch_id", "INTEGER"),))
    ensure_columns("stock_moves", (("branch_id", "INTEGER"),))
    ensure_columns("stock_counts", (("branch_id", "INTEGER"),))
    ensure_columns("quotes", (("branch_id", "INTEGER"),))
    ensure_columns("employees", (("username", "TEXT"),))
    ensure_columns("shifts", (("employee_id", "INTEGER"),))
    ensure_columns(
        "salaries",
        (
            ("attendance_deduct", "REAL NOT NULL DEFAULT 0"),
            ("absence_days", "REAL NOT NULL DEFAULT 0"),
            ("late_minutes", "INTEGER NOT NULL DEFAULT 0"),
        ),
    )
    ensure_columns(
        "invoices",
        (
            ("payment_method", "TEXT DEFAULT 'نقدي'"),
            ("labor_total", "REAL NOT NULL DEFAULT 0"),
            ("due_date", "TEXT"),
            ("salesperson", "TEXT"),
            ("coupon_code", "TEXT"),
            ("cash_account_id", "INTEGER"),
            ("shift_id", "INTEGER"),
        ),
    )
    ensure_columns("invoice_items", (("unit_name", "TEXT"),))
    ensure_columns(
        "payments",
        (("receipt_no", "TEXT"),),
    )
    ensure_columns(
        "daily_closures",
        (
            ("sales_total", "REAL NOT NULL DEFAULT 0"),
            ("sales_cash", "REAL NOT NULL DEFAULT 0"),
            ("returns_total", "REAL NOT NULL DEFAULT 0"),
            ("purchases_total", "REAL NOT NULL DEFAULT 0"),
            ("expenses_total", "REAL NOT NULL DEFAULT 0"),
            ("receipts_total", "REAL NOT NULL DEFAULT 0"),
            ("invoice_count", "INTEGER NOT NULL DEFAULT 0"),
            ("summary_json", "TEXT"),
        ),
    )
    conn.execute("UPDATE products SET barcode = sku WHERE barcode IS NULL OR barcode = ''")
    rows_need_loc = conn.execute(
        "SELECT id, location, warehouse, aisle, shelf, bin FROM products"
    ).fetchall()
    for row in rows_need_loc:
        loc, warehouse, aisle, shelf, bin_code = row[1], row[2], row[3], row[4], row[5]
        if not any([warehouse, aisle, shelf, bin_code]) and loc and "مخزن" not in loc and "-" in loc:
            left, right = loc.split("-", 1)
            warehouse, aisle, shelf = "الرئيسي", left.strip(), right.strip()
            built = format_location(warehouse, aisle, shelf, "")
            conn.execute(
                "UPDATE products SET warehouse=?, aisle=?, shelf=?, location=? WHERE id=?",
                (warehouse, aisle, shelf, built, row[0]),
            )
        else:
            built = format_location(warehouse, aisle, shelf, bin_code)
            if built and not loc:
                conn.execute("UPDATE products SET location=? WHERE id=?", (built, row[0]))
    if conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
        conn.execute(
            "INSERT INTO users (username, password, full_name, role, permissions) VALUES (?,?,?,?,?)",
            ("admin", "admin123", "مدير النظام", "مدير", "all"),
        )
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    if conn.execute("SELECT COUNT(*) FROM warehouses").fetchone()[0] == 0:
        conn.execute(
            "INSERT INTO warehouses (name, branch, created_at) VALUES (?,?,?)",
            ("الرئيسي", "الفرع الرئيسي", now),
        )
        conn.execute(
            "INSERT INTO warehouses (name, branch, created_at) VALUES (?,?,?)",
            ("الفرعي", "الورشة", now),
        )
    if conn.execute("SELECT COUNT(*) FROM cash_accounts").fetchone()[0] == 0:
        for cash_name, cash_type in (
            ("درج", "خزينة"),
            ("فيزا", "بنك"),
            ("إنستاباي", "محفظة"),
            ("فودافون كاش", "محفظة"),
        ):
            conn.execute(
                "INSERT INTO cash_accounts (name, account_type, opening_balance) VALUES (?,?,0)",
                (cash_name, cash_type),
            )
    if conn.execute("SELECT COUNT(*) FROM categories").fetchone()[0] == 0:
        for cat in CATEGORY_PREFIX:
            conn.execute("INSERT OR IGNORE INTO categories (name) VALUES (?)", (cat,))
    if conn.execute("SELECT COUNT(*) FROM unit_defs").fetchone()[0] == 0:
        for uname in ("قطعة", "طقم", "علبة", "زوج", "لتر", "متر"):
            conn.execute("INSERT OR IGNORE INTO unit_defs (name) VALUES (?)", (uname,))
    # Seed exactly the three initial branches; the legacy branch is always the default.
    branch_seed = (("الفرع الرئيسي", "MAIN"), ("فرع الورشة", "WORKSHOP"), ("فرع التصنيع", "MANUFACTURING"))
    for branch_name, branch_code in branch_seed:
        conn.execute(
            "INSERT OR IGNORE INTO branches (name, code, status, is_default, created_at, address, phone) VALUES (?,?,?,?,?,?,?)",
            (branch_name, branch_code, "نشط", 1 if branch_code == "MAIN" else 0, now, "المملكة العربية السعودية", "0120000000"),
        )
    conn.execute("UPDATE branches SET created_at=COALESCE(NULLIF(created_at,''),?)", (now,))
    conn.execute("UPDATE branches SET code=CASE id WHEN (SELECT MIN(id) FROM branches) THEN 'MAIN' ELSE COALESCE(code,'BR-'||id) END WHERE code IS NULL OR code='' ")
    conn.execute("UPDATE branches SET status='نشط' WHERE status IS NULL OR status='' ")
    conn.execute("UPDATE branches SET is_default=0")
    conn.execute("UPDATE branches SET is_default=1 WHERE name='الفرع الرئيسي' OR id=(SELECT MIN(id) FROM branches)")
    main_branch_id = conn.execute("SELECT id FROM branches WHERE is_default=1 ORDER BY id LIMIT 1").fetchone()[0]
    conn.execute(
        """INSERT OR IGNORE INTO user_branch_permissions (user_id, branch_id, permissions, status)
           SELECT id, COALESCE(branch_id, ?), permissions, 'نشط' FROM users""",
        (main_branch_id,),
    )
    # Legacy rows are assigned to the default branch.
    for table in ("users", "employees", "invoices", "quotes", "shifts", "production_orders", "stock_moves", "stock_counts"):
        if conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone()[0]:
            cols = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
            if "branch_id" in cols:
                conn.execute(f"UPDATE {table} SET branch_id=? WHERE branch_id IS NULL", (main_branch_id,))
    conn.execute("UPDATE users SET branch_id=? WHERE branch_id IS NULL", (main_branch_id,))
    conn.execute("UPDATE employees SET branch_id=? WHERE branch_id IS NULL", (main_branch_id,))
    if conn.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 0:
        products = [
            ("BRK-001", "تيل فرامل أمامي", "فرامل", "Brembo", "تويوتا كورولا", "طقم", 85, 140, 24, "A-1"),
            ("FLT-012", "فلتر زيت أصلي", "فلاتر", "Bosch", "هيونداي إلنترا", "قطعة", 12, 25, 60, "B-3"),
            ("FLT-020", "فلتر هواء", "فلاتر", "Mann", "نيسان صني", "قطعة", 18, 35, 40, "B-3"),
            ("OIL-5W30", "زيت محرك 5W30 4 لتر", "زيوت", "Castrol", "عام", "علبة", 55, 85, 36, "C-1"),
            ("BAT-70", "بطارية 70 أمبير", "كهرباء", "ACDelco", "تويوتا كامري", "قطعة", 280, 420, 12, "D-2"),
            ("SPK-04", "بواجي إيريديوم", "محرك", "NGK", "هوندا سيفيك", "طقم", 48, 90, 20, "E-4"),
            ("ALT-210", "دينامو 90 أمبير", "كهرباء", "Denso", "كيا سيراتو", "قطعة", 410, 650, 6, "D-5"),
            ("RAD-33", "رادياتير ألومنيوم", "تبريد", "Nissens", "شيفروليه أفيو", "قطعة", 320, 510, 8, "F-1"),
            ("SHK-FR", "مساعد أمامي", "تعليق", "KYB", "تويوتا يارس", "قطعة", 190, 310, 10, "G-2"),
            ("TIR-16", "إطار 205/55 R16", "إطارات", "Michelin", "عام", "قطعة", 260, 390, 16, "H-1"),
            ("LMP-H4", "لمبة أمامية H4", "إنارة", "Philips", "عام", "زوج", 22, 45, 30, "I-3"),
            ("CLH-KIT", "طقم كلاتش كامل", "ناقل حركة", "Valeo", "هيونداي أكسنت", "طقم", 540, 820, 5, "J-1"),
        ]
        conn.executemany(
            """INSERT INTO products (sku, name, category, brand, car_model, unit, cost, price, qty, location)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            products,
        )
        customers = [
            ("ورشة النور للصيانة", "0501112233", "جدة - حي الصفا", 0),
            ("مؤسسة الطريق الذهبي", "0553344556", "مكة - العزيزية", 850),
            ("أحمد سعيد الغامدي", "0547788990", "جدة - أبحر", 0),
            ("شركة أسطول النقل", "0126655443", "جدة - الصناعية", 2100),
        ]
        conn.executemany(
            "INSERT INTO customers (name, phone, address, balance) VALUES (?,?,?,?)",
            customers,
        )
        suppliers = [
            ("الوكلاء المتحدة لقطع الغيار", "0126112233", "جدة - الصناعية الأولى", 0),
            ("مستودع الشرق للزيوت", "0126889900", "جدة - الخمرة", 4200),
            ("شركة الإطارات الحديثة", "0114556677", "الرياض - السلي", 0),
        ]
        conn.executemany(
            "INSERT INTO suppliers (name, phone, address, balance) VALUES (?,?,?,?)",
            suppliers,
        )
        employees = [
            ("خالد عبدالعزيز", "مدير المحل", "0551112233", "2021-03-01", 7500),
            ("ماجد العتيبي", "فني صيانة", "0552223344", "2022-06-15", 4800),
            ("يوسف الحربي", "أمين مخزن", "0553334455", "2023-01-10", 4200),
            ("سعاد العمري", "محاسبة", "0554445566", "2022-11-01", 5200),
        ]
        conn.executemany(
            "INSERT INTO employees (name, job_title, phone, hire_date, salary) VALUES (?,?,?,?,?)",
            employees,
        )
        seed_demo_documents(conn)
    conn.commit()
    conn.close()


def product_map():
    return {str(r["id"]): r for r in query("SELECT * FROM products ORDER BY name")}


def party_name(kind, party_id):
    table = "customers" if kind in ("sale", "sale_return", "maintenance") else "suppliers"
    row = query(f"SELECT name FROM {table} WHERE id=?", (party_id,), one=True)
    return row["name"] if row else ""


def apply_stock(product_id, qty_delta, unit_cost, move_type, ref, notes=""):
    product = query("SELECT * FROM products WHERE id=?", (product_id,), one=True)
    if not product:
        return
    new_qty = float(product["qty"]) + float(qty_delta)
    if new_qty < 0:
        raise ValueError(f"الكمية غير كافية للمنتج: {product['name']}")
    new_cost = float(product["cost"])
    if qty_delta > 0 and unit_cost is not None:
        old_value = float(product["qty"]) * float(product["cost"])
        add_value = float(qty_delta) * float(unit_cost)
        new_cost = (old_value + add_value) / new_qty if new_qty else unit_cost
        execute(
            "UPDATE products SET qty=?, cost=? WHERE id=?",
            (new_qty, round(new_cost, 4), product_id),
        )
        add_lot(product_id, qty_delta, unit_cost)
    else:
        execute("UPDATE products SET qty=? WHERE id=?", (new_qty, product_id))
        if qty_delta < 0:
            consume_lots(product_id, abs(qty_delta))
    execute(
        """INSERT INTO stock_moves (product_id, qty, unit_cost, move_type, ref, date, notes)
           VALUES (?,?,?,?,?,?,?)""",
        (
            product_id,
            qty_delta,
            unit_cost or product["cost"],
            move_type,
            ref,
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            notes,
        ),
    )


def parse_items(form):
    ids = form.getlist("item_product_id[]")
    descs = form.getlist("item_desc[]")
    qtys = form.getlist("item_qty[]")
    prices = form.getlist("item_price[]")
    costs = form.getlist("item_cost[]")
    items = []
    for i, desc in enumerate(descs):
        desc = (desc or "").strip()
        if not desc and not (ids[i] if i < len(ids) else ""):
            continue
        qty = float(qtys[i] or 0)
        if qty <= 0:
            continue
        product_id = ids[i] if i < len(ids) and ids[i] else None
        price = float(prices[i] or 0)
        cost = float(costs[i] or 0)
        if product_id:
            prod = query("SELECT * FROM products WHERE id=?", (product_id,), one=True)
            if prod:
                desc = desc or prod["name"]
                if cost <= 0:
                    cost = float(prod["cost"])
                if price <= 0:
                    price = float(prod["price"])
        line_total = round(qty * price, 2)
        line_cost = round(qty * cost, 2)
        items.append(
            {
                "product_id": int(product_id) if product_id else None,
                "description": desc or "بند",
                "qty": qty,
                "unit_cost": cost,
                "unit_price": price,
                "line_total": line_total,
                "line_cost": line_cost,
                "line_profit": round(line_total - line_cost, 2),
            }
        )
    return items


def save_invoice(kind, form, related_id=None):
    items = parse_items(form)
    if not items:
        raise ValueError("أضف بنداً واحداً على الأقل")
    prefixes = {
        "sale": "SAL",
        "sale_return": "SRT",
        "purchase": "PUR",
        "purchase_return": "PRT",
        "maintenance": "MNT",
    }
    party_type = "customer" if kind in ("sale", "sale_return", "maintenance") else "supplier"
    party_id = form.get("party_id") or None
    name = form.get("party_name") or form.get("party_search")
    if party_id and not name:
        name = party_name(kind, party_id)
    subtotal = sum(i["line_total"] for i in items)
    coupon_extra, coupon_code = apply_coupon_code(form.get("coupon_code"), subtotal)
    discount = float(form.get("discount") or 0) + coupon_extra
    if party_id and party_type == "customer":
        cust = query("SELECT * FROM customers WHERE id=?", (party_id,), one=True)
        if cust and float(cust["discount_percent"] or 0) > 0 and not coupon_extra:
            discount += round(subtotal * float(cust["discount_percent"]) / 100.0, 2)
    tax = float(form.get("tax") or 0)
    total = round(subtotal - discount + tax, 2)
    cost_total = sum(i["line_cost"] for i in items)
    profit = round(subtotal - discount - cost_total, 2)
    if kind in ("sale_return", "purchase"):
        profit = round(-profit if kind == "sale_return" else 0, 2)
    if kind == "purchase_return":
        profit = 0
    payment_method = form.get("payment_method") or form.get("status") or "نقدي"
    status = "آجل" if payment_method in ("آجل", "شيك", "أقساط") else payment_method
    paid = total if payment_method not in ("آجل", "شيك", "أقساط") else float(form.get("paid") or 0)
    due_days = 0
    if party_id and party_type == "customer":
        cust = query("SELECT * FROM customers WHERE id=?", (party_id,), one=True)
        if cust:
            due_days = int(cust["payment_term_days"] or 0)
            limit = float(cust["credit_limit"] or 0)
            if payment_method in ("آجل", "أقساط") and limit > 0 and float(cust["balance"] or 0) + (total - paid) > limit:
                raise ValueError("تجاوز العميل حد الائتمان")
    due_date = form.get("due_date") or (
        (date.today() + timedelta(days=due_days)).isoformat() if due_days else None
    )
    number = next_number(prefixes[kind])
    labor_total = float(form.get("labor_total") or 0)
    salesperson = form.get("salesperson") or session.get("user")
    cash_acc = cash_account_for(payment_method)
    active_shift = current_shift()
    inv_id = execute(
        """INSERT INTO invoices
           (number, kind, party_type, party_id, party_name, date, status, subtotal, discount, tax,
            total, paid, cost_total, profit, vehicle, related_id, notes, created_by, payment_method, labor_total,
            due_date, salesperson, coupon_code, cash_account_id, shift_id, branch_id)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            number,
            kind,
            party_type,
            party_id,
            name,
            form.get("date") or date.today().isoformat(),
            status,
            subtotal,
            discount,
            tax,
            total,
            paid,
            cost_total,
            profit,
            form.get("vehicle"),
            related_id,
            form.get("notes"),
            session.get("user"),
            payment_method,
            labor_total,
            due_date,
            salesperson,
            coupon_code,
            cash_acc["id"] if cash_acc else None,
            active_shift["id"] if active_shift else None,
            current_branch_id(),
        ),
    )
    for item in items:
        execute(
            """INSERT INTO invoice_items
               (invoice_id, product_id, description, qty, unit_cost, unit_price, line_total, line_cost, line_profit)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (
                inv_id,
                item["product_id"],
                item["description"],
                item["qty"],
                item["unit_cost"],
                item["unit_price"],
                item["line_total"],
                item["line_cost"],
                item["line_profit"],
            ),
        )
        if item["product_id"]:
            if kind in ("sale", "maintenance"):
                apply_stock(item["product_id"], -item["qty"], item["unit_cost"], kind, number, item["description"])
            elif kind == "purchase":
                apply_stock(item["product_id"], item["qty"], item["unit_cost"], kind, number, item["description"])
            elif kind == "sale_return":
                apply_stock(item["product_id"], item["qty"], item["unit_cost"], kind, number, item["description"])
            elif kind == "purchase_return":
                apply_stock(item["product_id"], -item["qty"], item["unit_cost"], kind, number, item["description"])
    due = total - paid
    if party_id and due:
        if party_type == "customer":
            sign = 1 if kind in ("sale", "maintenance") else -1
            execute("UPDATE customers SET balance = balance + ? WHERE id=?", (sign * due, party_id))
            add_ledger(
                "customer",
                party_id,
                "فاتورة" if sign > 0 else "مرتجع",
                f"{KIND_LABELS.get(kind, kind)} {number}",
                debit=due if sign > 0 else 0,
                credit=due if sign < 0 else 0,
                invoice_id=inv_id,
                due_date=due_date,
            )
        else:
            sign = 1 if kind == "purchase" else -1
            execute("UPDATE suppliers SET balance = balance + ? WHERE id=?", (sign * due, party_id))
            add_ledger(
                "supplier",
                party_id,
                "فاتورة" if sign > 0 else "مرتجع",
                f"{KIND_LABELS.get(kind, kind)} {number}",
                debit=due if sign < 0 else 0,
                credit=due if sign > 0 else 0,
                invoice_id=inv_id,
                due_date=due_date,
            )
    if kind in ("sale", "maintenance") and party_id:
        award_points(party_id, total)
    if kind in ("sale", "maintenance"):
        add_journal(f"مبيعات {number}", "المبيعات", credit=total, reference_type="invoice", reference_id=inv_id)
        add_journal(f"تكلفة {number}", "تكلفة البضاعة", debit=cost_total, reference_type="invoice", reference_id=inv_id)
        add_journal(f"مخزون {number}", "المخزون", credit=cost_total, reference_type="invoice", reference_id=inv_id)
        if paid:
            post_cash(payment_method, paid, f"تحصيل {number}", "invoice", inv_id, "in")
            add_journal(f"تحصيل {number}", payment_method, debit=paid, reference_type="invoice", reference_id=inv_id)
        if due:
            add_journal(f"حسابات {number}", "العملاء", debit=due, reference_type="invoice", reference_id=inv_id)
    elif kind == "purchase":
        add_journal(f"مشتريات {number}", "المخزون", debit=total, reference_type="invoice", reference_id=inv_id)
        if paid:
            post_cash(payment_method, paid, f"سداد مشتريات {number}", "invoice", inv_id, "out")
            add_journal(f"سداد {number}", payment_method, credit=paid, reference_type="invoice", reference_id=inv_id)
        if due:
            add_journal(f"موردون {number}", "الموردون", credit=due, reference_type="invoice", reference_id=inv_id)
    elif kind == "sale_return":
        if paid:
            post_cash(payment_method, paid, f"رد مبلغ {number}", "invoice", inv_id, "out")
    elif kind == "purchase_return":
        if paid:
            post_cash(payment_method, paid, f"مرتجع شراء {number}", "invoice", inv_id, "in")
    inst_count = int(form.get("installments_count") or 0)
    if kind in ("sale", "maintenance") and party_id and inst_count > 1:
        per = round(total / inst_count, 2)
        start = date.today()
        for n in range(1, inst_count + 1):
            due_d = (start + timedelta(days=30 * n)).isoformat()
            amt = per if n < inst_count else round(total - per * (inst_count - 1), 2)
            execute(
                """INSERT INTO installments (customer_id, invoice_id, installment_no, due_date, amount, status)
                   VALUES (?,?,?,?,?,'مستحق')""",
                (party_id, inv_id, n, due_d, amt),
            )
    audit_log("إنشاء", "فاتورة", inv_id, number, f"نوع المستند: {KIND_LABELS.get(kind, kind)}، الإجمالي: {total:,.2f}")
    return inv_id


def cancel_invoice(inv_id):
    inv_sql, inv_args = branch_filter("invoices", include_all=True)
    inv = query("SELECT * FROM invoices WHERE id=?" + inv_sql, [inv_id] + inv_args, one=True)
    if not inv:
        raise ValueError("المستند غير موجود")
    if inv["status"] == "ملغاة":
        raise ValueError("المستند ملغى مسبقاً")
    items = query("SELECT * FROM invoice_items WHERE invoice_id=?", (inv_id,))
    kind = inv["kind"]
    number = inv["number"]
    for item in items:
        if not item["product_id"]:
            continue
        qty = float(item["qty"] or 0)
        cost = float(item["unit_cost"] or 0)
        if kind in ("sale", "maintenance"):
            apply_stock(item["product_id"], qty, cost, "cancel", number, "إلغاء فاتورة")
        elif kind == "purchase":
            apply_stock(item["product_id"], -qty, cost, "cancel", number, "إلغاء فاتورة")
        elif kind == "sale_return":
            apply_stock(item["product_id"], -qty, cost, "cancel", number, "إلغاء مرتجع")
        elif kind == "purchase_return":
            apply_stock(item["product_id"], qty, cost, "cancel", number, "إلغاء مرتجع شراء")
    paid = float(inv["paid"] or 0)
    total = float(inv["total"] or 0)
    due = total - paid
    party_id = inv["party_id"]
    party_type = inv["party_type"]
    if party_id and due:
        if party_type == "customer":
            sign = 1 if kind in ("sale", "maintenance") else -1
            execute("UPDATE customers SET balance = balance - ? WHERE id=?", (sign * due, party_id))
            add_ledger(
                "customer",
                party_id,
                "إلغاء",
                f"إلغاء {KIND_LABELS.get(kind, kind)} {number}",
                debit=due if sign < 0 else 0,
                credit=due if sign > 0 else 0,
                invoice_id=inv_id,
            )
        else:
            sign = 1 if kind == "purchase" else -1
            execute("UPDATE suppliers SET balance = balance - ? WHERE id=?", (sign * due, party_id))
            add_ledger(
                "supplier",
                party_id,
                "إلغاء",
                f"إلغاء {KIND_LABELS.get(kind, kind)} {number}",
                debit=due if sign > 0 else 0,
                credit=due if sign < 0 else 0,
                invoice_id=inv_id,
            )
    method = inv["payment_method"] or inv["status"] or "نقدي"
    if paid:
        if kind in ("sale", "maintenance"):
            post_cash(method, paid, f"إلغاء تحصيل {number}", "invoice_cancel", inv_id, "out")
        elif kind == "purchase":
            post_cash(method, paid, f"إلغاء سداد {number}", "invoice_cancel", inv_id, "in")
        elif kind == "sale_return":
            post_cash(method, paid, f"إلغاء رد مبلغ {number}", "invoice_cancel", inv_id, "in")
        elif kind == "purchase_return":
            post_cash(method, paid, f"إلغاء مرتجع شراء {number}", "invoice_cancel", inv_id, "out")
    if kind in ("sale", "maintenance") and party_id:
        rate = float(get_setting("points_per_100", "1") or 1)
        pts = round((total / 100.0) * rate, 2)
        if pts:
            execute("UPDATE customers SET points = MAX(COALESCE(points,0) - ?, 0) WHERE id=?", (pts, party_id))
    execute("UPDATE invoices SET status='ملغاة', notes=COALESCE(notes,'') || ? WHERE id=?", (f" | أُلغي بواسطة {session.get('user')}", inv_id))
    audit_log("إلغاء", "فاتورة", inv_id, number, f"إلغاء {KIND_LABELS.get(kind, kind)}")
    add_journal(f"إلغاء {number}", "إلغاء مستندات", debit=0, credit=0, reference_type="invoice_cancel", reference_id=inv_id)
    return inv


def day_summary(day=None):
    day = day or date.today().isoformat()
    like = day + "%"
    sales = query(
        """SELECT COALESCE(SUM(total),0) t, COALESCE(SUM(paid),0) p, COUNT(*) c
           FROM invoices WHERE kind IN ('sale','maintenance') AND status!='ملغاة' AND date LIKE ?""",
        (like,),
        one=True,
    )
    sale_returns = query(
        "SELECT COALESCE(SUM(total),0) t FROM invoices WHERE kind='sale_return' AND status!='ملغاة' AND date LIKE ?",
        (like,),
        one=True,
    )["t"]
    purchases = query(
        "SELECT COALESCE(SUM(total),0) t FROM invoices WHERE kind='purchase' AND status!='ملغاة' AND date LIKE ?",
        (like,),
        one=True,
    )["t"]
    purchase_returns = query(
        "SELECT COALESCE(SUM(total),0) t FROM invoices WHERE kind='purchase_return' AND status!='ملغاة' AND date LIKE ?",
        (like,),
        one=True,
    )["t"]
    expenses_total = query(
        "SELECT COALESCE(SUM(amount),0) t FROM expenses WHERE date LIKE ?",
        (like,),
        one=True,
    )["t"]
    receipts_in = query(
        "SELECT COALESCE(SUM(amount),0) t FROM payments WHERE party_type='customer' AND date LIKE ?",
        (like,),
        one=True,
    )["t"]
    receipts_out = query(
        "SELECT COALESCE(SUM(amount),0) t FROM payments WHERE party_type='supplier' AND date LIKE ?",
        (like,),
        one=True,
    )["t"]
    cash_in = query(
        """SELECT COALESCE(SUM(t.amount),0) t FROM cash_transactions t
           JOIN cash_accounts a ON a.id=t.account_id
           WHERE a.name='درج' AND t.transaction_type='وارد' AND t.transaction_date LIKE ?""",
        (like,),
        one=True,
    )["t"]
    cash_out = query(
        """SELECT COALESCE(SUM(t.amount),0) t FROM cash_transactions t
           JOIN cash_accounts a ON a.id=t.account_id
           WHERE a.name='درج' AND t.transaction_type='صادر' AND t.transaction_date LIKE ?""",
        (like,),
        one=True,
    )["t"]
    drawer = query(
        """SELECT a.opening_balance + COALESCE((
             SELECT SUM(CASE WHEN t.transaction_type='وارد' THEN t.amount ELSE -t.amount END)
             FROM cash_transactions t WHERE t.account_id=a.id),0) balance
           FROM cash_accounts a WHERE a.name='درج'""",
        one=True,
    )
    invoices = query(
        """SELECT id, number, kind, party_name, total, paid, status, date
           FROM invoices WHERE date LIKE ? ORDER BY id""",
        (like,),
    )
    return {
        "day": day,
        "sales_total": float(sales["t"] or 0),
        "sales_cash": float(sales["p"] or 0),
        "invoice_count": int(sales["c"] or 0),
        "returns_total": float(sale_returns or 0),
        "purchases_total": float(purchases or 0),
        "purchase_returns": float(purchase_returns or 0),
        "expenses_total": float(expenses_total or 0),
        "receipts_total": float(receipts_in or 0),
        "payments_total": float(receipts_out or 0),
        "cash_in": float(cash_in or 0),
        "cash_out": float(cash_out or 0),
        "drawer_balance": float(drawer["balance"] if drawer else 0),
        "net_cash": float(cash_in or 0) - float(cash_out or 0),
        "invoices": invoices,
    }


KIND_LABELS = {
    "sale": "فاتورة بيع / كاشير",
    "sale_return": "مرتجع بيع",
    "purchase": "فاتورة شراء",
    "purchase_return": "مرتجع شراء",
    "maintenance": "فاتورة صيانة",
}

JOB_STATUSES = ("قيد الانتظار", "جاري العمل", "تم الإصلاح", "تم التسليم")
PAY_METHODS = ("نقدي", "فيزا", "محفظة إلكترونية", "إنستاباي", "فودافون كاش", "شيك", "أقساط", "آجل")


def recalc_job(job_id):
    parts = query("SELECT COALESCE(SUM(line_total),0) v FROM job_parts WHERE job_id=?", (job_id,), one=True)["v"]
    labor = query("SELECT COALESCE(SUM(amount),0) v FROM job_labor WHERE job_id=?", (job_id,), one=True)["v"]
    job = query("SELECT * FROM job_cards WHERE id=?", (job_id,), one=True)
    discount = float(job["discount"] or 0)
    tax = float(job["tax"] or 0)
    total = round(float(parts) + float(labor) - discount + tax, 2)
    execute(
        "UPDATE job_cards SET parts_total=?, labor_total=?, total=? WHERE id=?",
        (parts, labor, total, job_id),
    )


def award_points(customer_id, total):
    if not customer_id:
        return
    rate = float(get_setting("points_per_100", "1") or 1)
    pts = round((float(total) / 100.0) * rate, 2)
    if pts:
        execute("UPDATE customers SET points = COALESCE(points,0) + ? WHERE id=?", (pts, customer_id))


def current_branch():
    """Return the active branch; administrators may select it, legacy sessions fall back safely."""
    try:
        if not has_table_column("branches", "id"):
            return None
        branch_id = session.get("branch_id")
        if session.get("role") != "مدير":
            user = query("SELECT branch_id FROM users WHERE username=?", (session.get("user"),), one=True) if session.get("user") else None
            branch_id = (user["branch_id"] if user and user["branch_id"] else branch_id) if user else branch_id
        branch = query("SELECT * FROM branches WHERE id=? AND status!='معطل'", (branch_id,), one=True) if branch_id else None
        if branch:
            return branch
        return query("SELECT * FROM branches WHERE is_default=1 AND status!='معطل' ORDER BY id LIMIT 1", one=True)
    except sqlite3.Error:
        return None

def current_branch_id():
    branch = current_branch()
    return branch["id"] if branch else None


def audit_log(action, entity_type, entity_id=None, entity_ref=None, details=""):
    if not session.get("user"):
        return
    execute(
        """INSERT INTO audit_log (username, full_name, branch_id, action, entity_type, entity_id, entity_ref, details, ip_address, created_at)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (
            session.get("user"), session.get("full_name"), current_branch_id(), action,
            entity_type, entity_id, entity_ref, details, request.remote_addr if request else "",
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ),
    )

def has_table_column(table, column):
    try:
        return any(r[1] == column for r in db().execute(f"PRAGMA table_info({table})").fetchall())
    except sqlite3.Error:
        return False

def branch_filter(alias=None, include_all=False):
    """Return SQL predicate/params while remaining compatible with pre-migration databases."""
    schema_table = {"s": "shifts", "v": "invoices", "e": "employees"}.get(alias, alias or "invoices")
    if not has_table_column(schema_table, "branch_id"):
        return "", []
    if include_all and session.get("role") == "مدير" and request.args.get("branch_id") in ("all", ""):
        return "", []
    bid = current_branch_id()
    if bid is None:
        return "", []
    prefix = (alias + ".") if alias else ""
    return f" AND {prefix}branch_id=?", [bid]


def current_shift(username=None):
    username = username or session.get("user")
    if not username:
        return None
    sql = "SELECT * FROM shifts WHERE username=? AND ended_at IS NULL"
    args = [username]
    extra, extra_args = branch_filter("shifts")
    return query(sql + extra + " ORDER BY id DESC LIMIT 1", args + extra_args, one=True)


def current_employee(username=None):
    username = username or session.get("user")
    if not username:
        return None
    extra, extra_args = branch_filter("employees")
    return query("SELECT * FROM employees WHERE username=? AND status='نشط'" + extra, [username] + extra_args, one=True)


def attendance_preview(employee_id, period):
    emp = query("SELECT salary FROM employees WHERE id=?", (employee_id,), one=True)
    if not emp:
        return {"deduct": 0.0, "absence_days": 0, "late_minutes": 0, "rows": []}
    daily = float(emp["salary"] or 0) / 30.0
    rows = query(
        "SELECT * FROM attendance WHERE employee_id=? AND attendance_date LIKE ? ORDER BY attendance_date",
        (employee_id, period + "%"),
    )
    absence_days = sum(1 for r in rows if r["status"] == "غائب")
    late_minutes = sum(int(r["late_minutes"] or 0) for r in rows)
    deduct = round(absence_days * daily + (late_minutes / 480.0) * daily, 2)
    return {"deduct": deduct, "absence_days": absence_days, "late_minutes": late_minutes, "rows": rows}


def current_late_notice():
    employee = current_employee()
    if not employee:
        return None
    return query(
        """SELECT * FROM attendance WHERE employee_id=? AND attendance_date=? AND late_minutes>0
           ORDER BY id DESC LIMIT 1""",
        (employee["id"], date.today().isoformat()), one=True,
    )


def add_journal(description, account_name, debit=0, credit=0, reference_type=None, reference_id=None):
    execute(
        """INSERT INTO journal_entries
           (date, description, account_name, debit, credit, reference_type, reference_id, created_by)
           VALUES (?,?,?,?,?,?,?,?)""",
        (
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            description,
            account_name,
            float(debit or 0),
            float(credit or 0),
            reference_type,
            reference_id,
            session.get("user"),
        ),
    )


def add_ledger(party_type, party_id, entry_type, description, debit=0, credit=0, invoice_id=None, due_date=None, payment_date=None):
    if not party_id:
        return
    execute(
        """INSERT INTO party_ledger
           (party_type, party_id, invoice_id, entry_type, description, debit, credit, due_date, payment_date, created_by, created_at)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (
            party_type,
            party_id,
            invoice_id,
            entry_type,
            description,
            float(debit or 0),
            float(credit or 0),
            due_date,
            payment_date,
            session.get("user"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
        ),
    )


def cash_account_for(method):
    mapping = {
        "نقدي": "درج",
        "فيزا": "فيزا",
        "محفظة إلكترونية": "إنستاباي",
        "إنستاباي": "إنستاباي",
        "فودافون كاش": "فودافون كاش",
        "آجل": None,
        "شيك": None,
    }
    name = mapping.get(method, method)
    if not name:
        return None
    row = query("SELECT * FROM cash_accounts WHERE name=?", (name,), one=True)
    return row


def post_cash(method, amount, description, reference_type=None, reference_id=None, direction="in"):
    amount = float(amount or 0)
    if amount <= 0:
        return
    acc = cash_account_for(method)
    if not acc:
        return
    execute(
        """INSERT INTO cash_transactions
           (account_id, transaction_type, amount, description, reference_type, reference_id, transaction_date, created_by)
           VALUES (?,?,?,?,?,?,?,?)""",
        (
            acc["id"],
            "وارد" if direction == "in" else "صادر",
            amount,
            description,
            reference_type,
            reference_id,
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            session.get("user"),
        ),
    )


def default_warehouse_id():
    row = query("SELECT id FROM warehouses ORDER BY id LIMIT 1", one=True)
    return row["id"] if row else None


def add_lot(product_id, qty, unit_cost, warehouse_id=None):
    execute(
        """INSERT INTO inventory_lots (product_id, warehouse_id, qty_remaining, unit_cost, received_at)
           VALUES (?,?,?,?,?)""",
        (
            product_id,
            warehouse_id or default_warehouse_id(),
            qty,
            unit_cost,
            datetime.now().strftime("%Y-%m-%d %H:%M"),
        ),
    )


def consume_lots(product_id, qty):
    remaining = float(qty)
    lots = query(
        "SELECT * FROM inventory_lots WHERE product_id=? AND qty_remaining>0 ORDER BY id",
        (product_id,),
    )
    cost_used = 0
    for lot in lots:
        if remaining <= 0:
            break
        take = min(float(lot["qty_remaining"]), remaining)
        execute(
            "UPDATE inventory_lots SET qty_remaining=qty_remaining-? WHERE id=?",
            (take, lot["id"]),
        )
        cost_used += take * float(lot["unit_cost"])
        remaining -= take
    if remaining > 0:
        prod = query("SELECT cost FROM products WHERE id=?", (product_id,), one=True)
        cost_used += remaining * float(prod["cost"] if prod else 0)
    return cost_used


def apply_coupon_code(code, subtotal):
    code = (code or "").strip()
    if not code:
        return 0, None
    coupon = query("SELECT * FROM coupons WHERE code=?", (code,), one=True)
    if not coupon or not coupon["active"]:
        raise ValueError("كود الخصم غير صالح")
    if coupon["expires_on"] and coupon["expires_on"] < date.today().isoformat():
        raise ValueError("كود الخصم منتهي")
    if int(coupon["used_count"] or 0) >= int(coupon["max_uses"] or 1):
        raise ValueError("تم استهلاك كود الخصم بالكامل")
    today = date.today().isoformat()
    if coupon["campaign_id"]:
        camp = query("SELECT * FROM campaigns WHERE id=?", (coupon["campaign_id"],), one=True)
        if camp and (not camp["active"] or camp["starts_on"] > today or camp["ends_on"] < today):
            raise ValueError("الحملة غير سارية")
    if coupon["discount_type"] == "نسبة":
        amount = round(float(subtotal) * float(coupon["discount_value"]) / 100.0, 2)
    else:
        amount = float(coupon["discount_value"])
    execute("UPDATE coupons SET used_count=used_count+1 WHERE code=?", (code,))
    return amount, code


def rows_to_dicts(rows):
    return [{k: r[k] for k in r.keys()} for r in rows]


PRODUCT_IMPORT_COLUMNS = (
    ("sku", "كود الصنف"),
    ("name", "الاسم"),
    ("category", "التصنيف"),
    ("brand", "الماركة"),
    ("car_model", "الموديلات المتوافقة"),
    ("unit", "الوحدة"),
    ("cost", "تكلفة الشراء"),
    ("price", "سعر البيع"),
    ("qty", "الكمية"),
    ("min_qty", "حد إعادة الطلب"),
    ("barcode", "الباركود"),
    ("warehouse", "المخزن"),
    ("aisle", "الممر"),
    ("shelf", "الرف"),
    ("bin", "الموضع"),
    ("year_from", "سنة من"),
    ("year_to", "سنة إلى"),
    ("item_type", "نوع الصنف"),
    ("base_unit", "الوحدة الأساسية"),
    ("notes", "ملاحظات"),
)

PRODUCT_HEADER_MAP = {label: key for key, label in PRODUCT_IMPORT_COLUMNS}
PRODUCT_HEADER_MAP.update({key: key for key, _ in PRODUCT_IMPORT_COLUMNS})
PRODUCT_HEADER_MAP.update({
    "الكود": "sku",
    "كود": "sku",
    "الصنف": "name",
    "اسم الصنف": "name",
    "الاسم": "name",
    "التكلفة": "cost",
    "السعر": "price",
    "سعر البيع": "price",
    "الكمية": "qty",
    "حد الطلب": "min_qty",
    "الماركة": "brand",
    "الموقع": "location",
})


def workbook_response(workbook, filename):
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_products_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "الأصناف"
    ws.sheet_view.rightToLeft = True
    headers = [label for _, label in PRODUCT_IMPORT_COLUMNS]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 18
    for r in rows:
        ws.append([
            r["sku"] or "",
            r["name"] or "",
            r["category"] or "",
            r["brand"] or "",
            r["car_model"] or "",
            r["unit"] or "",
            float(r["cost"] or 0),
            float(r["price"] or 0),
            float(r["qty"] or 0),
            float(r["min_qty"] or 0),
            r["barcode"] or "",
            r["warehouse"] or "",
            r["aisle"] or "",
            r["shelf"] or "",
            r["bin"] or "",
            r["year_from"] or "",
            r["year_to"] or "",
            r["item_type"] or "منتج",
            r["base_unit"] or "قطعة",
            r["notes"] or "",
        ])
    ws.freeze_panes = "A2"
    return wb


def read_products_workbook(file_storage):
    filename = (file_storage.filename or "").lower()
    if filename.endswith(".csv"):
        import csv
        raw = file_storage.read().decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(raw))
        rows = []
        for row in reader:
            rows.append({(k or "").strip(): (v or "").strip() for k, v in row.items()})
        return rows
    if not XLSX_OK:
        raise RuntimeError("دعم ملفات Excel غير متاح على الخادم. استخدم ملف CSV.")
    wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True)
    ws = wb.active
    data = list(ws.iter_rows(values_only=True))
    if not data:
        return []
    header = [str(c).strip() if c is not None else "" for c in data[0]]
    rows = []
    for line in data[1:]:
        if not any(v not in (None, "") for v in line):
            continue
        item = {}
        for i, h in enumerate(header):
            if not h:
                continue
            item[h] = line[i] if i < len(line) else None
        rows.append(item)
    return rows


def import_products(rows):
    added, updated, skipped = 0, 0, []
    for idx, raw in enumerate(rows, start=2):
        mapped = {}
        for header, value in raw.items():
            key = PRODUCT_HEADER_MAP.get(str(header).strip())
            if key:
                mapped[key] = value

        def text(key):
            v = mapped.get(key)
            if v is None:
                return ""
            return str(v).strip()

        def number(key):
            v = mapped.get(key)
            if v in (None, ""):
                return 0.0
            try:
                return float(str(v).replace(",", "").strip())
            except ValueError:
                return 0.0

        name = text("name")
        if not name:
            skipped.append(f"سطر {idx}: بدون اسم")
            continue
        sku = text("sku")
        warehouse = text("warehouse")
        aisle = text("aisle")
        shelf = text("shelf")
        bin_code = text("bin")
        location = format_location(warehouse, aisle, shelf, bin_code) or text("location")
        existing = None
        if sku:
            existing = query("SELECT id, sku FROM products WHERE sku=?", (sku,), one=True)
        if not existing:
            existing = query(
                "SELECT id, sku FROM products WHERE LOWER(TRIM(name))=LOWER(TRIM(?)) ORDER BY id LIMIT 1",
                (name,),
                one=True,
            )
        if not sku:
            sku = existing["sku"] if existing else next_sku(text("category"))
        values = (
            name,
            text("category"),
            text("brand"),
            text("car_model"),
            text("unit") or text("base_unit") or "قطعة",
            number("cost"),
            number("price"),
            number("qty"),
            number("min_qty"),
            location,
            warehouse,
            aisle,
            shelf,
            bin_code,
            text("barcode") or sku,
            text("year_from"),
            text("year_to"),
            text("item_type") or "منتج",
            text("base_unit") or "قطعة",
            text("notes"),
        )
        try:
            if existing:
                execute(
                    """UPDATE products SET name=?, category=?, brand=?, car_model=?, unit=?,
                       cost=?, price=?, qty=?, min_qty=?, location=?, warehouse=?, aisle=?, shelf=?, bin=?,
                       barcode=?, year_from=?, year_to=?, item_type=?, base_unit=?, notes=? WHERE id=?""",
                    values + (existing["id"],),
                )
                updated += 1
            else:
                execute(
                    """INSERT INTO products (name, category, brand, car_model, unit, cost, price, qty, min_qty,
                       location, warehouse, aisle, shelf, bin, barcode, year_from, year_to, item_type, base_unit, notes, sku)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    values + (sku,),
                )
                added += 1
        except sqlite3.IntegrityError:
            skipped.append(f"سطر {idx}: كود مكرر {sku}")
    return added, updated, skipped


def snapshot_payload():
    tables = (
        "products",
        "customers",
        "suppliers",
        "employees",
        "invoices",
        "invoice_items",
        "job_cards",
        "expenses",
        "settings",
        "cash_accounts",
        "warehouses",
        "coupons",
        "campaigns",
    )
    data = {}
    for table in tables:
        try:
            data[table] = rows_to_dicts(query(f"SELECT * FROM {table}"))
        except sqlite3.Error:
            data[table] = []
    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "device_id": get_setting("device_id"),
        "device_name": get_setting("device_name", "جهاز المحل"),
        "data": data,
    }


def apply_sync_op(op):
    op_type = op.get("op_type") or op.get("type")
    payload = op.get("payload") or {}
    if isinstance(payload, str):
        payload = json.loads(payload)
    if op_type == "sale":
        class Form(dict):
            def getlist(self, key):
                return self.get(key, [])
        form = Form(payload)
        form.setdefault("payment_method", payload.get("payment_method") or "نقدي")
        form.setdefault("party_name", payload.get("party_name") or "عميل نقدي")
        save_invoice(payload.get("kind") or "sale", form)
        return "فاتورة"
    if op_type == "expense":
        execute(
            "INSERT INTO expenses (category, amount, date, paid_to, notes, created_by) VALUES (?,?,?,?,?,?)",
            (
                payload.get("category") or "أخرى",
                float(payload.get("amount") or 0),
                payload.get("date") or date.today().isoformat(),
                payload.get("paid_to"),
                payload.get("notes") or "مزامنة بدون اتصال",
                payload.get("created_by") or session.get("user") or "offline",
            ),
        )
        return "مصروف"
    if op_type == "customer":
        execute(
            "INSERT INTO customers (name, phone, address, tax_no, notes, credit_limit, payment_term_days, discount_percent) VALUES (?,?,?,?,?,?,?,?)",
            (
                payload.get("name"),
                payload.get("phone"),
                payload.get("address"),
                payload.get("tax_no"),
                payload.get("notes"),
                float(payload.get("credit_limit") or 0),
                int(payload.get("payment_term_days") or 0),
                float(payload.get("discount_percent") or 0),
            ),
        )
        return "عميل"
    if op_type == "product":
        sku = payload.get("sku") or next_sku(payload.get("category"))
        execute(
            """INSERT INTO products (sku, name, category, brand, car_model, unit, cost, price, qty, min_qty, location, warehouse, aisle, shelf, bin, barcode, notes)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                sku,
                payload.get("name"),
                payload.get("category"),
                payload.get("brand"),
                payload.get("car_model"),
                payload.get("unit") or "قطعة",
                float(payload.get("cost") or 0),
                float(payload.get("price") or 0),
                float(payload.get("qty") or 0),
                float(payload.get("min_qty") or 2),
                payload.get("location"),
                payload.get("warehouse"),
                payload.get("aisle"),
                payload.get("shelf"),
                payload.get("bin"),
                payload.get("barcode") or sku,
                payload.get("notes"),
            ),
        )
        return "صنف"
    if op_type == "job":
        number = next_number("JOB", "job_cards")
        execute(
            """INSERT INTO job_cards (number, customer_name, phone, plate, vehicle_type, vehicle_model, complaint, status, opened_at, notes, created_by)
               VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            (
                number,
                payload.get("customer_name"),
                payload.get("phone"),
                payload.get("plate"),
                payload.get("vehicle_type"),
                payload.get("vehicle_model"),
                payload.get("complaint"),
                "قيد الانتظار",
                datetime.now().strftime("%Y-%m-%d %H:%M"),
                payload.get("notes") or "مزامنة بدون اتصال",
                payload.get("created_by") or session.get("user") or "offline",
            ),
        )
        return "أمر شغل"
    raise ValueError("نوع عملية غير مدعوم")


def log_sync(direction, status, message):
    execute(
        "INSERT INTO sync_log (direction, status, message, created_at) VALUES (?,?,?,?)",
        (direction, status, message, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )


def nav_active(endpoint):
    return "active" if request.endpoint == endpoint else ""


def quick_stats():
    try:
        return {
            "stock_value": query("SELECT COALESCE(SUM(qty*cost),0) v FROM products", one=True)["v"],
            "customers_due": query("SELECT COALESCE(SUM(balance),0) v FROM customers", one=True)["v"],
            "suppliers_due": query("SELECT COALESCE(SUM(balance),0) v FROM suppliers", one=True)["v"],
            "low_stock": query("SELECT COUNT(*) c FROM products WHERE qty<=min_qty", one=True)["c"],
        }
    except sqlite3.Error:
        return {}


@app.context_processor
def inject():
    return {
        "APP_NAME": APP_NAME,
        "quick_stats": quick_stats() if session.get("user") else {},
        "PERMISSION_MODULES": PERMISSION_MODULES,
        "can": user_has_permission,
        "KIND_LABELS": KIND_LABELS,
        "nav_active": nav_active,
        "today": date.today().isoformat(),
        "user": session.get("user"),
        "role": session.get("role"),
        "current_branch": current_branch() if session.get("user") else None,
        "branches": accessible_branches_for_user() if session.get("user") else [],
        "kind": (request.view_args or {}).get("kind"),
        "print_settings": all_settings() if session.get("user") else {},
        "open_shift": current_shift() if session.get("user") else None,
        "late_notice": current_late_notice() if session.get("user") else None,
        "sync_cfg": {
            "enabled": get_setting("sync_enabled", "1"),
            "server": get_setting("sync_server_url", ""),
            "token": get_setting("sync_token", "alasly-sync-local"),
            "device_id": get_setting("device_id", ""),
            "device_name": get_setting("device_name", "جهاز المحل"),
            "last_sync_at": get_setting("last_sync_at", ""),
            "last_sync_status": get_setting("last_sync_status", ""),
        } if session.get("user") else {},
    }


@app.route("/health")
def health():
    return jsonify({"ok": True, "status": "ok"}), 200


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = query(
            "SELECT * FROM users WHERE username=? AND password=?",
            (request.form.get("username"), request.form.get("password")),
            one=True,
        )
        if user:
            session.permanent = True
            session["user"] = user["username"]
            session["full_name"] = user["full_name"]
            session["role"] = user["role"]
            session["user_id"] = user["id"]
            session["permissions"] = user["permissions"] or ""
            session["branch_id"] = user["branch_id"] if "branch_id" in user.keys() and user["branch_id"] else (query("SELECT id FROM branches WHERE is_default=1 ORDER BY id LIMIT 1", one=True)["id"])
            return redirect(url_for("dashboard"))
        flash("بيانات الدخول غير صحيحة", "err")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def dashboard():
    sales = query(
        "SELECT COALESCE(SUM(total),0) s, COALESCE(SUM(profit),0) p FROM invoices WHERE kind IN ('sale','maintenance') AND status!='ملغاة'",
        one=True,
    )
    purchases = query(
        "SELECT COALESCE(SUM(total),0) s FROM invoices WHERE kind='purchase' AND status!='ملغاة'",
        one=True,
    )
    stock_value = query("SELECT COALESCE(SUM(qty*cost),0) v FROM products", one=True)["v"]
    customers_due = query("SELECT COALESCE(SUM(balance),0) v FROM customers", one=True)["v"]
    suppliers_due = query("SELECT COALESCE(SUM(balance),0) v FROM suppliers", one=True)["v"]
    payroll = query("SELECT COALESCE(SUM(net),0) v FROM salaries", one=True)["v"]
    expenses_total = query("SELECT COALESCE(SUM(amount),0) v FROM expenses", one=True)["v"]
    open_jobs = query("SELECT COUNT(*) c FROM job_cards WHERE status NOT IN ('تم التسليم')", one=True)["c"]
    low = query("SELECT * FROM products WHERE qty <= min_qty ORDER BY qty ASC LIMIT 8")
    top_parts = query(
        """SELECT p.name, p.sku, SUM(i.qty) qty, SUM(i.line_total) sales, SUM(i.line_profit) profit
           FROM invoice_items i JOIN invoices v ON v.id=i.invoice_id
           JOIN products p ON p.id=i.product_id
           WHERE v.kind IN ('sale','maintenance') AND v.status!='ملغاة'
           GROUP BY p.id ORDER BY qty DESC LIMIT 6"""
    )
    recent = query("SELECT * FROM invoices ORDER BY id DESC LIMIT 8")
    month = datetime.now().strftime("%Y-%m")
    month_sales = query(
        "SELECT COALESCE(SUM(total),0) s, COALESCE(SUM(profit),0) p FROM invoices WHERE kind IN ('sale','maintenance') AND date LIKE ?",
        (month + "%",),
        one=True,
    )
    month_purchases = query(
        "SELECT COALESCE(SUM(total),0) s FROM invoices WHERE kind='purchase' AND date LIKE ?",
        (month + "%",),
        one=True,
    )
    month_returns = query(
        "SELECT COALESCE(SUM(total),0) s FROM invoices WHERE kind IN ('sale_return','purchase_return') AND date LIKE ?",
        (month + "%",),
        one=True,
    )
    by_kind = query(
        """SELECT kind, COUNT(*) c, COALESCE(SUM(total),0) t, COALESCE(SUM(profit),0) p
           FROM invoices WHERE status!='ملغاة' GROUP BY kind"""
    )

    # --- Analytical data for the main dashboard screen ---
    total_items = query("SELECT COUNT(*) c FROM products", one=True)["c"]
    suppliers_count = query("SELECT COUNT(*) c FROM suppliers", one=True)["c"]
    customers_count = query("SELECT COUNT(*) c FROM customers", one=True)["c"]
    total_qty = query("SELECT COALESCE(SUM(qty),0) q FROM products", one=True)["q"]
    low_count = query("SELECT COUNT(*) c FROM products WHERE qty<=min_qty", one=True)["c"]
    out_count = query("SELECT COUNT(*) c FROM products WHERE qty<=0", one=True)["c"]

    months_ar = [
        "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
        "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر",
    ]
    inv_series = query(
        """SELECT substr(date,1,7) m,
                  COALESCE(SUM(CASE WHEN kind IN ('sale','maintenance') THEN total END),0) sales,
                  COALESCE(SUM(CASE WHEN kind='purchase' THEN total END),0) purchases
           FROM invoices WHERE status!='ملغاة' GROUP BY m ORDER BY m"""
    )
    moves_series = query(
        """SELECT substr(date,1,7) m,
                  COALESCE(SUM(CASE WHEN qty>0 THEN qty END),0) stock_in,
                  COALESCE(SUM(CASE WHEN qty<0 THEN -qty END),0) stock_out
           FROM stock_moves GROUP BY m ORDER BY m"""
    )
    inv_map = {r["m"]: r for r in inv_series}
    move_map = {r["m"]: r for r in moves_series}
    chart_labels, chart_sales, chart_purchases, chart_in, chart_out = [], [], [], [], []
    now = datetime.now()
    for offset in range(11, -1, -1):
        y, m = now.year, now.month - offset
        while m <= 0:
            m += 12
            y -= 1
        key = f"{y:04d}-{m:02d}"
        inv_row = inv_map.get(key)
        move_row = move_map.get(key)
        chart_labels.append(months_ar[m - 1])
        chart_sales.append(round(inv_row["sales"], 2) if inv_row else 0)
        chart_purchases.append(round(inv_row["purchases"], 2) if inv_row else 0)
        chart_in.append(round(move_row["stock_in"], 2) if move_row else 0)
        chart_out.append(round(move_row["stock_out"], 2) if move_row else 0)

    top_debtors = query(
        "SELECT name, balance FROM customers WHERE balance>0 ORDER BY balance DESC LIMIT 6"
    )
    top_creditors = query(
        "SELECT name, balance FROM suppliers WHERE balance>0 ORDER BY balance DESC LIMIT 6"
    )
    cash_balance = query(
        "SELECT COALESCE(SUM(CASE WHEN transaction_type='وارد' THEN amount ELSE -amount END),0) v FROM cash_transactions",
        one=True,
    )["v"]

    chart_data = {
        "labels": chart_labels,
        "sales": chart_sales,
        "purchases": chart_purchases,
        "stockIn": chart_in,
        "stockOut": chart_out,
        "status": {
            "labels": ["متوفر", "منخفض", "نفد"],
            "values": [
                max(total_items - low_count, 0),
                max(low_count - out_count, 0),
                out_count,
            ],
        },
    }

    return render_template(
        "dashboard.html",
        sales=sales,
        purchases=purchases,
        stock_value=stock_value,
        customers_due=customers_due,
        suppliers_due=suppliers_due,
        payroll=payroll,
        low=low,
        recent=recent,
        month_sales=month_sales,
        month_purchases=month_purchases,
        month_returns=month_returns,
        by_kind=by_kind,
        expenses_total=expenses_total,
        open_jobs=open_jobs,
        top_parts=top_parts,
        total_items=total_items,
        suppliers_count=suppliers_count,
        customers_count=customers_count,
        total_qty=total_qty,
        low_count=low_count,
        out_count=out_count,
        top_debtors=top_debtors,
        top_creditors=top_creditors,
        cash_balance=cash_balance,
        chart_data=chart_data,
        client_code=get_setting("client_code", "102"),
    )


def save_product_from_form():
    pid = request.form.get("id")
    sku = (request.form.get("sku") or "").strip()
    if not sku:
        sku = next_sku(request.form.get("category"))
    warehouse = (request.form.get("warehouse") or "").strip()
    aisle = (request.form.get("aisle") or "").strip()
    shelf = (request.form.get("shelf") or "").strip()
    bin_code = (request.form.get("bin") or "").strip()
    location = format_location(warehouse, aisle, shelf, bin_code) or (request.form.get("location") or "").strip()
    data = (
        sku,
        request.form["name"].strip(),
        request.form.get("category"),
        request.form.get("brand"),
        request.form.get("car_model"),
        request.form.get("unit") or "قطعة",
        float(request.form.get("cost") or 0),
        float(request.form.get("price") or 0),
        float(request.form.get("qty") or 0),
        float(request.form.get("min_qty") or 2),
        location,
        warehouse,
        aisle,
        shelf,
        bin_code,
        (request.form.get("barcode") or sku).strip(),
        request.form.get("year_from"),
        request.form.get("year_to"),
        request.form.get("notes"),
        request.form.get("item_type") or "منتج",
        request.form.get("base_unit") or request.form.get("unit") or "قطعة",
        1 if request.form.get("serial_tracking") else 0,
    )
    try:
        if pid:
            execute(
                """UPDATE products SET sku=?, name=?, category=?, brand=?, car_model=?, unit=?,
                   cost=?, price=?, qty=?, min_qty=?, location=?, warehouse=?, aisle=?, shelf=?, bin=?,
                   barcode=?, year_from=?, year_to=?, notes=?, item_type=?, base_unit=?, serial_tracking=? WHERE id=?""",
                data + (pid,),
            )
            audit_log("تعديل", "صنف", pid, sku, f"تعديل بيانات الصنف والموقع: {location or 'غير محدد'}")
            flash("تم تحديث الصنف", "ok")
        else:
            new_id = execute(
                """INSERT INTO products (sku, name, category, brand, car_model, unit, cost, price, qty, min_qty, location, warehouse, aisle, shelf, bin, barcode, year_from, year_to, notes, item_type, base_unit, serial_tracking)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                data,
            )
            audit_log("إنشاء", "صنف", new_id, sku, f"إضافة الصنف وموقعه: {location or 'غير محدد'}")
            flash(f"تم إضافة الصنف بالكود {sku}", "ok")
    except sqlite3.IntegrityError:
        flash("رقم الصنف موجود مسبقاً", "err")
    return redirect(url_for("inventory"))


@app.route("/products", methods=["GET", "POST"])
@login_required
def products():
    if request.method == "POST":
        return save_product_from_form()
    return redirect(url_for("inventory", **request.args))


@app.route("/products/delete/<int:pid>")
@login_required
def delete_product(pid):
    product = query("SELECT name, sku FROM products WHERE id=?", (pid,), one=True)
    if not product:
        flash("الصنف غير موجود", "err")
        return redirect(url_for("inventory"))
    try:
        execute("DELETE FROM products WHERE id=?", (pid,))
        audit_log("حذف", "صنف", pid, product["sku"], f"حذف الصنف: {product['name']}")
        flash("تم حذف الصنف", "ok")
    except sqlite3.IntegrityError:
        flash("لا يمكن حذف الصنف لأنه مرتبط بحركات أو فواتير؛ استخدم تعديلًا أو اجعله غير نشط", "err")
    return redirect(url_for("inventory"))


@app.route("/inventory", methods=["GET", "POST"])
@login_required
def inventory():
    if request.method == "POST":
        return save_product_from_form()
    q = request.args.get("q", "").strip()
    sql = "SELECT * FROM products"
    args = []
    if q:
        sql += """ WHERE name LIKE ? OR sku LIKE ? OR brand LIKE ? OR car_model LIKE ?
                   OR IFNULL(location,'') LIKE ? OR IFNULL(warehouse,'') LIKE ?
                   OR IFNULL(aisle,'') LIKE ? OR IFNULL(shelf,'') LIKE ? OR IFNULL(bin,'') LIKE ?
                   OR IFNULL(barcode,'') LIKE ?"""
        args = [f"%{q}%"] * 10
    sql += " ORDER BY name"
    rows = query(sql, args)
    stats_rows = query(
        """SELECT i.product_id,
                  AVG(CASE WHEN v.kind='purchase' THEN i.unit_cost END) avg_purchase,
                  AVG(CASE WHEN v.kind IN ('sale','maintenance') THEN i.unit_price END) avg_sale,
                  MIN(CASE WHEN v.kind IN ('sale','maintenance') AND i.unit_price>0 THEN i.unit_price END) min_sale
           FROM invoice_items i JOIN invoices v ON v.id=i.invoice_id
           WHERE v.status!='ملغاة' AND i.product_id IS NOT NULL
           GROUP BY i.product_id"""
    )
    stats = {s["product_id"]: s for s in stats_rows}
    edit = None
    if request.args.get("edit"):
        edit = query("SELECT * FROM products WHERE id=?", (request.args["edit"],), one=True)
    moves = query(
        """SELECT m.*, p.name, p.sku FROM stock_moves m
           JOIN products p ON p.id=m.product_id ORDER BY m.id DESC LIMIT 40"""
    )
    value = query("SELECT COALESCE(SUM(qty*cost),0) v FROM products", one=True)["v"]
    warehouses = query("SELECT * FROM warehouses ORDER BY name")
    lots = query(
        """SELECT l.*, p.name, p.sku FROM inventory_lots l
           JOIN products p ON p.id=l.product_id WHERE l.qty_remaining>0 ORDER BY l.id DESC LIMIT 20"""
    )
    transfers = query(
        """SELECT t.*, s.name src, d.name dst FROM stock_transfers t
           JOIN warehouses s ON s.id=t.source_id JOIN warehouses d ON d.id=t.dest_id
           ORDER BY t.id DESC LIMIT 15"""
    )
    counts = query(
        """SELECT c.*, w.name warehouse_name FROM stock_counts c
           JOIN warehouses w ON w.id=c.warehouse_id ORDER BY c.id DESC LIMIT 10"""
    )
    return render_template(
        "inventory.html",
        rows=rows,
        moves=moves,
        value=value,
        edit=edit,
        q=q,
        warehouses=warehouses,
        lots=lots,
        transfers=transfers,
        counts=counts,
        stats=stats,
    )


def barcode_pdf_response(products, quantities):
    if not REPORTLAB_OK or code128 is None:
        flash("طباعة الباركود غير متاحة حاليًا على الخادم", "err")
        return redirect(url_for("barcode_center"))
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BarcodeArabic", fontName="DejaVu", fontSize=8, leading=10, alignment=TA_RIGHT))
    cells = []
    for product in products:
        value = str(product["barcode"] or product["sku"] or product["id"]).strip()
        count = max(1, min(int(quantities.get(str(product["id"]), 1) or 1), 500))
        for _ in range(count):
            bc = code128.Code128(value, barHeight=28, barWidth=0.75)
            cells.append([bc, Paragraph(rtl_pdf(product["name"]), styles["BarcodeArabic"]), Paragraph(rtl_pdf(value), styles["BarcodeArabic"])])
    if not cells:
        flash("اختر صنفًا واحدًا على الأقل", "err")
        return redirect(url_for("barcode_center"))
    rows = []
    for i in range(0, len(cells), 2):
        row = []
        for cell in cells[i:i + 2]:
            row.append(Table([cell], colWidths=[250], rowHeights=[62], style=TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccd6df"))])))
        if len(row) == 1:
            row.append("")
        rows.append(row)
    doc.build([Paragraph(rtl_pdf("مركز طباعة باركود الأصناف"), styles["Title"]), Spacer(1, 12), Table(rows, colWidths=[260, 260], hAlign="CENTER")])
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name="barcodes.pdf", mimetype="application/pdf")


@app.route("/inventory/barcodes", methods=["GET", "POST"])
@login_required
def barcode_center():
    if request.method == "POST":
        ids = [int(x) for x in request.form.getlist("product_ids") if str(x).isdigit()]
        products = query("SELECT id,name,sku,barcode FROM products WHERE id IN (%s) ORDER BY name" % ",".join("?" * len(ids)), ids) if ids else []
        return barcode_pdf_response(products, request.form)
    q = (request.args.get("q") or "").strip()
    sql = "SELECT id,name,sku,barcode,qty FROM products WHERE (item_type IS NULL OR item_type!='خدمة')"
    args = []
    if q:
        sql += " AND (name LIKE ? OR sku LIKE ? OR IFNULL(barcode,'') LIKE ? OR IFNULL(brand,'') LIKE ? OR IFNULL(location,'') LIKE ?)"
        args.extend([f"%{q}%"] * 5)
    products = query(sql + " ORDER BY name", args)
    return render_template("barcode_center.html", products=products, q=q)


@app.route("/inventory/<int:pid>/barcode.pdf")
@login_required
def product_barcode_pdf(pid):
    product = query("SELECT id,name,sku,barcode FROM products WHERE id=?", (pid,), one=True)
    if not product:
        flash("الصنف غير موجود", "err")
        return redirect(url_for("inventory"))
    return barcode_pdf_response([product], {str(pid): request.args.get("qty", "1")})


@app.route("/inventory/shortages")
@login_required
def inventory_shortages():
    q = (request.args.get("q") or "").strip()
    status = request.args.get("status") or "all"
    sql = """SELECT * FROM products WHERE qty <= min_qty"""
    args = []
    if status == "out":
        sql += " AND qty <= 0"
    elif status == "low":
        sql += " AND qty > 0"
    if q:
        sql += """ AND (name LIKE ? OR sku LIKE ? OR brand LIKE ? OR car_model LIKE ?
                    OR IFNULL(category,'') LIKE ? OR IFNULL(warehouse,'') LIKE ?
                    OR IFNULL(barcode,'') LIKE ?)"""
        args.extend([f"%{q}%"] * 7)
    sql += " ORDER BY CASE WHEN qty<=0 THEN 0 ELSE 1 END, (min_qty - qty) DESC, name"
    rows = query(sql, args)
    out_count = query("SELECT COUNT(*) c FROM products WHERE qty<=0", one=True)["c"]
    low_count = query("SELECT COUNT(*) c FROM products WHERE qty>0 AND qty<=min_qty", one=True)["c"]
    needed_value = 0.0
    needed_qty = 0.0
    for r in rows:
        need = max(float(r["min_qty"] or 0) - float(r["qty"] or 0), 0)
        needed_qty += need
        needed_value += need * float(r["cost"] or 0)
    return render_template(
        "shortages.html",
        rows=rows,
        q=q,
        status=status,
        out_count=out_count,
        low_count=low_count,
        total_count=out_count + low_count,
        needed_qty=needed_qty,
        needed_value=needed_value,
    )


@app.route("/inventory/export")
@login_required
def inventory_export():
    q = request.args.get("q", "").strip()
    sql = "SELECT * FROM products"
    args = []
    if q:
        sql += """ WHERE name LIKE ? OR sku LIKE ? OR brand LIKE ? OR car_model LIKE ?
                   OR IFNULL(location,'') LIKE ? OR IFNULL(barcode,'') LIKE ?"""
        args = [f"%{q}%"] * 6
    sql += " ORDER BY name"
    rows = query(sql, args)
    if not XLSX_OK:
        output = io.StringIO()
        headers = [label for _, label in PRODUCT_IMPORT_COLUMNS]
        output.write(",".join(headers) + "\n")
        for r in rows:
            values = [str(r[key] if r[key] is not None else "") for key, _ in PRODUCT_IMPORT_COLUMNS]
            output.write(",".join('"' + v.replace('"', '""') + '"' for v in values) + "\n")
        data = io.BytesIO(output.getvalue().encode("utf-8-sig"))
        return send_file(data, as_attachment=True, download_name="الأصناف.csv", mimetype="text/csv")
    wb = build_products_workbook(rows)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    return workbook_response(wb, f"الاصناف-{stamp}.xlsx")


@app.route("/inventory/template")
@login_required
def inventory_template():
    headers = [label for _, label in PRODUCT_IMPORT_COLUMNS]
    if not XLSX_OK:
        data = io.BytesIO((",".join(headers) + "\n").encode("utf-8-sig"))
        return send_file(data, as_attachment=True, download_name="قالب-الأصناف.csv", mimetype="text/csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "قالب الأصناف"
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = 18
    example = [
        "BRK-1001", "تيل فرامل خلفي", "فرامل", "Bosch", "تويوتا كورولا 2018", "طقم",
        90, 150, 20, 5, "6281000000", "الرئيسي", "A", "3", "يمين", "2015", "2024", "منتج", "قطعة", "مثال",
    ]
    ws.append(example)
    ws.freeze_panes = "A2"
    return workbook_response(wb, "قالب-الأصناف.xlsx")


@app.route("/inventory/import", methods=["POST"])
@login_required
def inventory_import():
    file = request.files.get("file")
    if not file or not file.filename:
        flash("اختر ملف الأصناف أولاً", "err")
        return redirect(url_for("inventory"))
    if file.filename.lower().endswith((".xlsx", ".xlsm")) and not XLSX_OK:
        flash("دعم Excel غير متاح، استخدم ملف CSV", "err")
        return redirect(url_for("inventory"))
    try:
        rows = read_products_workbook(file)
    except Exception as exc:
        flash(f"تعذر قراءة الملف: {exc}", "err")
        return redirect(url_for("inventory"))
    if not rows:
        flash("الملف لا يحتوي بيانات", "err")
        return redirect(url_for("inventory"))
    if request.form.get("mode") == "replace":
        execute("DELETE FROM products")
    added, updated, skipped = import_products(rows)
    msg = f"تمت الاستيراد: إضافة {added}، تحديث {updated}"
    if skipped:
        msg += f"، تخطي {len(skipped)}"
    flash(msg, "ok" if (added or updated) else "err")
    if skipped:
        flash(" · ".join(skipped[:5]), "err")
    return redirect(url_for("inventory"))


@app.route("/inventory/adjust", methods=["POST"])
@login_required
def inventory_adjust():
    pid = int(request.form["product_id"])
    qty = float(request.form["qty"])
    notes = request.form.get("notes") or "تسوية مخزون"
    product = query("SELECT * FROM products WHERE id=?", (pid,), one=True)
    delta = qty - float(product["qty"])
    apply_stock(pid, delta, product["cost"], "تسوية", "ADJ", notes)
    flash("تم تعديل المخزون", "ok")
    return redirect(url_for("inventory"))


@app.route("/inventory/location", methods=["POST"])
@login_required
def inventory_location():
    pid = int(request.form["product_id"])
    warehouse = (request.form.get("warehouse") or "").strip()
    aisle = (request.form.get("aisle") or "").strip()
    shelf = (request.form.get("shelf") or "").strip()
    bin_code = (request.form.get("bin") or "").strip()
    location = format_location(warehouse, aisle, shelf, bin_code)
    execute(
        "UPDATE products SET warehouse=?, aisle=?, shelf=?, bin=?, location=? WHERE id=?",
        (warehouse, aisle, shelf, bin_code, location, pid),
    )
    flash("تم حفظ مكان الصنف", "ok")
    return redirect(url_for("inventory"))


def sync_linked_party(source_table, source_id, target_table, want_link, form, extra):
    """Create/update the counterpart record when a party is both customer and supplier."""
    source = query(f"SELECT * FROM {source_table} WHERE id=?", (source_id,), one=True)
    if not source:
        return
    existing = None
    if source["linked_id"]:
        existing = query(
            f"SELECT * FROM {target_table} WHERE id=?", (source["linked_id"],), one=True
        )
    if not existing:
        existing = query(
            f"SELECT * FROM {target_table} WHERE name=? AND IFNULL(phone,'')=? AND IFNULL(linked_id,0)=0",
            (source["name"], source["phone"] or ""),
            one=True,
        )
    contact = (
        source["name"],
        source["phone"],
        source["address"],
        source["tax_no"],
    )
    if want_link:
        if existing:
            execute(
                f"UPDATE {target_table} SET name=?, phone=?, address=?, tax_no=?, {extra}=1 WHERE id=?",
                contact + (existing["id"],),
            )
            target_id = existing["id"]
        else:
            target_id = execute(
                f"INSERT INTO {target_table} (name, phone, address, tax_no, {extra}) VALUES (?,?,?,?,1)",
                contact,
            )
        execute(
            f"UPDATE {target_table} SET linked_id=? WHERE id=?", (source_id, target_id)
        )
        execute(
            f"UPDATE {source_table} SET linked_id=?, {extra}=1 WHERE id=?", (target_id, source_id)
        )
    else:
        execute(f"UPDATE {source_table} SET {extra}=0 WHERE id=?", (source_id,))
        if existing:
            execute(f"UPDATE {target_table} SET linked_id=NULL, {extra}=0 WHERE id=?", (existing["id"],))


@app.route("/customers", methods=["GET", "POST"])
@login_required
def customers():
    if request.method == "POST":
        cid = request.form.get("id")
        is_supplier = 1 if request.form.get("is_supplier") else 0
        data = (
            request.form["name"].strip(),
            request.form.get("phone"),
            request.form.get("address"),
            request.form.get("tax_no"),
            request.form.get("notes"),
            float(request.form.get("credit_limit") or 0),
            int(request.form.get("payment_term_days") or 0),
            float(request.form.get("discount_percent") or 0),
        )
        if cid:
            execute(
                "UPDATE customers SET name=?, phone=?, address=?, tax_no=?, notes=?, credit_limit=?, payment_term_days=?, discount_percent=? WHERE id=?",
                data + (cid,),
            )
            sync_linked_party("customers", int(cid), "suppliers", is_supplier, request.form, "is_customer")
            flash("تم تحديث العميل", "ok")
        else:
            new_id = execute(
                "INSERT INTO customers (name, phone, address, tax_no, notes, credit_limit, payment_term_days, discount_percent) VALUES (?,?,?,?,?,?,?,?)",
                data,
            )
            sync_linked_party("customers", new_id, "suppliers", is_supplier, request.form, "is_customer")
            flash("تم إضافة العميل", "ok")
        return redirect(url_for("customers"))
    rows = query("SELECT * FROM customers ORDER BY name")
    edit = None
    if request.args.get("edit"):
        edit = query("SELECT * FROM customers WHERE id=?", (request.args["edit"],), one=True)
    return render_template(
        "parties.html", title="العملاء", endpoint="customers", rows=rows, edit=edit,
        dual_label="يُعامل أيضاً كمورد (له حساب مشتريات)",
    )


@app.route("/customers/delete/<int:cid>")
@login_required
def delete_customer(cid):
    execute("DELETE FROM customers WHERE id=?", (cid,))
    flash("تم حذف العميل", "ok")
    return redirect(url_for("customers"))


@app.route("/suppliers", methods=["GET", "POST"])
@login_required
def suppliers():
    if request.method == "POST":
        sid = request.form.get("id")
        is_customer = 1 if request.form.get("is_customer") else 0
        data = (
            request.form["name"].strip(),
            request.form.get("phone"),
            request.form.get("address"),
            request.form.get("tax_no"),
            request.form.get("notes"),
            float(request.form.get("credit_limit") or 0),
            int(request.form.get("payment_term_days") or 0),
            float(request.form.get("discount_percent") or 0),
        )
        if sid:
            execute(
                "UPDATE suppliers SET name=?, phone=?, address=?, tax_no=?, notes=?, credit_limit=?, payment_term_days=?, discount_percent=? WHERE id=?",
                data + (sid,),
            )
            sync_linked_party("suppliers", int(sid), "customers", is_customer, request.form, "is_supplier")
            flash("تم تحديث المورد", "ok")
        else:
            new_id = execute(
                "INSERT INTO suppliers (name, phone, address, tax_no, notes, credit_limit, payment_term_days, discount_percent) VALUES (?,?,?,?,?,?,?,?)",
                data,
            )
            sync_linked_party("suppliers", new_id, "customers", is_customer, request.form, "is_supplier")
            flash("تم إضافة المورد", "ok")
        return redirect(url_for("suppliers"))
    rows = query("SELECT * FROM suppliers ORDER BY name")
    edit = None
    if request.args.get("edit"):
        edit = query("SELECT * FROM suppliers WHERE id=?", (request.args["edit"],), one=True)
    return render_template(
        "parties.html", title="الموردون", endpoint="suppliers", rows=rows, edit=edit,
        dual_label="يُعامل أيضاً كعميل (له حساب مبيعات)",
    )


@app.route("/suppliers/delete/<int:sid>")
@login_required
def delete_supplier(sid):
    execute("DELETE FROM suppliers WHERE id=?", (sid,))
    flash("تم حذف المورد", "ok")
    return redirect(url_for("suppliers"))


@app.route("/branches", methods=["GET", "POST"])
@login_required
@admin_required
def branches():
    if request.method == "POST":
        action = request.form.get("action", "save")
        try:
            if action == "assign_user":
                execute("UPDATE users SET branch_id=? WHERE id=?", (request.form.get("branch_id") or None, request.form.get("user_id")))
                uid = int(request.form.get("user_id") or 0)
                bid = int(request.form.get("branch_id") or 0)
                target_perms = ",".join(k for k in request.form.getlist("branch_permissions") if k in PERMISSION_KEYS)
                execute("INSERT INTO user_branch_permissions (user_id, branch_id, permissions, status) VALUES (?,?,?,'نشط') ON CONFLICT(user_id, branch_id) DO UPDATE SET permissions=excluded.permissions, status='نشط'", (uid, bid, target_perms))
                flash("تم ربط المستخدم بالفرع", "ok")
            elif action == "save_branch_permissions":
                uid = int(request.form.get("user_id") or 0)
                bid = int(request.form.get("branch_id") or 0)
                target_perms = ",".join(k for k in request.form.getlist("branch_permissions") if k in PERMISSION_KEYS)
                execute("INSERT INTO user_branch_permissions (user_id, branch_id, permissions, status) VALUES (?,?,?,'نشط') ON CONFLICT(user_id, branch_id) DO UPDATE SET permissions=excluded.permissions, status='نشط'", (uid, bid, target_perms))
                flash("تم حفظ صلاحيات المستخدم لهذا الفرع", "ok")
            elif action == "toggle":
                bid = int(request.form["branch_id"])
                if bid == current_branch_id():
                    flash("لا يمكن تعطيل الفرع الحالي", "err")
                else:
                    execute("UPDATE branches SET status=? WHERE id=?", ("معطل" if request.form.get("status") == "نشط" else "نشط", bid))
                    flash("تم تحديث حالة الفرع", "ok")
            elif action == "default":
                bid = int(request.form["branch_id"])
                execute("UPDATE branches SET is_default=0")
                execute("UPDATE branches SET is_default=1,status='نشط' WHERE id=?", (bid,))
                flash("تم تعيين الفرع الافتراضي", "ok")
            else:
                name = (request.form.get("name") or "").strip()
                if not name: raise ValueError("اسم الفرع مطلوب")
                bid = request.form.get("id")
                if bid:
                    execute("UPDATE branches SET name=?,code=?,status=? WHERE id=?", (name, request.form.get("code") or None, request.form.get("status") or "نشط", bid))
                else:
                    execute("INSERT INTO branches (name,code,status,is_default,created_at) VALUES (?,?,?,0,?)", (name, request.form.get("code") or None, request.form.get("status") or "نشط", datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
                flash("تم حفظ الفرع", "ok")
        except (ValueError, sqlite3.IntegrityError) as exc:
            flash(str(exc), "err")
        return redirect(url_for("branches"))
    branch_rows = query("SELECT ubp.*, b.name branch_name, u.username, u.full_name FROM user_branch_permissions ubp JOIN branches b ON b.id=ubp.branch_id JOIN users u ON u.id=ubp.user_id ORDER BY u.full_name, b.id")
    return render_template("branches.html", branches=query("SELECT * FROM branches ORDER BY id"), users=query("SELECT id,username,full_name,role,branch_id FROM users ORDER BY full_name"), warehouses=query("SELECT id,name,branch FROM warehouses ORDER BY id"), branch_rows=branch_rows, permission_modules=PERMISSION_MODULES)

@app.route("/branches/select", methods=["POST"])
@login_required
def branch_select():
    bid = request.form.get("branch_id")
    branch = query("SELECT id FROM branches WHERE id=? AND status!='معطل'", (bid,), one=True)
    allowed = session.get("role") == "مدير" or query("SELECT 1 FROM user_branch_permissions WHERE user_id=? AND branch_id=? AND status='نشط'", (session.get("user_id"), bid), one=True)
    if branch and allowed:
        session["branch_id"] = branch["id"]
        flash("تم تغيير الفرع الحالي", "ok")
    else:
        flash("لا تملك صلاحية الوصول إلى هذا الفرع", "err")
    return redirect(request.referrer or url_for("dashboard"))

@app.route("/branches/transfer", methods=["POST"])
@login_required
def branch_transfer():
    try:
        product_id, source_id, dest_id = int(request.form["product_id"]), int(request.form["source_id"]), int(request.form["dest_id"])
        qty = float(request.form["qty"])
        if qty <= 0 or source_id == dest_id: raise ValueError("تحقق من المخازن والكمية")
        src = query("SELECT qty FROM stock_balances WHERE warehouse_id=? AND product_id=?", (source_id, product_id), one=True)
        available = float(src["qty"] if src else 0)
        if available < qty: raise ValueError("الكمية المتاحة في المصدر غير كافية")
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        execute("UPDATE stock_balances SET qty=qty-? WHERE warehouse_id=? AND product_id=?", (qty, source_id, product_id))
        execute("INSERT INTO stock_balances (warehouse_id,product_id,qty) VALUES (?,?,?) ON CONFLICT(warehouse_id,product_id) DO UPDATE SET qty=qty+excluded.qty", (dest_id, product_id, qty))
        execute("INSERT INTO stock_moves (product_id,qty,unit_cost,move_type,ref,date,notes,branch_id) VALUES (?,?,?,?,?,?,?,?)", (product_id,-qty,0,"تحويل فرعي",f"{source_id}->{dest_id}",now,"تحويل بين الفروع",current_branch_id()))
        execute("INSERT INTO stock_moves (product_id,qty,unit_cost,move_type,ref,date,notes,branch_id) VALUES (?,?,?,?,?,?,?,?)", (product_id,qty,0,"تحويل فرعي",f"{source_id}->{dest_id}",now,"استلام تحويل بين الفروع",current_branch_id()))
        flash("تم نقل الكمية وتسجيل الحركتين", "ok")
    except (ValueError, KeyError, sqlite3.Error) as exc:
        flash(str(exc), "err")
    return redirect(url_for("branches"))

@app.route("/invoices/<kind>")
@login_required
def invoices_list(kind):
    if kind not in KIND_LABELS:
        return redirect(url_for("dashboard"))
    inv_sql, inv_args = branch_filter("invoices", include_all=True)
    rows = query("SELECT * FROM invoices WHERE kind=?" + inv_sql + " ORDER BY id DESC", [kind] + inv_args)
    return render_template("invoices.html", kind=kind, rows=rows)


@app.route("/invoices/<kind>/new", methods=["GET", "POST"])
@login_required
def invoice_new(kind):
    if kind not in KIND_LABELS:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        try:
            inv_id = save_invoice(kind, request.form)
            flash("تم حفظ المستند بنجاح", "ok")
            return redirect(url_for("invoice_view", inv_id=inv_id))
        except Exception as exc:
            flash(str(exc), "err")
    parties = query(
        "SELECT * FROM customers ORDER BY name"
        if kind in ("sale", "sale_return", "maintenance")
        else "SELECT * FROM suppliers ORDER BY name"
    )
    products = query("SELECT * FROM products ORDER BY name")
    related = None
    if request.args.get("from"):
        related = query("SELECT * FROM invoices WHERE id=?", (request.args["from"],), one=True)
        related_items = query("SELECT * FROM invoice_items WHERE invoice_id=?", (request.args["from"],))
        related = dict(related) if related else None
        if related:
            related["items"] = related_items
    return render_template(
        "invoice_form.html",
        kind=kind,
        parties=parties,
        products=products,
        related=related,
    )


@app.route("/invoice/<int:inv_id>")
@login_required
def invoice_view(inv_id):
    inv = query("SELECT * FROM invoices WHERE id=?", (inv_id,), one=True)
    if not inv:
        flash("المستند غير موجود", "err")
        return redirect(url_for("dashboard"))
    items = query("SELECT * FROM invoice_items WHERE invoice_id=?", (inv_id,))
    return render_template("invoice_view.html", inv=inv, items=items)


@app.route("/api/product/<int:pid>")
@login_required
def api_product(pid):
    row = query("SELECT * FROM products WHERE id=?", (pid,), one=True)
    if not row:
        return jsonify({}), 404
    return jsonify({k: row[k] for k in row.keys()})


@app.route("/api/next-sku")
@login_required
def api_next_sku():
    category = request.args.get("category", "")
    return jsonify({"sku": next_sku(category)})


@app.route("/api/search")
@login_required
def api_search():
    q = (request.args.get("q") or "").strip()
    kind = request.args.get("type") or "product"
    if not q:
        return jsonify([])
    like = f"%{q}%"
    if kind == "product":
        rows = query(
            """SELECT id, sku, name, category, brand, car_model, qty, cost, price, unit, min_qty, notes,
                      location, warehouse, aisle, shelf, bin, barcode, year_from, year_to, item_type
               FROM products
               WHERE name LIKE ? OR sku LIKE ? OR brand LIKE ? OR car_model LIKE ? OR category LIKE ?
                  OR IFNULL(location,'') LIKE ? OR IFNULL(warehouse,'') LIKE ? OR IFNULL(aisle,'') LIKE ?
                  OR IFNULL(shelf,'') LIKE ? OR IFNULL(bin,'') LIKE ? OR IFNULL(barcode,'') LIKE ?
               ORDER BY
                 CASE WHEN sku = ? THEN 0 WHEN name = ? THEN 1
                      WHEN name LIKE ? THEN 2 WHEN sku LIKE ? THEN 3
                      WHEN category LIKE ? THEN 4 WHEN brand LIKE ? THEN 5 ELSE 6 END,
                 name
               LIMIT 18""",
            (like, like, like, like, like, like, like, like, like, like, like, q, q, f"{q}%", f"{q}%", like, like),
        )
        return jsonify(
            [
                {
                    "id": r["id"],
                    "sku": r["sku"],
                    "name": r["name"],
                    "category": r["category"],
                    "brand": r["brand"],
                    "car_model": r["car_model"],
                    "qty": r["qty"],
                    "min_qty": r["min_qty"],
                    "cost": r["cost"],
                    "price": r["price"],
                    "unit": r["unit"],
                    "notes": r["notes"],
                    "item_type": r["item_type"],
                    "location": r["location"],
                    "warehouse": r["warehouse"],
                    "aisle": r["aisle"],
                    "shelf": r["shelf"],
                    "bin": r["bin"],
                    "barcode": r["barcode"],
                    "year_from": r["year_from"],
                    "year_to": r["year_to"],
                    "label": r["name"],
                    "hint": " · ".join(
                        x
                        for x in [
                            r["sku"],
                            r["category"],
                            r["brand"],
                            r["car_model"],
                            f"كمية {r['qty']}",
                            r["location"] or " · ".join(
                                p for p in [
                                    f"مخزن {r['warehouse']}" if r["warehouse"] else "",
                                    f"ممر {r['aisle']}" if r["aisle"] else "",
                                    f"رف {r['shelf']}" if r["shelf"] else "",
                                    f"موضع {r['bin']}" if r["bin"] else "",
                                ] if p
                            ),
                        ]
                        if x
                    ),
                }
                for r in rows
            ]
        )
    if kind in ("category", "brand", "warehouse", "aisle", "shelf", "bin", "unit", "car_model"):
        col = {
            "category": "category",
            "brand": "brand",
            "warehouse": "warehouse",
            "aisle": "aisle",
            "shelf": "shelf",
            "bin": "bin",
            "unit": "unit",
            "car_model": "car_model",
        }[kind]
        rows = query(
            f"""SELECT DISTINCT {col} v, COUNT(*) c FROM products
                WHERE IFNULL({col},'') != '' AND {col} LIKE ?
                GROUP BY {col} ORDER BY c DESC, v LIMIT 12""",
            (like,),
        )
        return jsonify(
            [{"id": r["v"], "name": r["v"], "label": r["v"], "hint": f"{r['c']} أصناف في المخزن"} for r in rows]
        )
    if kind == "customer":
        rows = query(
            """SELECT id, name, phone, address, balance FROM customers
               WHERE name LIKE ? OR IFNULL(phone,'') LIKE ? OR IFNULL(address,'') LIKE ?
               ORDER BY name LIMIT 12""",
            (like, like, like),
        )
        return jsonify(
            [
                {
                    "id": r["id"],
                    "name": r["name"],
                    "phone": r["phone"],
                    "label": r["name"],
                    "hint": " · ".join(x for x in [r["phone"], r["address"]] if x),
                    "balance": r["balance"],
                }
                for r in rows
            ]
        )
    if kind == "supplier":
        rows = query(
            """SELECT id, name, phone, address, balance FROM suppliers
               WHERE name LIKE ? OR IFNULL(phone,'') LIKE ? OR IFNULL(address,'') LIKE ?
               ORDER BY name LIMIT 12""",
            (like, like, like),
        )
        return jsonify(
            [
                {
                    "id": r["id"],
                    "name": r["name"],
                    "phone": r["phone"],
                    "label": r["name"],
                    "hint": " · ".join(x for x in [r["phone"], r["address"]] if x),
                    "balance": r["balance"],
                }
                for r in rows
            ]
        )
    if kind == "employee":
        rows = query(
            """SELECT id, name, job_title, phone, salary FROM employees
               WHERE name LIKE ? OR IFNULL(job_title,'') LIKE ? OR IFNULL(phone,'') LIKE ?
               ORDER BY name LIMIT 12""",
            (like, like, like),
        )
        return jsonify(
            [
                {
                    "id": r["id"],
                    "name": r["name"],
                    "label": r["name"],
                    "hint": " · ".join(
                        x for x in [r["job_title"], r["phone"], f"راتب {r['salary']}"] if x
                    ),
                    "salary": r["salary"],
                }
                for r in rows
            ]
        )
    if kind == "invoice":
        rows = query(
            """SELECT id, number, kind, party_name, date, total FROM invoices
               WHERE number LIKE ? OR IFNULL(party_name,'') LIKE ? OR IFNULL(vehicle,'') LIKE ?
               ORDER BY id DESC LIMIT 12""",
            (like, like, like),
        )
        return jsonify(
            [
                {
                    "id": r["id"],
                    "label": r["number"],
                    "hint": " · ".join(
                        x
                        for x in [
                            KIND_LABELS.get(r["kind"], r["kind"]),
                            r["party_name"],
                            r["date"],
                        ]
                        if x
                    ),
                    "url": url_for("invoice_view", inv_id=r["id"]),
                }
                for r in rows
            ]
        )
    return jsonify([])


@app.route("/settings", methods=["GET", "POST"])
@login_required
@admin_required
def settings():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add_user":
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""
            full_name = (request.form.get("full_name") or "").strip()
            role = request.form.get("role") or "موظف"
            perms = ",".join(
                k for k in request.form.getlist("permissions") if k in PERMISSION_KEYS
            )
            if role == "مدير":
                perms = "all"
            if not username or not password or not full_name:
                flash("أكمل بيانات المستخدم", "err")
            elif len(password) < 4:
                flash("كلمة المرور قصيرة", "err")
            else:
                try:
                    execute(
                        "INSERT INTO users (username, password, full_name, role, permissions, branch_id) VALUES (?,?,?,?,?,?)",
                        (username, password, full_name, role, perms, current_branch_id()),
                    )
                    new_user = query("SELECT id FROM users WHERE username=?", (username,), one=True)
                    if new_user and current_branch_id():
                        execute("INSERT OR IGNORE INTO user_branch_permissions (user_id, branch_id, permissions, status) VALUES (?,?,?,'نشط')", (new_user["id"], current_branch_id(), perms))
                    flash("تم إضافة المستخدم", "ok")
                except sqlite3.IntegrityError:
                    flash("اسم المستخدم موجود مسبقاً", "err")
        elif action == "update_user":
            uid = request.form.get("user_id")
            full_name = (request.form.get("full_name") or "").strip()
            role = request.form.get("role") or "موظف"
            password = request.form.get("password") or ""
            perms = ",".join(
                k for k in request.form.getlist("permissions") if k in PERMISSION_KEYS
            )
            if role == "مدير":
                perms = "all"
            if uid:
                execute(
                    "UPDATE users SET full_name=?, role=?, permissions=? WHERE id=?",
                    (full_name, role, perms, uid),
                )
                if password:
                    execute("UPDATE users SET password=? WHERE id=?", (password, uid))
                if str(uid) == str(session.get("user_id")):
                    session["role"] = role
                    session["full_name"] = full_name
                    session["permissions"] = perms
                flash("تم تحديث المستخدم وصلاحياته", "ok")
        elif action == "save_permissions":
            uid = request.form.get("user_id")
            target = query("SELECT * FROM users WHERE id=?", (uid,), one=True)
            if not target:
                flash("المستخدم غير موجود", "err")
            elif target["role"] == "مدير":
                flash("المدير لديه كل الصلاحيات بالكامل", "err")
            else:
                perms = ",".join(
                    k for k in request.form.getlist("permissions") if k in PERMISSION_KEYS
                )
                execute("UPDATE users SET permissions=? WHERE id=?", (perms, uid))
                if str(uid) == str(session.get("user_id")):
                    session["permissions"] = perms
                flash("تم تحديث صلاحيات المستخدم", "ok")
        elif action == "delete_user":
            uid = int(request.form.get("user_id") or 0)
            target = query("SELECT * FROM users WHERE id=?", (uid,), one=True)
            if not target:
                flash("المستخدم غير موجود", "err")
            elif target["username"] == session.get("user"):
                flash("لا يمكن حذف المستخدم الحالي", "err")
            elif target["username"] == "admin" and query(
                "SELECT COUNT(*) c FROM users WHERE role='مدير'", one=True
            )["c"] <= 1:
                flash("لا يمكن حذف آخر مدير", "err")
            else:
                execute("DELETE FROM users WHERE id=?", (uid,))
                flash("تم حذف المستخدم", "ok")
        elif action == "save_print":
            keys = (
                "shop_name",
                "shop_subtitle",
                "phone",
                "address",
                "tax_no",
                "print_footer",
                "paper_size",
                "print_copies",
                "sku_prefix",
                "vat_rate",
                "points_per_100",
                "default_labor_rate",
                "work_start_time",
                "attendance_grace_minutes",
                "sync_server_url",
                "sync_token",
                "device_name",
            )
            for key in keys:
                execute(
                    "INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                    (key, request.form.get(key) or ""),
                )
            logo_file = request.files.get("company_logo")
            if logo_file and logo_file.filename:
                raw_logo = logo_file.read()
                if len(raw_logo) > 2 * 1024 * 1024:
                    flash("حجم الشعار يجب ألا يتجاوز 2 ميجابايت", "err")
                elif logo_file.mimetype not in ("image/png", "image/jpeg", "image/jpg", "image/gif"):
                    flash("صيغة الشعار يجب أن تكون PNG أو JPG أو GIF", "err")
                else:
                    logo_data = "data:%s;base64,%s" % (logo_file.mimetype, base64.b64encode(raw_logo).decode("ascii"))
                    execute("INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", ("company_logo_data", logo_data))
            execute(
                "INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                ("show_cost", "1" if request.form.get("show_cost") else "0"),
            )
            execute(
                "INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                ("show_profit", "1" if request.form.get("show_profit") else "0"),
            )
            execute(
                "INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                ("sync_enabled", "1" if request.form.get("sync_enabled") else "0"),
            )
            flash("تم حفظ إعدادات الطباعة والمزامنة والشعار", "ok")
        return redirect(url_for("settings"))
    users = query("SELECT * FROM users ORDER BY id")
    perm_groups = []
    current_group = None
    for key, label, group in PERMISSION_MODULES:
        if group != current_group:
            perm_groups.append({"name": group, "mods": []})
            current_group = group
        perm_groups[-1]["mods"].append({"key": key, "label": label})
    return render_template(
        "settings.html",
        users=users,
        settings=all_settings(),
        perm_groups=perm_groups,
    )


@app.route("/invoice/<int:inv_id>/print")
@login_required
def invoice_print(inv_id):
    inv = query("SELECT * FROM invoices WHERE id=?", (inv_id,), one=True)
    if not inv:
        flash("المستند غير موجود", "err")
        return redirect(url_for("dashboard"))
    items = query("SELECT * FROM invoice_items WHERE invoice_id=?", (inv_id,))
    book_type = "نقدي"
    method = (inv["payment_method"] or inv["status"] or "نقدي")
    salesperson = (inv["salesperson"] or "")
    if method in ("آجل", "شيك", "أقساط"):
        book_type = "آجل"
    elif "مندوب" in salesperson:
        book_type = "مندوب"
    return render_template(
        "invoice_print.html",
        inv=inv,
        items=items,
        settings=all_settings(),
        book_type=book_type,
    )


def pdf_table_style():
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f6fb5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8c6d4")),
        ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f7fb")]),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
    ])


def pdf_company_header(styles, title):
    settings = all_settings()
    header = []
    logo_data = settings.get("company_logo_data", "")
    if logo_data.startswith("data:image/") and "," in logo_data:
        try:
            logo = PdfImage(io.BytesIO(base64.b64decode(logo_data.split(",", 1)[1])), width=58, height=58)
            header.append([logo, Paragraph(rtl_pdf(settings.get("shop_name", APP_NAME)), styles["ArabicTitle"])])
            table = Table(header, colWidths=[70, 560], style=TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "RIGHT")]))
            result = [table]
        except Exception:
            result = [Paragraph(rtl_pdf(settings.get("shop_name", APP_NAME)), styles["ArabicTitle"])]
    else:
        result = [Paragraph(rtl_pdf(settings.get("shop_name", APP_NAME)), styles["ArabicTitle"])]
    company_line = " | ".join(x for x in [settings.get("shop_subtitle"), settings.get("phone"), settings.get("address"), ("الرقم الضريبي: " + settings.get("tax_no", "")) if settings.get("tax_no") else ""] if x)
    if company_line:
        result.append(Paragraph(rtl_pdf(company_line), styles["Arabic"]))
    result.append(Paragraph(rtl_pdf(title), styles["Arabic"]))
    return result


def invoice_pdf_response(inv, items, filename_prefix="فاتورة"):
    if not REPORTLAB_OK:
        flash("تصدير PDF غير متاح حاليًا على الخادم", "err")
        return redirect(url_for("invoice_view", inv_id=inv["id"]))
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Arabic", fontName="DejaVu", fontSize=9, leading=13, alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="ArabicTitle", fontName="DejaVu", fontSize=15, leading=20, alignment=TA_RIGHT, textColor=colors.HexColor("#173653")))
    story = pdf_company_header(styles, f"{KIND_LABELS.get(inv['kind'], inv['kind'])} — رقم {inv['number']}") + [
        Paragraph(rtl_pdf(f"التاريخ: {inv['date']}   |   الطرف: {inv['party_name'] or '-'}   |   الحالة: {inv['status']}"), styles["Arabic"]),
        Spacer(1, 10),
    ]
    data = [[rtl_pdf(x) for x in ["الوصف", "الكمية", "سعر الوحدة", "إجمالي الصنف", "التكلفة", "الربح"]]]
    for item in items:
        data.append([
            Paragraph(rtl_pdf(item["description"]), styles["Arabic"]),
            str(item["qty"]), money(item["unit_price"]), money(item["line_total"]),
            money(item["line_cost"]), money(item["line_profit"]),
        ])
    story.append(Table(data, repeatRows=1, colWidths=[260, 65, 85, 95, 85, 85], style=pdf_table_style()))
    story.extend([
        Spacer(1, 10),
        Paragraph(rtl_pdf(f"المجموع: {money(inv['subtotal'])}   |   الخصم: {money(inv['discount'])}   |   الضريبة: {money(inv['tax'])}   |   الصافي: {money(inv['total'])}"), styles["Arabic"]),
        Paragraph(rtl_pdf(f"المدفوع: {money(inv['paid'])}   |   المتبقي: {money(float(inv['total']) - float(inv['paid']))}"), styles["Arabic"]),
    ])
    doc.build(story)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"{filename_prefix}-{inv['number']}.pdf", mimetype="application/pdf")


@app.route("/invoice/<int:inv_id>/pdf")
@login_required
def invoice_pdf(inv_id):
    inv = query("SELECT * FROM invoices WHERE id=?", (inv_id,), one=True)
    if not inv:
        flash("المستند غير موجود", "err")
        return redirect(url_for("dashboard"))
    items = query("SELECT * FROM invoice_items WHERE invoice_id=?", (inv_id,))
    return invoice_pdf_response(inv, items)


@app.route("/quotes/<int:qid>/pdf")
@login_required
def quote_pdf(qid):
    quote = query("SELECT * FROM quotes WHERE id=?", (qid,), one=True)
    items = query("SELECT * FROM quote_items WHERE quote_id=?", (qid,))
    if not quote:
        flash("عرض السعر غير موجود", "err")
        return redirect(url_for("quotes"))
    if not REPORTLAB_OK:
        flash("تصدير PDF غير متاح حاليًا على الخادم", "err")
        return redirect(url_for("quotes"))
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="QuoteArabic", fontName="DejaVu", fontSize=9, leading=13, alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="QuoteTitle", fontName="DejaVu", fontSize=15, leading=20, alignment=TA_RIGHT, textColor=colors.HexColor("#173653")))
    story = pdf_company_header({"ArabicTitle": styles["QuoteTitle"], "Arabic": styles["QuoteArabic"]}, f"عرض سعر — رقم {quote['number']}") + [
        Paragraph(rtl_pdf(f"التاريخ: {quote['date']}   |   العميل: {quote['customer_name'] or '-'}   |   الحالة: {'منفذ' if quote['invoice_id'] else 'غير منفذ'}"), styles["QuoteArabic"]),
        Spacer(1, 10),
    ]
    data = [[rtl_pdf(x) for x in ["الوصف", "الكمية", "سعر الوحدة", "الإجمالي"]]]
    for item in items:
        data.append([Paragraph(rtl_pdf(item["description"]), styles["QuoteArabic"]), str(item["qty"]), money(item["unit_price"]), money(item["line_total"])])
    story.append(Table(data, repeatRows=1, colWidths=[330, 80, 100, 110], style=pdf_table_style()))
    story.extend([
        Spacer(1, 10),
        Paragraph(rtl_pdf(f"المجموع: {money(quote['subtotal'])}   |   الخصم: {money(quote['discount'])}   |   الضريبة: {money(quote['tax'])}   |   الإجمالي: {money(quote['total'])}"), styles["QuoteArabic"]),
        Paragraph(rtl_pdf(f"ملاحظات: {quote['notes'] or '-'}"), styles["QuoteArabic"]),
    ])
    doc.build(story)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"عرض-سعر-{quote['number']}.pdf", mimetype="application/pdf")


@app.route("/invoice/<int:inv_id>/cancel", methods=["POST"])
@login_required
def invoice_cancel(inv_id):
    try:
        cancel_invoice(inv_id)
        flash("تم إلغاء المستند وإرجاع المخزون والخزينة", "ok")
    except Exception as exc:
        flash(str(exc), "err")
    return redirect(url_for("invoice_view", inv_id=inv_id))


@app.route("/employees", methods=["GET", "POST"])
@login_required
def employees():
    if request.method == "POST":
        eid = request.form.get("id")
        data = (
            request.form["name"].strip(),
            request.form.get("job_title"),
            request.form.get("phone"),
            request.form.get("hire_date") or date.today().isoformat(),
            float(request.form.get("salary") or 0),
            request.form.get("status") or "نشط",
            request.form.get("notes"),
            (request.form.get("username") or "").strip() or None,
        )
        if eid:
            execute(
                """UPDATE employees SET name=?, job_title=?, phone=?, hire_date=?, salary=?, status=?, notes=?, username=?
                   WHERE id=?""",
                data + (eid,),
            )
            flash("تم تحديث الموظف", "ok")
        else:
            execute(
                """INSERT INTO employees (name, job_title, phone, hire_date, salary, status, notes, username)
                   VALUES (?,?,?,?,?,?,?,?)""",
                data,
            )
            flash("تم إضافة الموظف", "ok")
        return redirect(url_for("employees"))
    rows = query("SELECT * FROM employees ORDER BY name")
    edit = None
    if request.args.get("edit"):
        edit = query("SELECT * FROM employees WHERE id=?", (request.args["edit"],), one=True)
    return render_template("employees.html", rows=rows, edit=edit)


@app.route("/employees/delete/<int:eid>")
@login_required
def delete_employee(eid):
    execute("DELETE FROM employees WHERE id=?", (eid,))
    flash("تم حذف الموظف", "ok")
    return redirect(url_for("employees"))


@app.route("/advances", methods=["GET", "POST"])
@login_required
def advances():
    if request.method == "POST":
        emp_id = int(request.form["employee_id"])
        amount = float(request.form["amount"])
        execute(
            "INSERT INTO advances (employee_id, amount, remaining, date, notes) VALUES (?,?,?,?,?)",
            (
                emp_id,
                amount,
                amount,
                request.form.get("date") or date.today().isoformat(),
                request.form.get("notes"),
            ),
        )
        flash("تم تسجيل السلفة", "ok")
        return redirect(url_for("advances"))
    rows = query(
        """SELECT a.*, e.name FROM advances a JOIN employees e ON e.id=a.employee_id
           ORDER BY a.id DESC"""
    )
    emps = query("SELECT * FROM employees WHERE status='نشط' ORDER BY name")
    total_open = query("SELECT COALESCE(SUM(remaining),0) v FROM advances", one=True)["v"]
    return render_template("advances.html", rows=rows, emps=emps, total_open=total_open)


@app.route("/attendance", methods=["GET", "POST"])
@login_required
def attendance():
    if request.method == "POST":
        employee_id = int(request.form["employee_id"])
        emp_branch_sql, emp_branch_args = branch_filter("employees")
        employee = query("SELECT id FROM employees WHERE id=?" + emp_branch_sql, [employee_id] + emp_branch_args, one=True)
        if not employee:
            flash("لا يمكن تسجيل حضور موظف خارج الفرع الحالي", "err")
            return redirect(url_for("attendance"))
        day = request.form.get("attendance_date") or date.today().isoformat()
        existing = query("SELECT id FROM attendance WHERE employee_id=? AND attendance_date=?", (employee_id, day), one=True)
        if existing:
            execute("UPDATE attendance SET status='غائب', notes=? WHERE id=?", (request.form.get("notes"), existing["id"]))
        else:
            execute("INSERT INTO attendance (employee_id, attendance_date, status, notes) VALUES (?,?,?,?)", (employee_id, day, "غائب", request.form.get("notes")))
        audit_log("تسجيل غياب", "حضور", employee_id, day, request.form.get("notes") or "")
        flash("تم تسجيل الغياب، وسيظهر الخصم للمراجعة في كشف الراتب", "ok")
        return redirect(url_for("attendance"))
    period = request.args.get("period") or date.today().strftime("%Y-%m")
    emp_branch_sql, emp_branch_args = branch_filter("e")
    rows = query("""SELECT a.*, e.name, e.salary FROM attendance a JOIN employees e ON e.id=a.employee_id
                    WHERE a.attendance_date LIKE ?""" + emp_branch_sql + " ORDER BY a.attendance_date DESC, e.name", [period + "%"] + emp_branch_args)
    emps = query("SELECT * FROM employees WHERE status='نشط'" + emp_branch_sql.replace("e.", "") + " ORDER BY name", emp_branch_args)
    return render_template("attendance.html", rows=rows, emps=emps, period=period)


@app.route("/salaries", methods=["GET", "POST"])
@login_required
def salaries():
    if request.method == "POST":
        emp = query("SELECT * FROM employees WHERE id=?", (request.form["employee_id"],), one=True)
        period = request.form["period"]
        bonus = float(request.form.get("bonus") or 0)
        preview = attendance_preview(emp["id"], period)
        attendance_deduct = preview["deduct"] if request.form.get("attendance_approve") == "1" else 0.0
        open_adv = query(
            "SELECT COALESCE(SUM(remaining),0) v FROM advances WHERE employee_id=?",
            (emp["id"],),
            one=True,
        )["v"]
        deduct = min(float(request.form.get("advance_deduct") or open_adv), open_adv)
        net = float(emp["salary"]) + bonus - deduct - attendance_deduct
        execute(
            """INSERT INTO salaries (employee_id, period, base_salary, bonus, advance_deduct, net, status, paid_at, notes, attendance_deduct, absence_days, late_minutes)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                emp["id"],
                period,
                emp["salary"],
                bonus,
                deduct,
                net,
                "مدفوع",
                date.today().isoformat(),
                request.form.get("notes"),
                attendance_deduct,
                preview["absence_days"] if attendance_deduct else 0,
                preview["late_minutes"] if attendance_deduct else 0,
            ),
        )
        remain = deduct
        if remain:
            opens = query(
                "SELECT * FROM advances WHERE employee_id=? AND remaining>0 ORDER BY id",
                (emp["id"],),
            )
            for adv in opens:
                if remain <= 0:
                    break
                take = min(remain, float(adv["remaining"]))
                execute("UPDATE advances SET remaining=remaining-? WHERE id=?", (take, adv["id"]))
                remain -= take
        flash("تم صرف الراتب", "ok")
        return redirect(url_for("salaries"))
    rows = query(
        """SELECT s.*, e.name, e.job_title FROM salaries s
           JOIN employees e ON e.id=s.employee_id ORDER BY s.id DESC"""
    )
    emps = query("SELECT * FROM employees WHERE status='نشط' ORDER BY name")
    previews = {str(e["id"]): attendance_preview(e["id"], datetime.now().strftime("%Y-%m")) for e in emps}
    total = query("SELECT COALESCE(SUM(net),0) v FROM salaries", one=True)["v"]
    return render_template("salaries.html", rows=rows, emps=emps, previews=previews, total=total)


@app.route("/costs")
@login_required
def costs():
    products = query(
        """SELECT p.*, (p.price-p.cost) margin,
           CASE WHEN p.price>0 THEN ROUND((p.price-p.cost)*100.0/p.price,1) ELSE 0 END margin_pct
           FROM products p ORDER BY margin DESC"""
    )
    sales_cost = query(
        """SELECT date, number, kind, party_name, total, cost_total, profit
           FROM invoices WHERE kind IN ('sale','maintenance','sale_return')
           ORDER BY id DESC LIMIT 30"""
    )
    avg = query(
        "SELECT COALESCE(AVG(cost),0) c, COALESCE(AVG(price),0) p FROM products",
        one=True,
    )
    return render_template("costs.html", products=products, sales_cost=sales_cost, avg=avg)


@app.route("/pnl")
@login_required
def pnl():
    ops = query(
        """SELECT * FROM invoices WHERE status!='ملغاة' ORDER BY date DESC, id DESC"""
    )
    sales = sum(r["total"] for r in ops if r["kind"] in ("sale", "maintenance"))
    sale_returns = sum(r["total"] for r in ops if r["kind"] == "sale_return")
    purchases = sum(r["total"] for r in ops if r["kind"] == "purchase")
    purchase_returns = sum(r["total"] for r in ops if r["kind"] == "purchase_return")
    cogs = sum(r["cost_total"] for r in ops if r["kind"] in ("sale", "maintenance"))
    cogs_return = sum(r["cost_total"] for r in ops if r["kind"] == "sale_return")
    salaries_total = query("SELECT COALESCE(SUM(net),0) v FROM salaries", one=True)["v"]
    expenses_total = query("SELECT COALESCE(SUM(amount),0) v FROM expenses", one=True)["v"]
    gross = (sales - sale_returns) - (cogs - cogs_return)
    net_profit = gross - salaries_total - expenses_total
    by_op = query(
        """SELECT id, number, kind, date, party_name, total, cost_total, profit, notes
           FROM invoices WHERE status!='ملغاة' ORDER BY id DESC"""
    )
    return render_template(
        "pnl.html",
        ops=ops,
        sales=sales,
        sale_returns=sale_returns,
        purchases=purchases,
        purchase_returns=purchase_returns,
        cogs=cogs - cogs_return,
        salaries_total=salaries_total,
        expenses_total=expenses_total,
        gross=gross,
        net_profit=net_profit,
        by_op=by_op,
    )


@app.route("/accounts")
@login_required
def accounts():
    customers = query("SELECT * FROM customers ORDER BY balance DESC")
    suppliers = query("SELECT * FROM suppliers ORDER BY balance DESC")
    ledger = query("SELECT * FROM party_ledger ORDER BY id DESC LIMIT 40")
    journals = query("SELECT * FROM journal_entries ORDER BY id DESC LIMIT 30")
    return render_template(
        "accounts.html",
        customers=customers,
        suppliers=suppliers,
        ledger=ledger,
        journals=journals,
    )


@app.route("/pos", methods=["GET", "POST"])
@login_required
def pos():
    if request.method == "POST":
        try:
            inv_id = save_invoice("sale", request.form)
            flash("تم إصدار فاتورة الكاشير", "ok")
            return redirect(url_for("invoice_print", inv_id=inv_id))
        except Exception as exc:
            flash(str(exc), "err")
    customers = query("SELECT * FROM customers ORDER BY name")
    vat = float(get_setting("vat_rate", "15") or 0)
    cash_accounts = query("SELECT * FROM cash_accounts ORDER BY id")
    return render_template(
        "pos.html",
        customers=customers,
        vat=vat,
        pay_methods=PAY_METHODS,
        cash_accounts=cash_accounts,
        shift=current_shift(),
    )


@app.route("/jobs")
@login_required
def jobs():
    rows = query("SELECT * FROM job_cards ORDER BY id DESC")
    return render_template("jobs.html", rows=rows, statuses=JOB_STATUSES)


@app.route("/jobs/new", methods=["GET", "POST"])
@login_required
def job_new():
    if request.method == "POST":
        number = next_number("JOB", "job_cards")
        tech_id = request.form.get("technician_id") or None
        tech = query("SELECT name FROM employees WHERE id=?", (tech_id,), one=True) if tech_id else None
        cust_id = request.form.get("customer_id") or None
        name = request.form.get("customer_name") or ""
        if cust_id and not name:
            c = query("SELECT name FROM customers WHERE id=?", (cust_id,), one=True)
            name = c["name"] if c else ""
        jid = execute(
            """INSERT INTO job_cards
               (number, customer_id, customer_name, phone, plate, vehicle_type, vehicle_model, vehicle_year,
                complaint, technician_id, technician_name, status, opened_at, notes, created_by)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                number,
                cust_id,
                name,
                request.form.get("phone"),
                request.form.get("plate"),
                request.form.get("vehicle_type"),
                request.form.get("vehicle_model"),
                request.form.get("vehicle_year"),
                request.form.get("complaint"),
                tech_id,
                tech["name"] if tech else request.form.get("technician_name"),
                "قيد الانتظار",
                datetime.now().strftime("%Y-%m-%d %H:%M"),
                request.form.get("notes"),
                session.get("user"),
            ),
        )
        flash("تم فتح أمر الشغل", "ok")
        return redirect(url_for("job_view", job_id=jid))
    customers = query("SELECT * FROM customers ORDER BY name")
    techs = query("SELECT * FROM employees WHERE status='نشط' ORDER BY name")
    return render_template("job_form.html", customers=customers, techs=techs)


@app.route("/jobs/<int:job_id>", methods=["GET", "POST"])
@login_required
def job_view(job_id):
    job = query("SELECT * FROM job_cards WHERE id=?", (job_id,), one=True)
    if not job:
        flash("أمر الشغل غير موجود", "err")
        return redirect(url_for("jobs"))
    if request.method == "POST":
        action = request.form.get("action")
        try:
            if action == "add_part":
                pid = request.form.get("product_id")
                qty = float(request.form.get("qty") or 1)
                prod = query("SELECT * FROM products WHERE id=?", (pid,), one=True)
                if not prod:
                    raise ValueError("اختر قطعة غيار")
                execute(
                    """INSERT INTO job_parts (job_id, product_id, description, qty, unit_cost, unit_price, line_total, line_cost)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (
                        job_id,
                        prod["id"],
                        prod["name"],
                        qty,
                        prod["cost"],
                        prod["price"],
                        qty * float(prod["price"]),
                        qty * float(prod["cost"]),
                    ),
                )
                apply_stock(prod["id"], -qty, prod["cost"], "job", job["number"], prod["name"])
                flash("تم سحب القطعة من المخزن إلى أمر الشغل", "ok")
            elif action == "add_labor":
                emp_id = request.form.get("employee_id") or None
                emp = query("SELECT * FROM employees WHERE id=?", (emp_id,), one=True) if emp_id else None
                hours = float(request.form.get("hours") or 1)
                rate = float(request.form.get("rate") or get_setting("default_labor_rate", "80") or 80)
                amount = round(hours * rate, 2)
                execute(
                    """INSERT INTO job_labor (job_id, employee_id, service_name, hours, rate, amount)
                       VALUES (?,?,?,?,?,?)""",
                    (
                        job_id,
                        emp_id,
                        request.form.get("service_name") or "مصنعية",
                        hours,
                        rate,
                        amount,
                    ),
                )
                flash("تم إضافة المصنعية", "ok")
            elif action == "status":
                execute("UPDATE job_cards SET status=? WHERE id=?", (request.form.get("status"), job_id))
                flash("تم تحديث حالة الإصلاح", "ok")
            elif action == "totals":
                execute(
                    "UPDATE job_cards SET discount=?, tax=? WHERE id=?",
                    (float(request.form.get("discount") or 0), float(request.form.get("tax") or 0), job_id),
                )
                flash("تم تحديث الخصم والضريبة", "ok")
            recalc_job(job_id)
        except Exception as exc:
            flash(str(exc), "err")
        return redirect(url_for("job_view", job_id=job_id))
    parts = query("SELECT * FROM job_parts WHERE job_id=?", (job_id,))
    labor = query("SELECT * FROM job_labor WHERE job_id=?", (job_id,))
    techs = query("SELECT * FROM employees WHERE status='نشط' ORDER BY name")
    vat = float(get_setting("vat_rate", "15") or 0)
    return render_template(
        "job_view.html",
        job=job,
        parts=parts,
        labor=labor,
        techs=techs,
        statuses=JOB_STATUSES,
        vat=vat,
        pay_methods=PAY_METHODS,
    )


@app.route("/jobs/<int:job_id>/invoice", methods=["POST"])
@login_required
def job_to_invoice(job_id):
    job = query("SELECT * FROM job_cards WHERE id=?", (job_id,), one=True)
    if not job:
        flash("أمر الشغل غير موجود", "err")
        return redirect(url_for("jobs"))
    if job["invoice_id"]:
        return redirect(url_for("invoice_view", inv_id=job["invoice_id"]))
    if job["status"] not in ("تم الإصلاح", "تم التسليم"):
        flash("حوّل الحالة إلى تم الإصلاح قبل إصدار الفاتورة", "err")
        return redirect(url_for("job_view", job_id=job_id))
    recalc_job(job_id)
    job = query("SELECT * FROM job_cards WHERE id=?", (job_id,), one=True)
    parts = query("SELECT * FROM job_parts WHERE job_id=?", (job_id,))
    labor = query("SELECT * FROM job_labor WHERE job_id=?", (job_id,))
    payment_method = request.form.get("payment_method") or "نقدي"
    items = []
    for p in parts:
        items.append(
            {
                "product_id": p["product_id"],
                "description": p["description"],
                "qty": p["qty"],
                "unit_cost": p["unit_cost"],
                "unit_price": p["unit_price"],
                "line_total": p["line_total"],
                "line_cost": p["line_cost"],
                "line_profit": float(p["line_total"]) - float(p["line_cost"]),
            }
        )
    labor_total = sum(float(x["amount"]) for x in labor)
    if labor_total:
        items.append(
            {
                "product_id": None,
                "description": "مصنعية / خدمات الورشة",
                "qty": 1,
                "unit_cost": 0,
                "unit_price": labor_total,
                "line_total": labor_total,
                "line_cost": 0,
                "line_profit": labor_total,
            }
        )
    if not items:
        flash("أضف قطعاً أو مصنعية قبل الفوترة", "err")
        return redirect(url_for("job_view", job_id=job_id))
    class Form(dict):
        def getlist(self, key):
            return self.get(key, [])
    form = Form()
    form["party_id"] = job["customer_id"] or ""
    form["party_name"] = job["customer_name"]
    form["discount"] = job["discount"]
    form["tax"] = job["tax"]
    form["payment_method"] = payment_method
    form["status"] = payment_method
    form["vehicle"] = " ".join(x for x in [job["vehicle_type"], job["vehicle_model"], job["plate"]] if x)
    form["notes"] = f"فاتورة أمر شغل {job['number']}"
    form["labor_total"] = labor_total
    form["item_product_id[]"] = [str(i["product_id"] or "") for i in items]
    form["item_desc[]"] = [i["description"] for i in items]
    form["item_qty[]"] = [str(i["qty"]) for i in items]
    form["item_price[]"] = [str(i["unit_price"]) for i in items]
    form["item_cost[]"] = [str(i["unit_cost"]) for i in items]
    try:
        for p in parts:
            prod = query("SELECT qty FROM products WHERE id=?", (p["product_id"],), one=True)
            if prod:
                execute("UPDATE products SET qty = qty + ? WHERE id=?", (p["qty"], p["product_id"]))
        inv_id = save_invoice("maintenance", form, related_id=job_id)
        execute(
            "UPDATE job_cards SET invoice_id=?, status=?, closed_at=? WHERE id=?",
            (inv_id, "تم التسليم", datetime.now().strftime("%Y-%m-%d %H:%M"), job_id),
        )
        flash("تم تحويل أمر الشغل إلى فاتورة نهائية", "ok")
        return redirect(url_for("invoice_view", inv_id=inv_id))
    except Exception as exc:
        flash(str(exc), "err")
        return redirect(url_for("job_view", job_id=job_id))


@app.route("/expenses", methods=["GET", "POST"])
@login_required
def expenses():
    if request.method == "POST":
        execute(
            "INSERT INTO expenses (category, amount, date, paid_to, notes, created_by) VALUES (?,?,?,?,?,?)",
            (
                request.form.get("category") or "أخرى",
                float(request.form.get("amount") or 0),
                request.form.get("date") or date.today().isoformat(),
                request.form.get("paid_to"),
                request.form.get("notes"),
                session.get("user"),
            ),
        )
        last = query("SELECT id FROM expenses ORDER BY id DESC LIMIT 1", one=True)
        post_cash("نقدي", float(request.form.get("amount") or 0), request.form.get("category") or "مصروف", "expense", last["id"] if last else None, "out")
        add_journal(request.form.get("category") or "مصروف", "المصروفات", debit=float(request.form.get("amount") or 0), reference_type="expense", reference_id=last["id"] if last else None)
        flash("تم تسجيل المصروف", "ok")
        return redirect(url_for("expenses"))
    rows = query("SELECT * FROM expenses ORDER BY id DESC")
    total = query("SELECT COALESCE(SUM(amount),0) v FROM expenses", one=True)["v"]
    by_cat = query("SELECT category, SUM(amount) v FROM expenses GROUP BY category ORDER BY v DESC")
    return render_template("expenses.html", rows=rows, total=total, by_cat=by_cat)


@app.route("/reports")
@login_required
def reports():
    period = request.args.get("period") or "month"
    report_branch = request.args.get("branch_id")
    today = date.today()
    if period == "day":
        label, like = "اليوم", today.isoformat() + "%"
    elif period == "year":
        label, like = f"سنة {today.year}", f"{today.year}%"
    else:
        label, like = today.strftime("%Y-%m"), today.strftime("%Y-%m") + "%"
    branch_sql, branch_args = branch_filter("invoices", include_all=True)
    if session.get("role") == "مدير" and report_branch and report_branch.isdigit():
        branch_sql, branch_args = " AND invoices.branch_id=?", [int(report_branch)]
    sales = query("SELECT COALESCE(SUM(total),0) s, COALESCE(SUM(profit),0) p FROM invoices WHERE kind IN ('sale','maintenance') AND status!='ملغاة' AND date LIKE ?" + branch_sql, [like] + branch_args, one=True)
    purchases = query("SELECT COALESCE(SUM(total),0) s FROM invoices WHERE kind='purchase' AND status!='ملغاة' AND date LIKE ?" + branch_sql, [like] + branch_args, one=True)
    expenses_v = query("SELECT COALESCE(SUM(amount),0) v FROM expenses WHERE date LIKE ?", (like,), one=True)["v"]
    salaries_v = query("SELECT COALESCE(SUM(net),0) v FROM salaries WHERE paid_at LIKE ?", (like,), one=True)["v"]
    top_sql = " AND v.branch_id=?" if branch_sql else ""
    top_args = [like] + (branch_args if branch_sql else [])
    top_qty = query("SELECT p.name, p.sku, SUM(i.qty) qty, SUM(i.line_total) sales, SUM(i.line_profit) profit FROM invoice_items i JOIN invoices v ON v.id=i.invoice_id JOIN products p ON p.id=i.product_id WHERE v.kind IN ('sale','maintenance') AND v.date LIKE ?" + top_sql + " GROUP BY p.id ORDER BY qty DESC LIMIT 10", top_args)
    top_profit = query("SELECT p.name, p.sku, SUM(i.qty) qty, SUM(i.line_total) sales, SUM(i.line_profit) profit FROM invoice_items i JOIN invoices v ON v.id=i.invoice_id JOIN products p ON p.id=i.product_id WHERE v.kind IN ('sale','maintenance') AND v.date LIKE ?" + top_sql + " GROUP BY p.id ORDER BY profit DESC LIMIT 10", top_args)
    customers = query("SELECT * FROM customers WHERE balance>0 ORDER BY balance DESC")
    suppliers = query("SELECT * FROM suppliers WHERE balance>0 ORDER BY balance DESC")
    return render_template("reports.html", period=period, label=label, sales=sales, purchases=purchases, expenses_v=expenses_v, salaries_v=salaries_v, net=float(sales["p"])-float(expenses_v)-float(salaries_v), top_qty=top_qty, top_profit=top_profit, customers=customers, suppliers=suppliers, branches=query("SELECT * FROM branches WHERE status!='معطل' ORDER BY id"), selected_branch_id=int(report_branch) if report_branch and report_branch.isdigit() else current_branch_id())


@app.route("/reports/export.pdf")
@login_required
def reports_export_pdf():
    if not REPORTLAB_OK:
        flash("تصدير PDF غير متاح حاليًا على الخادم", "err")
        return redirect(url_for("reports"))
    period = request.args.get("period") or "month"
    today = date.today()
    if period == "day":
        label, like = "اليوم", today.isoformat() + "%"
    elif period == "year":
        label, like = f"سنة {today.year}", f"{today.year}%"
    else:
        label, like = today.strftime("%Y-%m"), today.strftime("%Y-%m") + "%"
    branch_sql, branch_args = branch_filter("invoices", include_all=True)
    sales = query("SELECT COALESCE(SUM(total),0) s, COALESCE(SUM(profit),0) p FROM invoices WHERE kind IN ('sale','maintenance') AND status!='ملغاة' AND date LIKE ?" + branch_sql, [like] + branch_args, one=True)
    purchases = query("SELECT COALESCE(SUM(total),0) s FROM invoices WHERE kind='purchase' AND status!='ملغاة' AND date LIKE ?" + branch_sql, [like] + branch_args, one=True)
    expenses_v = query("SELECT COALESCE(SUM(amount),0) v FROM expenses WHERE date LIKE ?", (like,), one=True)["v"]
    salaries_v = query("SELECT COALESCE(SUM(net),0) v FROM salaries WHERE paid_at LIKE ?", (like,), one=True)["v"]
    top_sql = " AND v.branch_id=?" if branch_sql else ""
    top_args = [like] + (branch_args if branch_sql else [])
    top_rows = query("SELECT p.name, p.sku, SUM(i.qty) qty, SUM(i.line_total) sales, SUM(i.line_profit) profit FROM invoice_items i JOIN invoices v ON v.id=i.invoice_id JOIN products p ON p.id=i.product_id WHERE v.kind IN ('sale','maintenance') AND v.status!='ملغاة' AND v.date LIKE ?" + top_sql + " GROUP BY p.id ORDER BY sales DESC LIMIT 15", top_args)
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportArabic", fontName="DejaVu", fontSize=9, leading=13, alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="ReportTitle", fontName="DejaVu", fontSize=15, leading=20, alignment=TA_RIGHT, textColor=colors.HexColor("#173653")))
    story = pdf_company_header({"ArabicTitle": styles["ReportTitle"], "Arabic": styles["ReportArabic"]}, f"التقرير المالي — {label}")
    summary = [[rtl_pdf(x) for x in ["المبيعات", "المشتريات", "المصروفات", "الرواتب", "صافي الدخل"]], [money(sales["s"]), money(purchases["s"]), money(expenses_v), money(salaries_v), money(float(sales["p"]) - float(expenses_v) - float(salaries_v))]]
    story += [Table(summary, colWidths=[130] * 5, style=pdf_table_style()), Spacer(1, 14), Paragraph(rtl_pdf("الأصناف الأكثر مبيعًا"), styles["ReportTitle"])]
    data = [[rtl_pdf(x) for x in ["الصنف", "الكود", "الكمية", "المبيعات", "الربح"]]]
    data += [[Paragraph(rtl_pdf(r["name"]), styles["ReportArabic"]), str(r["sku"]), str(r["qty"]), money(r["sales"]), money(r["profit"])] for r in top_rows]
    story.append(Table(data, repeatRows=1, colWidths=[290, 110, 90, 110, 110], style=pdf_table_style()))
    doc.build(story)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"التقرير-المالي-{period}.pdf", mimetype="application/pdf")
@app.route("/reports/shifts")
@login_required
def shift_reports():
    report_day = request.args.get("date") or date.today().isoformat()
    report_branch = request.args.get("branch_id")
    shift_sql, shift_args = branch_filter("s", include_all=True)
    invoice_sql, invoice_args = branch_filter("v", include_all=True)
    if session.get("role") == "مدير" and report_branch and report_branch.isdigit():
        shift_sql, shift_args = " AND s.branch_id=?", [int(report_branch)]
        invoice_sql, invoice_args = " AND v.branch_id=?", [int(report_branch)]
    shift_rows = query(
        """SELECT s.*, b.name branch_name, COALESCE(e.name, s.username) employee_name, e.job_title,
                  COUNT(v.id) invoice_count, COALESCE(SUM(v.total),0) sales_total,
                  COALESCE(SUM(v.paid),0) sales_paid, COALESCE(SUM(v.profit),0) sales_profit
           FROM shifts s
           LEFT JOIN employees e ON e.id=s.employee_id
           LEFT JOIN branches b ON b.id=s.branch_id
           LEFT JOIN invoices v ON v.shift_id=s.id AND v.kind IN ('sale','maintenance') AND v.status!='ملغاة'
           WHERE substr(s.started_at,1,10)=?""" + shift_sql + " GROUP BY s.id ORDER BY s.started_at DESC",
        [report_day] + shift_args,
    )
    employee_totals = query(
        """SELECT COALESCE(e.name, v.salesperson) employee_name, v.salesperson,
                  COUNT(v.id) invoice_count, COALESCE(SUM(v.total),0) sales_total,
                  COALESCE(SUM(v.paid),0) sales_paid, COALESCE(SUM(v.profit),0) sales_profit
           FROM invoices v LEFT JOIN employees e ON e.username=v.salesperson
           WHERE v.kind IN ('sale','maintenance') AND v.status!='ملغاة' AND v.date LIKE ?""" + invoice_sql + " GROUP BY v.salesperson, e.name ORDER BY sales_total DESC",
        [report_day + "%"] + invoice_args,
    )
    invoices = query(
        """SELECT v.number, v.date, v.kind, v.party_name, v.total, v.paid, v.profit,
                  b.name branch_name, v.salesperson, COALESCE(e.name, v.salesperson) employee_name, v.shift_id
           FROM invoices v LEFT JOIN employees e ON e.username=v.salesperson LEFT JOIN branches b ON b.id=v.branch_id
           WHERE v.kind IN ('sale','maintenance') AND v.status!='ملغاة' AND v.date LIKE ?""" + invoice_sql + " ORDER BY v.id DESC",
        [report_day + "%"] + invoice_args,
    )
    return render_template(
        "shift_reports.html", report_day=report_day, shift_rows=shift_rows,
        employee_totals=employee_totals, invoices=invoices,
    )


def shift_report_data(report_day):
    shift_rows = query(
        """SELECT s.*, b.name branch_name, COALESCE(e.name, s.username) employee_name, e.job_title,
                  COUNT(v.id) invoice_count, COALESCE(SUM(v.total),0) sales_total,
                  COALESCE(SUM(v.paid),0) sales_paid, COALESCE(SUM(v.profit),0) sales_profit
           FROM shifts s LEFT JOIN employees e ON e.id=s.employee_id
           LEFT JOIN branches b ON b.id=s.branch_id
           LEFT JOIN invoices v ON v.shift_id=s.id AND v.kind IN ('sale','maintenance') AND v.status!='ملغاة'
           WHERE substr(s.started_at,1,10)=? GROUP BY s.id ORDER BY s.started_at DESC""", (report_day,)
    )
    employee_totals = query(
        """SELECT COALESCE(e.name, v.salesperson) employee_name, v.salesperson,
                  COUNT(v.id) invoice_count, COALESCE(SUM(v.total),0) sales_total,
                  COALESCE(SUM(v.paid),0) sales_paid, COALESCE(SUM(v.profit),0) sales_profit
           FROM invoices v LEFT JOIN employees e ON e.username=v.salesperson
           WHERE v.kind IN ('sale','maintenance') AND v.status!='ملغاة' AND v.date LIKE ?
           GROUP BY v.salesperson, e.name ORDER BY sales_total DESC""", (report_day + "%",)
    )
    invoices = query(
        """SELECT v.number, v.date, v.kind, v.party_name, v.total, v.paid, v.profit,
                  v.salesperson, COALESCE(e.name, v.salesperson) employee_name, v.shift_id
           FROM invoices v LEFT JOIN employees e ON e.username=v.salesperson
           WHERE v.kind IN ('sale','maintenance') AND v.status!='ملغاة' AND v.date LIKE ?
           ORDER BY v.id DESC""", (report_day + "%",)
    )
    return shift_rows, employee_totals, invoices


@app.route("/reports/shifts/export.xlsx")
@login_required
def shift_report_export_excel():
    report_day = request.args.get("date") or date.today().isoformat()
    shift_rows, employee_totals, invoices = shift_report_data(report_day)
    wb = Workbook()
    ws = wb.active
    ws.title = "ملخص الموظفين"
    ws.append(["تقرير الورديات والمبيعات", report_day])
    ws.append(["الموظف", "عدد الفواتير", "إجمالي المبيعات", "المحصل", "الربح"])
    for r in employee_totals:
        ws.append([r["employee_name"] or "غير محدد", r["invoice_count"], r["sales_total"], r["sales_paid"], r["sales_profit"]])
    ws2 = wb.create_sheet("الورديات")
    ws2.append(["الموظف", "بدء الوردية", "إنهاء الوردية", "الفواتير", "المبيعات", "المحصل", "الربح"])
    for r in shift_rows:
        ws2.append([r["employee_name"], time12(r["started_at"]), time12(r["ended_at"]), r["invoice_count"], r["sales_total"], r["sales_paid"], r["sales_profit"]])
    ws3 = wb.create_sheet("الفواتير")
    ws3.append(["الفاتورة", "التاريخ", "النوع", "الموظف", "العميل", "الوردية", "الإجمالي", "المحصل", "الربح"])
    for r in invoices:
        ws3.append([r["number"], r["date"], KIND_LABELS.get(r["kind"], r["kind"]), r["employee_name"], r["party_name"], r["shift_id"] or "غير مرتبطة", r["total"], r["paid"], r["profit"]])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A3"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    return workbook_response(wb, f"تقرير-الورديات-{report_day}.xlsx")


@app.route("/reports/shifts/export.pdf")
@login_required
def shift_report_export_pdf():
    if not REPORTLAB_OK:
        flash("تصدير PDF غير متاح حاليًا، استخدم Excel أو ثبّت مكتبة PDF", "err")
        return redirect(url_for("shift_reports", date=request.args.get("date")))
    report_day = request.args.get("date") or date.today().isoformat()
    shift_rows, employee_totals, invoices = shift_report_data(report_day)
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Arabic", parent=styles["Normal"], fontName="DejaVu", fontSize=9, leading=12, alignment=2))
    story = [Paragraph(rtl_pdf(f"تقرير الورديات والمبيعات — {report_day}"), styles["Title"]), Spacer(1, 12)]
    table_data = [[rtl_pdf(x) for x in ["الموظف", "الفواتير", "المبيعات", "المحصل", "الربح"]]]
    for r in employee_totals:
        table_data.append([rtl_pdf(str(r["employee_name"] or "غير محدد")), str(r["invoice_count"]), money(r["sales_total"]), money(r["sales_paid"]), money(r["sales_profit"])])
    story.append(Table(table_data, repeatRows=1, hAlign="RIGHT"))
    story.append(Spacer(1, 14))
    shift_data = [[rtl_pdf(x) for x in ["الموظف", "البداية", "النهاية", "الفواتير", "المبيعات", "المحصل", "الربح"]]]
    for r in shift_rows:
        shift_data.append([rtl_pdf(str(r["employee_name"])), time12(r["started_at"]), time12(r["ended_at"]), str(r["invoice_count"]), money(r["sales_total"]), money(r["sales_paid"]), money(r["sales_profit"])])
    story.append(Table(shift_data, repeatRows=1, hAlign="RIGHT"))
    doc.build(story)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"تقرير-الورديات-{report_day}.pdf", mimetype="application/pdf")


@app.route("/reports/payroll-deductions")
@login_required
def payroll_deductions():
    period = request.args.get("period") or date.today().strftime("%Y-%m")
    emp_sql = "SELECT * FROM employees WHERE status='نشط'"
    emp_branch_sql, emp_branch_args = branch_filter("employees")
    emps = query(emp_sql + emp_branch_sql + " ORDER BY name", emp_branch_args)
    rows = []
    for emp in emps:
        preview = attendance_preview(emp["id"], period)
        rows.append({"employee": emp, **preview, "daily": round(float(emp["salary"] or 0) / 30.0, 2), "net_before_bonus": round(float(emp["salary"] or 0) - preview["deduct"], 2)})
    return render_template("payroll_deductions.html", period=period, rows=rows, total_deduct=sum(r["deduct"] for r in rows))


@app.route("/reports/payroll-deductions/export.pdf")
@login_required
def payroll_deductions_export_pdf():
    if not REPORTLAB_OK:
        flash("تصدير PDF غير متاح حاليًا على الخادم", "err")
        return redirect(url_for("payroll_deductions"))
    period = request.args.get("period") or date.today().strftime("%Y-%m")
    emp_sql = "SELECT * FROM employees WHERE status='نشط'"
    emp_branch_sql, emp_branch_args = branch_filter("employees")
    emps = query(emp_sql + emp_branch_sql + " ORDER BY name", emp_branch_args)
    rows = []
    for emp in emps:
        preview = attendance_preview(emp["id"], period)
        rows.append({"employee": emp, **preview, "daily": round(float(emp["salary"] or 0) / 30.0, 2), "net_before_bonus": round(float(emp["salary"] or 0) - preview["deduct"], 2)})
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PayrollArabic", fontName="DejaVu", fontSize=8, leading=12, alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="PayrollTitle", fontName="DejaVu", fontSize=15, leading=20, alignment=TA_RIGHT, textColor=colors.HexColor("#173653")))
    story = pdf_company_header({"ArabicTitle": styles["PayrollTitle"], "Arabic": styles["PayrollArabic"]}, f"مراجعة الخصومات الشهرية — {period}")
    data = [[rtl_pdf(x) for x in ["الموظف", "الراتب", "اليومية", "أيام الغياب", "دقائق التأخير", "الخصم", "الصافي المتوقع"]]]
    data += [[Paragraph(rtl_pdf(r["employee"]["name"]), styles["PayrollArabic"]), money(r["employee"]["salary"]), money(r["daily"]), str(r["absence_days"]), str(r["late_minutes"]), money(r["deduct"]), money(r["net_before_bonus"])] for r in rows]
    story += [Table(data, repeatRows=1, colWidths=[180, 95, 90, 95, 110, 95, 110], style=pdf_table_style()), Spacer(1, 10), Paragraph(rtl_pdf(f"إجمالي الخصومات المقترحة: {money(sum(r['deduct'] for r in rows))}"), styles["PayrollArabic"])]
    doc.build(story)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=f"خصومات-الرواتب-{period}.pdf", mimetype="application/pdf")


@app.route("/api/barcode")
@login_required
def api_barcode():
    code = (request.args.get("code") or "").strip()
    if not code:
        return jsonify({}), 404
    row = query(
        "SELECT * FROM products WHERE barcode=? OR sku=? LIMIT 1",
        (code, code),
        one=True,
    )
    if not row:
        return jsonify({}), 404
    return jsonify({k: row[k] for k in row.keys()})


@app.route("/api/coupon")
@login_required
def api_coupon():
    code = (request.args.get("code") or "").strip()
    subtotal = float(request.args.get("subtotal") or 0)
    coupon = query("SELECT * FROM coupons WHERE code=?", (code,), one=True)
    if not coupon or not coupon["active"]:
        return jsonify({"ok": False, "error": "كود غير صالح"}), 404
    if coupon["expires_on"] and coupon["expires_on"] < date.today().isoformat():
        return jsonify({"ok": False, "error": "منتهي"}), 400
    if int(coupon["used_count"] or 0) >= int(coupon["max_uses"] or 1):
        return jsonify({"ok": False, "error": "مستهلك"}), 400
    if coupon["discount_type"] == "نسبة":
        amount = round(subtotal * float(coupon["discount_value"]) / 100.0, 2)
    else:
        amount = float(coupon["discount_value"])
    return jsonify({"ok": True, "amount": amount, "type": coupon["discount_type"], "value": coupon["discount_value"]})


@app.route("/shift/open", methods=["POST"])
@login_required
def shift_open():
    if current_shift():
        flash("لديك وردية مفتوحة", "err")
        return redirect(request.referrer or url_for("pos"))
    employee = current_employee()
    if not employee:
        flash("اربط حساب المستخدم بموظف نشط أولاً من شاشة الموظفين", "err")
        return redirect(request.referrer or url_for("dashboard"))
    started = datetime.now()
    today = started.date().isoformat()
    start_time = datetime.strptime(get_setting("work_start_time", "09:00"), "%H:%M").replace(
        year=started.year, month=started.month, day=started.day
    )
    grace = int(get_setting("attendance_grace_minutes", "15") or 15)
    late_minutes = max(0, int((started - start_time).total_seconds() // 60) - grace)
    attendance = query(
        "SELECT id FROM attendance WHERE employee_id=? AND attendance_date=?",
        (employee["id"], today), one=True,
    )
    execute(
        "INSERT INTO shifts (username, started_at, opening_cash, employee_id, branch_id) VALUES (?,?,?,?,?)",
        (session.get("user"), started.strftime("%Y-%m-%d %H:%M"), float(request.form.get("opening_cash") or 0), employee["id"], current_branch_id()),
    )
    shift_id = db().execute("SELECT last_insert_rowid()").fetchone()[0]
    if attendance:
        execute(
            "UPDATE attendance SET shift_id=?, status='حاضر', started_at=?, late_minutes=?, notes=NULL WHERE id=?",
            (shift_id, started.strftime("%Y-%m-%d %H:%M"), late_minutes, attendance["id"]),
        )
    else:
        execute(
            """INSERT INTO attendance (employee_id, attendance_date, shift_id, status, started_at, late_minutes)
               VALUES (?,?,?,?,?,?)""",
            (employee["id"], today, shift_id, "حاضر", started.strftime("%Y-%m-%d %H:%M"), late_minutes),
        )
    audit_log("بدء وردية", "وردية", shift_id, str(shift_id), f"الموظف: {employee['name']}، التأخير: {late_minutes} دقيقة")
    flash("تم فتح الوردية", "ok")
    return redirect(request.referrer or url_for("pos"))


@app.route("/shift/close", methods=["POST"])
@login_required
def shift_close():
    sh = current_shift()
    if not sh:
        flash("لا توجد وردية مفتوحة", "err")
        return redirect(url_for("pos"))
    actual = float(request.form.get("actual_cash") or 0)
    expected = query(
        """SELECT COALESCE(SUM(CASE WHEN transaction_type='وارد' THEN amount ELSE -amount END),0) v
           FROM cash_transactions t JOIN cash_accounts a ON a.id=t.account_id
           WHERE a.name='درج' AND t.transaction_date >= ?""",
        (sh["started_at"],),
        one=True,
    )["v"]
    expected = float(expected) + float(sh["opening_cash"] or 0)
    ended = datetime.now().strftime("%Y-%m-%d %H:%M")
    execute("UPDATE shifts SET ended_at=? WHERE id=?", (ended, sh["id"]))
    execute("UPDATE attendance SET ended_at=? WHERE shift_id=?", (ended, sh["id"]))
    execute(
        """INSERT INTO daily_closures (shift_id, closing_date, expected_cash, actual_cash, difference, closed_by, closed_at, notes)
           VALUES (?,?,?,?,?,?,?,?)""",
        (
            sh["id"],
            date.today().isoformat(),
            expected,
            actual,
            round(actual - expected, 2),
            session.get("user"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            request.form.get("notes"),
        ),
    )
    audit_log("إنهاء وردية", "وردية", sh["id"], str(sh["id"]), f"الفرق النقدي: {round(actual - expected, 2):,.2f}")
    flash(f"أُغلقت الوردية. الفرق {round(actual - expected, 2):,.2f}", "ok")
    return redirect(url_for("treasury"))


@app.route("/treasury", methods=["GET", "POST"])
@login_required
def treasury():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "move":
            acc_id = int(request.form["account_id"])
            amount = float(request.form.get("amount") or 0)
            direction = request.form.get("direction") or "وارد"
            execute(
                """INSERT INTO cash_transactions
                   (account_id, transaction_type, amount, description, reference_type, transaction_date, created_by)
                   VALUES (?,?,?,?,?,?,?)""",
                (
                    acc_id,
                    direction,
                    amount,
                    request.form.get("description") or direction,
                    "manual",
                    datetime.now().strftime("%Y-%m-%d %H:%M"),
                    session.get("user"),
                ),
            )
            flash("تم تسجيل حركة الخزينة", "ok")
        elif action == "account":
            try:
                execute(
                    "INSERT INTO cash_accounts (name, account_type, opening_balance) VALUES (?,?,?)",
                    (request.form["name"].strip(), request.form.get("account_type") or "خزينة", float(request.form.get("opening_balance") or 0)),
                )
                flash("تمت إضافة الخزينة", "ok")
            except sqlite3.IntegrityError:
                flash("الاسم موجود", "err")
        elif action == "day_close":
            day = request.form.get("closing_date") or date.today().isoformat()
            existing = query(
                "SELECT id FROM daily_closures WHERE closing_date=? AND COALESCE(invoice_count,0)>0",
                (day,),
                one=True,
            )
            if existing:
                flash("هذا اليوم مغلق ومؤرشف مسبقاً", "err")
                return redirect(url_for("day_archive", close_id=existing["id"]))
            summary = day_summary(day)
            actual = float(request.form.get("actual_cash") or summary["drawer_balance"])
            expected = summary["drawer_balance"]
            open_sh = current_shift()
            execute(
                """INSERT INTO daily_closures
                   (shift_id, closing_date, expected_cash, actual_cash, difference, closed_by, closed_at, notes,
                    sales_total, sales_cash, returns_total, purchases_total, expenses_total, receipts_total,
                    invoice_count, summary_json)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    open_sh["id"] if open_sh else None,
                    day,
                    expected,
                    actual,
                    round(actual - expected, 2),
                    session.get("user"),
                    datetime.now().strftime("%Y-%m-%d %H:%M"),
                    request.form.get("notes") or "إغلاق يومية",
                    summary["sales_total"],
                    summary["sales_cash"],
                    summary["returns_total"],
                    summary["purchases_total"],
                    summary["expenses_total"],
                    summary["receipts_total"],
                    summary["invoice_count"],
                    json.dumps(
                        {
                            "cash_in": summary["cash_in"],
                            "cash_out": summary["cash_out"],
                            "net_cash": summary["net_cash"],
                            "payments_total": summary["payments_total"],
                            "purchase_returns": summary["purchase_returns"],
                        },
                        ensure_ascii=False,
                    ),
                ),
            )
            sh = current_shift()
            if sh:
                execute("UPDATE shifts SET ended_at=? WHERE id=?", (datetime.now().strftime("%Y-%m-%d %H:%M"), sh["id"]))
            flash(f"تم إغلاق يومية {day} وأرشفتها", "ok")
            last = query("SELECT id FROM daily_closures ORDER BY id DESC LIMIT 1", one=True)
            return redirect(url_for("day_archive", close_id=last["id"]))
        return redirect(url_for("treasury"))
    accounts = query(
        """SELECT a.*, a.opening_balance + COALESCE((
             SELECT SUM(CASE WHEN t.transaction_type='وارد' THEN t.amount ELSE -t.amount END)
             FROM cash_transactions t WHERE t.account_id=a.id),0) balance
           FROM cash_accounts a ORDER BY a.id"""
    )
    moves = query(
        """SELECT t.*, a.name account_name FROM cash_transactions t
           JOIN cash_accounts a ON a.id=t.account_id ORDER BY t.id DESC LIMIT 40"""
    )
    closures = query("SELECT * FROM daily_closures ORDER BY id DESC LIMIT 15")
    shifts = query("SELECT * FROM shifts ORDER BY id DESC LIMIT 15")
    summary = day_summary()
    return render_template(
        "treasury.html",
        accounts=accounts,
        moves=moves,
        closures=closures,
        shifts=shifts,
        shift=current_shift(),
        day=summary,
    )


@app.route("/treasury/archive/<int:close_id>")
@login_required
def day_archive(close_id):
    closure = query("SELECT * FROM daily_closures WHERE id=?", (close_id,), one=True)
    if not closure:
        flash("الإغلاق غير موجود", "err")
        return redirect(url_for("treasury"))
    day = closure["closing_date"]
    invoices = query(
        """SELECT * FROM invoices WHERE date LIKE ? ORDER BY id""",
        (day + "%",),
    )
    receipts = query(
        "SELECT * FROM payments WHERE date LIKE ? ORDER BY id",
        (day + "%",),
    )
    extra = {}
    if closure["summary_json"]:
        try:
            extra = json.loads(closure["summary_json"])
        except (TypeError, ValueError):
            extra = {}
    return render_template(
        "day_archive.html",
        closure=closure,
        invoices=invoices,
        receipts=receipts,
        extra=extra,
        settings=all_settings(),
    )


@app.route("/quotes", methods=["GET", "POST"])
@login_required
def quotes():
    if request.method == "POST":
        items = parse_items(request.form)
        if not items:
            flash("أضف بنداً", "err")
            return redirect(url_for("quotes"))
        subtotal = sum(i["line_total"] for i in items)
        discount = float(request.form.get("discount") or 0)
        tax = float(request.form.get("tax") or 0)
        total = round(subtotal - discount + tax, 2)
        number = next_number("QTE", "quotes")
        qid = execute(
            """INSERT INTO quotes (number, date, customer_id, customer_name, subtotal, discount, tax, total, notes, created_by, branch_id)
               VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            (
                number,
                date.today().isoformat(),
                request.form.get("party_id") or None,
                request.form.get("party_name") or request.form.get("party_search") or "عميل",
                subtotal,
                discount,
                tax,
                total,
                request.form.get("notes"),
                session.get("user"),
                current_branch_id(),
            ),
        )
        for item in items:
            execute(
                "INSERT INTO quote_items (quote_id, product_id, description, qty, unit_price, line_total) VALUES (?,?,?,?,?,?)",
                (qid, item["product_id"], item["description"], item["qty"], item["unit_price"], item["line_total"]),
            )
        audit_log("إنشاء", "عرض سعر", qid, number, f"العميل: {request.form.get('party_name') or request.form.get('party_search') or 'عميل'}، الإجمالي: {total:,.2f}")
        flash("تم حفظ عرض السعر", "ok")
        return redirect(url_for("quotes"))
    rows = query("SELECT * FROM quotes WHERE branch_id=? OR ? IS NULL ORDER BY id DESC", (current_branch_id(), current_branch_id()))
    customers = query("SELECT * FROM customers ORDER BY name")
    return render_template("quotes.html", rows=rows, customers=customers, pay_methods=PAY_METHODS, quote_form_only=False)


@app.route("/quotes/new", methods=["GET", "POST"])
@login_required
def quote_new():
    if request.method == "POST":
        return quotes()
    customers = query("SELECT * FROM customers ORDER BY name")
    return render_template("quotes.html", rows=[], customers=customers, pay_methods=PAY_METHODS, quote_form_only=True)


@app.route("/quotes/<int:qid>/invoice", methods=["POST"])
@login_required
def quote_to_invoice(qid):
    quote = query("SELECT * FROM quotes WHERE id=?", (qid,), one=True)
    items = query("SELECT * FROM quote_items WHERE quote_id=?", (qid,))
    if not quote or not items:
        flash("عرض السعر غير صالح", "err")
        return redirect(url_for("quotes"))
    class Form(dict):
        def getlist(self, key):
            return self.get(key, [])
    form = Form()
    form["party_id"] = quote["customer_id"] or ""
    form["party_name"] = quote["customer_name"]
    form["discount"] = quote["discount"]
    form["tax"] = quote["tax"]
    form["payment_method"] = request.form.get("payment_method") or "نقدي"
    form["item_product_id[]"] = [str(i["product_id"] or "") for i in items]
    form["item_desc[]"] = [i["description"] for i in items]
    form["item_qty[]"] = [str(i["qty"]) for i in items]
    form["item_price[]"] = [str(i["unit_price"]) for i in items]
    form["item_cost[]"] = ["0"] * len(items)
    try:
        inv_id = save_invoice("sale", form, related_id=qid)
        execute("UPDATE quotes SET status=?, invoice_id=? WHERE id=?", ("محوّل", inv_id, qid))
        audit_log("تحويل إلى فاتورة", "عرض سعر", qid, quote["number"], f"رقم الفاتورة الناتجة: {query('SELECT number FROM invoices WHERE id=?', (inv_id,), one=True)['number']}")
        flash("تم تحويل العرض إلى فاتورة", "ok")
        return redirect(url_for("invoice_view", inv_id=inv_id))
    except Exception as exc:
        flash(str(exc), "err")
        return redirect(url_for("quotes"))


@app.route("/marketing", methods=["GET", "POST"])
@login_required
def marketing():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "campaign":
            execute(
                """INSERT INTO campaigns (name, discount_type, discount_value, starts_on, ends_on, active)
                   VALUES (?,?,?,?,?,1)""",
                (
                    request.form["name"].strip(),
                    request.form.get("discount_type") or "نسبة",
                    float(request.form.get("discount_value") or 0),
                    request.form.get("starts_on") or date.today().isoformat(),
                    request.form.get("ends_on") or date.today().isoformat(),
                ),
            )
            flash("تمت إضافة الحملة", "ok")
        elif action == "coupon":
            try:
                execute(
                    """INSERT INTO coupons (code, campaign_id, discount_type, discount_value, max_uses, expires_on, active)
                       VALUES (?,?,?,?,?,?,1)""",
                    (
                        request.form["code"].strip().upper(),
                        request.form.get("campaign_id") or None,
                        request.form.get("discount_type") or "نسبة",
                        float(request.form.get("discount_value") or 0),
                        int(request.form.get("max_uses") or 1),
                        request.form.get("expires_on") or None,
                    ),
                )
                flash("تمت إضافة الكوبون", "ok")
            except sqlite3.IntegrityError:
                flash("الكود موجود", "err")
        elif action == "target":
            execute(
                """INSERT INTO sales_targets (username, target_month, target_amount, commission_percent)
                   VALUES (?,?,?,?)
                   ON CONFLICT(username, target_month) DO UPDATE SET target_amount=excluded.target_amount, commission_percent=excluded.commission_percent""",
                (
                    request.form.get("username") or session.get("user"),
                    request.form.get("target_month") or date.today().strftime("%Y-%m-01"),
                    float(request.form.get("target_amount") or 0),
                    float(request.form.get("commission_percent") or 0),
                ),
            )
            flash("تم حفظ المستهدف", "ok")
        return redirect(url_for("marketing"))
    campaigns = query("SELECT * FROM campaigns ORDER BY id DESC")
    coupons = query("SELECT * FROM coupons ORDER BY code")
    targets = query("SELECT * FROM sales_targets ORDER BY target_month DESC")
    users = query("SELECT username, full_name FROM users ORDER BY username")
    return render_template("marketing.html", campaigns=campaigns, coupons=coupons, targets=targets, users=users)


@app.route("/finance", methods=["GET", "POST"])
@login_required
def finance():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "check":
            execute(
                """INSERT INTO checks (customer_id, supplier_id, check_number, check_type, bank_name, amount, due_date, status, notes, created_by, created_at)
                   VALUES (?,?,?,?,?,?,?,'معلق',?,?,?)""",
                (
                    request.form.get("customer_id") or None,
                    request.form.get("supplier_id") or None,
                    request.form["check_number"].strip(),
                    request.form.get("check_type") or "وارد",
                    request.form.get("bank_name"),
                    float(request.form.get("amount") or 0),
                    request.form.get("due_date") or date.today().isoformat(),
                    request.form.get("notes"),
                    session.get("user"),
                    datetime.now().strftime("%Y-%m-%d %H:%M"),
                ),
            )
            flash("تم تسجيل الشيك", "ok")
        elif action == "check_status":
            execute("UPDATE checks SET status=? WHERE id=?", (request.form.get("status"), request.form.get("check_id")))
            flash("تم تحديث حالة الشيك", "ok")
        elif action == "pay_installment":
            iid = int(request.form["installment_id"])
            amt = float(request.form.get("amount") or 0)
            inst = query("SELECT * FROM installments WHERE id=?", (iid,), one=True)
            if inst:
                paid = float(inst["paid_amount"]) + amt
                status = "مسدد" if paid >= float(inst["amount"]) else "جزئي"
                execute("UPDATE installments SET paid_amount=?, status=? WHERE id=?", (paid, status, iid))
                post_cash("نقدي", amt, f"قسط رقم {inst['installment_no']}", "installment", iid, "in")
                if inst["customer_id"]:
                    execute("UPDATE customers SET balance = balance - ? WHERE id=?", (amt, inst["customer_id"]))
                    add_ledger("customer", inst["customer_id"], "سداد", f"قسط {inst['installment_no']}", credit=amt, payment_date=date.today().isoformat())
                flash("تم تحصيل القسط", "ok")
        elif action == "receipt":
            party_type = request.form.get("party_type") or "customer"
            party_id = int(request.form["party_id"])
            amount = float(request.form.get("amount") or 0)
            method = request.form.get("method") or "نقدي"
            if party_type == "customer":
                execute("UPDATE customers SET balance = balance - ? WHERE id=?", (amount, party_id))
                name = query("SELECT name FROM customers WHERE id=?", (party_id,), one=True)
                add_ledger("customer", party_id, "سند قبض", request.form.get("notes") or "تحصيل", credit=amount, payment_date=date.today().isoformat())
                post_cash(method, amount, f"سند قبض {name['name'] if name else ''}", "receipt", party_id, "in")
            else:
                execute("UPDATE suppliers SET balance = balance - ? WHERE id=?", (amount, party_id))
                name = query("SELECT name FROM suppliers WHERE id=?", (party_id,), one=True)
                add_ledger("supplier", party_id, "سند صرف", request.form.get("notes") or "سداد", debit=amount, payment_date=date.today().isoformat())
                post_cash(method, amount, f"سند صرف {name['name'] if name else ''}", "payment", party_id, "out")
            year = datetime.now().strftime("%Y")
            last_rcp = query(
                "SELECT receipt_no FROM payments WHERE receipt_no LIKE ? ORDER BY id DESC LIMIT 1",
                (f"RCP-{year}-%",),
                one=True,
            )
            seq = 1
            if last_rcp and last_rcp["receipt_no"]:
                try:
                    seq = int(str(last_rcp["receipt_no"]).split("-")[-1]) + 1
                except ValueError:
                    seq = 1
            receipt_no = f"RCP-{year}-{seq:04d}"
            pay_id = execute(
                """INSERT INTO payments (party_type, party_id, party_name, amount, method, date, notes, receipt_no)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (
                    party_type,
                    party_id,
                    name["name"] if name else "",
                    amount,
                    method,
                    date.today().isoformat(),
                    request.form.get("notes"),
                    receipt_no,
                ),
            )
            flash("تم تسجيل السند", "ok")
            return redirect(url_for("receipt_print", pay_id=pay_id))
        return redirect(url_for("finance"))
    checks = query("SELECT * FROM checks ORDER BY due_date")
    installments = query(
        """SELECT i.*, c.name customer_name FROM installments i
           JOIN customers c ON c.id=i.customer_id ORDER BY i.due_date"""
    )
    customers = query("SELECT * FROM customers ORDER BY name")
    suppliers = query("SELECT * FROM suppliers ORDER BY name")
    payments = query("SELECT * FROM payments ORDER BY id DESC LIMIT 20")
    return render_template(
        "finance.html",
        checks=checks,
        installments=installments,
        customers=customers,
        suppliers=suppliers,
        payments=payments,
        pay_methods=PAY_METHODS,
    )


@app.route("/receipt/<int:pay_id>/print")
@login_required
def receipt_print(pay_id):
    pay = query("SELECT * FROM payments WHERE id=?", (pay_id,), one=True)
    if not pay:
        flash("السند غير موجود", "err")
        return redirect(url_for("finance"))
    party = None
    if pay["party_type"] == "customer" and pay["party_id"]:
        party = query("SELECT * FROM customers WHERE id=?", (pay["party_id"],), one=True)
    elif pay["party_type"] == "supplier" and pay["party_id"]:
        party = query("SELECT * FROM suppliers WHERE id=?", (pay["party_id"],), one=True)
    return render_template(
        "receipt_print.html",
        pay=pay,
        party=party,
        settings=all_settings(),
    )


@app.route("/inventory/transfer", methods=["POST"])
@login_required
def inventory_transfer():
    src = int(request.form["source_id"])
    dst = int(request.form["dest_id"])
    pid = int(request.form["product_id"])
    qty = float(request.form.get("qty") or 0)
    if src == dst or qty <= 0:
        flash("بيانات التحويل غير صحيحة", "err")
        return redirect(url_for("inventory"))
    number = next_number("TRN", "stock_transfers")
    tid = execute(
        """INSERT INTO stock_transfers (number, source_id, dest_id, status, requested_by, created_at, shipped_at, received_at)
           VALUES (?,?,?,'مستلم',?,?,?,?)""",
        (
            number,
            src,
            dst,
            session.get("user"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
        ),
    )
    execute("INSERT INTO stock_transfer_lines (transfer_id, product_id, qty) VALUES (?,?,?)", (tid, pid, qty))
    src_w = query("SELECT name FROM warehouses WHERE id=?", (src,), one=True)
    dst_w = query("SELECT name FROM warehouses WHERE id=?", (dst,), one=True)
    apply_stock(pid, 0, 0, "تحويل", number, f"من {src_w['name'] if src_w else src} إلى {dst_w['name'] if dst_w else dst} كمية {qty}")
    flash("تم تحويل الكمية بين المخازن", "ok")
    return redirect(url_for("inventory"))


@app.route("/inventory/count", methods=["POST"])
@login_required
def inventory_count():
    wid = int(request.form["warehouse_id"])
    pid = int(request.form["product_id"])
    counted = float(request.form.get("counted_qty") or 0)
    prod = query("SELECT * FROM products WHERE id=?", (pid,), one=True)
    expected = float(prod["qty"]) if prod else 0
    cid = execute(
        """INSERT INTO stock_counts (warehouse_id, count_type, status, counted_by, created_at, closed_at)
           VALUES (?,?,'مغلق',?,?,?)""",
        (
            wid,
            request.form.get("count_type") or "دوري",
            session.get("user"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
        ),
    )
    execute(
        """INSERT INTO stock_count_lines (count_id, product_id, expected_qty, counted_qty, difference)
           VALUES (?,?,?,?,?)""",
        (cid, pid, expected, counted, counted - expected),
    )
    if prod and counted != expected:
        apply_stock(pid, counted - expected, prod["cost"], "جرد", f"CNT-{cid}", "تسوية جرد")
    flash("تم حفظ الجرد", "ok")
    return redirect(url_for("inventory"))


@app.route("/inventory/serial", methods=["POST"])
@login_required
def inventory_serial():
    execute(
        "INSERT INTO serial_numbers (product_id, serial_number, status, notes) VALUES (?,?, 'متاح', ?)",
        (request.form["product_id"], request.form["serial_number"].strip(), request.form.get("notes")),
    )
    flash("تم تسجيل الرقم التسلسلي", "ok")
    return redirect(url_for("inventory"))


@app.route("/journal", methods=["GET", "POST"])
@login_required
def journal():
    if request.method == "POST":
        add_journal(
            request.form.get("description") or "قيد يدوي",
            request.form["account_name"].strip(),
            debit=float(request.form.get("debit") or 0),
            credit=float(request.form.get("credit") or 0),
            reference_type="manual",
        )
        flash("تم ترحيل القيد", "ok")
        return redirect(url_for("journal"))
    rows = query("SELECT * FROM journal_entries ORDER BY id DESC LIMIT 80")
    totals = query("SELECT COALESCE(SUM(debit),0) d, COALESCE(SUM(credit),0) c FROM journal_entries", one=True)
    return render_template("journal.html", rows=rows, totals=totals)


def manufacturing_now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def manufacturing_order_cost(order_id):
    row = query(
        """SELECT COALESCE(SUM((-qty) * unit_cost),0) material_cost
           FROM raw_material_moves
           WHERE reference_type='production_order' AND reference_id=? AND move_type='صرف'""",
        (order_id,), one=True,
    )
    waste = query("SELECT waste_cost, labor_cost, overhead_cost FROM production_orders WHERE id=?", (order_id,), one=True)
    return {
        "material": float(row["material_cost"] if row else 0),
        "labor": float(waste["labor_cost"] if waste else 0),
        "overhead": float(waste["overhead_cost"] if waste else 0),
        "waste": float(waste["waste_cost"] if waste else 0),
    }


@app.route("/manufacturing", methods=["GET", "POST"])
@login_required
def manufacturing():
    if request.method == "POST":
        action = request.form.get("action")
        try:
            if action == "add_material":
                code = (request.form.get("code") or "").strip()
                name = (request.form.get("name") or "").strip()
                if not code or not name:
                    raise ValueError("أدخل كود واسم المادة الخام")
                execute(
                    """INSERT INTO raw_materials (code,name,unit,cost,qty,min_qty,status,notes)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (code, name, request.form.get("unit") or "قطعة", float(request.form.get("cost") or 0),
                     float(request.form.get("qty") or 0), float(request.form.get("min_qty") or 0),
                     request.form.get("status") or "فعال", request.form.get("notes")),
                )
                flash("تمت إضافة المادة الخام", "ok")
            elif action == "delete_material":
                mid = int(request.form.get("material_id") or 0)
                if query("SELECT 1 FROM manufacturing_bom_items WHERE raw_material_id=?", (mid,), one=True):
                    raise ValueError("لا يمكن حذف مادة مستخدمة في قائمة مواد")
                execute("DELETE FROM raw_materials WHERE id=?", (mid,))
                flash("تم حذف المادة الخام", "ok")
            elif action == "save_bom":
                product_id = int(request.form.get("bom_product_id") or 0)
                if not query("SELECT id FROM products WHERE id=?", (product_id,), one=True):
                    raise ValueError("اختر المنتج النهائي")
                bom_id = execute(
                    """INSERT INTO manufacturing_boms (product_id,version,status,labor_cost,overhead_cost,notes)
                       VALUES (?,?,?,?,?,?)""",
                    (product_id, request.form.get("version") or "1", request.form.get("bom_status") or "فعال",
                     float(request.form.get("bom_labor_cost") or 0), float(request.form.get("bom_overhead_cost") or 0), request.form.get("bom_notes")),
                )
                material_ids = request.form.getlist("bom_material_id[]")
                quantities = request.form.getlist("bom_material_qty[]")
                for idx, material_id in enumerate(material_ids):
                    qty = float(quantities[idx] or 0) if idx < len(quantities) else 0
                    if qty > 0 and query("SELECT id FROM raw_materials WHERE id=?", (material_id,), one=True):
                        execute("INSERT INTO manufacturing_bom_items (bom_id,raw_material_id,qty) VALUES (?,?,?)", (bom_id, material_id, qty))
                flash("تم حفظ قائمة المواد", "ok")
            elif action == "open_order":
                product_id = int(request.form.get("product_id") or 0)
                product = query("SELECT * FROM products WHERE id=?", (product_id,), one=True)
                if not product:
                    raise ValueError("اختر المنتج النهائي")
                planned = float(request.form.get("planned_qty") or 0)
                if planned <= 0:
                    raise ValueError("الكمية المخططة يجب أن تكون أكبر من صفر")
                bom = query("SELECT * FROM manufacturing_boms WHERE product_id=? AND status='فعال' ORDER BY id DESC LIMIT 1", (product_id,), one=True)
                number = next_number("MO", "production_orders")
                order_id = execute(
                    """INSERT INTO production_orders
                       (number,product_id,planned_qty,status,current_stage,opened_at,labor_cost,overhead_cost,warehouse_id,notes,created_by,branch_id)
                       VALUES (?,?,?,'مفتوح','تجهيز',?,?,?,?,?,?,?)""",
                    (number, product_id, planned, manufacturing_now(),
                     float(request.form.get("labor_cost") or (bom["labor_cost"] if bom else 0)),
                     float(request.form.get("overhead_cost") or (bom["overhead_cost"] if bom else 0)),
                     request.form.get("warehouse_id") or None, request.form.get("order_notes"), session.get("user"), current_branch_id()),
                )
                for sequence, stage in enumerate(("تجهيز", "تصنيع", "فحص"), 1):
                    execute("INSERT INTO production_stages (order_id,name,sequence,status) VALUES (?,?,?,'معلق')", (order_id, stage, sequence))
                flash(f"تم فتح أمر التصنيع {number}", "ok")
            else:
                raise ValueError("إجراء تصنيع غير معروف")
        except (ValueError, sqlite3.IntegrityError) as exc:
            flash(str(exc) if isinstance(exc, ValueError) else "تعذر حفظ البيانات (قد يكون الكود أو الإصدار مكرراً)", "err")
        return redirect(url_for("manufacturing", order_id=request.form.get("order_id") or None))
    materials = query("SELECT * FROM raw_materials ORDER BY name")
    products = query("SELECT id,sku,name,unit,cost,qty FROM products WHERE item_type IS NULL OR item_type!='خدمة' ORDER BY name")
    boms = query("""SELECT b.*, p.name product_name, p.sku FROM manufacturing_boms b JOIN products p ON p.id=b.product_id ORDER BY b.id DESC""")
    mo_sql, mo_args = branch_filter("o", include_all=True)
    orders = query("SELECT o.*, p.name product_name, p.sku FROM production_orders o JOIN products p ON p.id=o.product_id WHERE 1=1" + mo_sql + " ORDER BY o.id DESC", mo_args)
    selected_id = request.args.get("order_id", type=int) or (orders[0]["id"] if orders else None)
    selected = query("SELECT o.*, p.name product_name, p.sku FROM production_orders o JOIN products p ON p.id=o.product_id WHERE o.id=?" + mo_sql, [selected_id] + mo_args, one=True) if selected_id else None
    stages = query("SELECT * FROM production_stages WHERE order_id=? ORDER BY sequence", (selected_id,)) if selected_id else []
    moves = query("""SELECT m.*, r.name material_name, r.code FROM raw_material_moves m JOIN raw_materials r ON r.id=m.raw_material_id
                    WHERE m.reference_type='production_order' AND m.reference_id=? ORDER BY m.id DESC""", (selected_id,)) if selected_id else []
    cost = manufacturing_order_cost(selected_id) if selected_id else {"material": 0, "labor": 0, "overhead": 0, "waste": 0}
    warehouses = query("SELECT id,name FROM warehouses ORDER BY id")
    return render_template("manufacturing.html", materials=materials, products=products, boms=boms, orders=orders,
                           selected=selected, stages=stages, moves=moves, cost=cost, warehouses=warehouses)


@app.route("/manufacturing/material/<int:material_id>/delete", methods=["POST"])
@login_required
def manufacturing_material_delete(material_id):
    if query("SELECT 1 FROM manufacturing_bom_items WHERE raw_material_id=?", (material_id,), one=True):
        flash("لا يمكن حذف مادة مستخدمة في قائمة مواد", "err")
    else:
        execute("DELETE FROM raw_materials WHERE id=?", (material_id,))
        flash("تم حذف المادة الخام", "ok")
    return redirect(url_for("manufacturing"))


@app.route("/manufacturing/order/<int:order_id>/start", methods=["POST"])
@login_required
def manufacturing_start(order_id):
    order = query("SELECT * FROM production_orders WHERE id=?", (order_id,), one=True)
    if not order or order["status"] in ("مكتمل", "ملغي"):
        flash("أمر التصنيع غير صالح للبدء", "err")
    else:
        stamp = manufacturing_now()
        execute("UPDATE production_orders SET status='قيد التشغيل', started_at=?, current_stage='تجهيز' WHERE id=?", (stamp, order_id))
        execute("UPDATE production_stages SET status='قيد التشغيل', started_at=? WHERE order_id=? AND sequence=1", (stamp, order_id))
        flash("تم بدء أمر التصنيع", "ok")
    return redirect(url_for("manufacturing", order_id=order_id))


@app.route("/manufacturing/order/<int:order_id>/issue", methods=["POST"])
@login_required
def manufacturing_issue(order_id):
    order = query("SELECT * FROM production_orders WHERE id=?", (order_id,), one=True)
    try:
        if not order or order["status"] not in ("قيد التشغيل", "مفتوح"):
            raise ValueError("ابدأ أمر التصنيع أولاً")
        bom = query("SELECT * FROM manufacturing_boms WHERE product_id=? AND status='فعال' ORDER BY id DESC LIMIT 1", (order["product_id"],), one=True)
        if not bom:
            raise ValueError("لا توجد قائمة مواد فعالة لهذا المنتج")
        if query("SELECT 1 FROM raw_material_moves WHERE reference_type='production_order' AND reference_id=? AND move_type='صرف'", (order_id,), one=True):
            raise ValueError("تم صرف مواد هذا الأمر مسبقاً")
        items = query("SELECT * FROM manufacturing_bom_items WHERE bom_id=?", (bom["id"],))
        if not items:
            raise ValueError("قائمة المواد فارغة")
        stamp = manufacturing_now()
        for item in items:
            material = query("SELECT * FROM raw_materials WHERE id=?", (item["raw_material_id"],), one=True)
            required = float(item["qty"]) * float(order["planned_qty"])
            if not material or float(material["qty"]) < required:
                raise ValueError(f"كمية المادة غير كافية: {material['name'] if material else item['raw_material_id']}")
        for item in items:
            material = query("SELECT * FROM raw_materials WHERE id=?", (item["raw_material_id"],), one=True)
            required = float(item["qty"]) * float(order["planned_qty"])
            execute("UPDATE raw_materials SET qty=qty-? WHERE id=?", (required, material["id"]))
            execute("""INSERT INTO raw_material_moves (raw_material_id,qty,move_type,reference_type,reference_id,unit_cost,date,created_by)
                       VALUES (?,?, 'صرف','production_order',?,?,?,?)""", (material["id"], -required, order_id, material["cost"], stamp, session.get("user")))
        flash("تم صرف مواد التصنيع حسب قائمة المواد", "ok")
    except ValueError as exc:
        flash(str(exc), "err")
    return redirect(url_for("manufacturing", order_id=order_id))


@app.route("/manufacturing/order/<int:order_id>/waste", methods=["POST"])
@login_required
def manufacturing_waste(order_id):
    try:
        order = query("SELECT * FROM production_orders WHERE id=?", (order_id,), one=True)
        material_id = int(request.form.get("raw_material_id") or 0)
        qty = float(request.form.get("qty") or 0)
        material = query("SELECT * FROM raw_materials WHERE id=?", (material_id,), one=True)
        if not order or order["status"] not in ("قيد التشغيل", "مفتوح"):
            raise ValueError("الأمر غير متاح لتسجيل الهالك")
        if not material or qty <= 0 or float(material["qty"]) < qty:
            raise ValueError("تحقق من المادة وكمية الهالك")
        stamp = manufacturing_now()
        execute("UPDATE raw_materials SET qty=qty-? WHERE id=?", (qty, material_id))
        execute("""INSERT INTO raw_material_moves (raw_material_id,qty,move_type,reference_type,reference_id,unit_cost,date,created_by)
                   VALUES (?,?, 'هالك','production_order',?,?,?,?)""", (material_id, -qty, order_id, material["cost"], stamp, session.get("user")))
        execute("UPDATE production_orders SET waste_cost=waste_cost+? WHERE id=?", (qty * float(material["cost"]), order_id))
        flash("تم تسجيل الهالك وخصمه من المادة الخام", "ok")
    except ValueError as exc:
        flash(str(exc), "err")
    return redirect(url_for("manufacturing", order_id=order_id))


@app.route("/manufacturing/order/<int:order_id>/finish", methods=["POST"])
@login_required
def manufacturing_finish(order_id):
    try:
        order = query("SELECT * FROM production_orders WHERE id=?", (order_id,), one=True)
        completed_qty = float(request.form.get("completed_qty") or (order["planned_qty"] if order else 0))
        if not order or order["status"] == "مكتمل" or completed_qty <= 0:
            raise ValueError("أمر التصنيع أو الكمية غير صالح")
        costs = manufacturing_order_cost(order_id)
        total = costs["material"] + costs["labor"] + costs["overhead"] + costs["waste"]
        unit_cost = total / completed_qty if completed_qty else 0
        apply_stock(order["product_id"], completed_qty, unit_cost, "إنتاج", order["number"], "إضافة منتج تام")
        execute("UPDATE products SET cost=? WHERE id=?", (round(unit_cost, 4), order["product_id"]))
        stamp = manufacturing_now()
        execute("UPDATE production_orders SET completed_qty=?, status='مكتمل', completed_at=?, current_stage='فحص' WHERE id=?", (completed_qty, stamp, order_id))
        execute("UPDATE production_stages SET status='مكتمل', ended_at=? WHERE order_id=?", (stamp, order_id))
        flash(f"تم إنهاء الأمر. تكلفة الوحدة: {money(unit_cost)}", "ok")
    except ValueError as exc:
        flash(str(exc), "err")
    return redirect(url_for("manufacturing", order_id=order_id))


@app.route("/manifest.webmanifest")
def web_manifest():
    body = {
        "name": APP_NAME,
        "short_name": "الأصلي",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#0f2744",
        "theme_color": "#0f2744",
        "lang": "ar",
        "dir": "rtl",
        "icons": [
            {"src": url_for("static", filename="icons/icon-192.png"), "sizes": "192x192", "type": "image/png"},
            {"src": url_for("static", filename="icons/icon-512.png"), "sizes": "512x512", "type": "image/png"},
        ],
    }
    resp = make_response(json.dumps(body, ensure_ascii=False))
    resp.headers["Content-Type"] = "application/manifest+json"
    return resp


@app.route("/sw.js")
def service_worker():
    return send_from_directory(os.path.join(app.root_path, "static"), "sw.js", mimetype="application/javascript")


@app.route("/api/offline-catalog")
@login_required
def api_offline_catalog():
    products = rows_to_dicts(query(
        "SELECT id, sku, name, category, brand, car_model, qty, cost, price, unit, location, warehouse, aisle, shelf, bin, barcode, min_qty FROM products ORDER BY name"
    ))
    customers = rows_to_dicts(query("SELECT id, name, phone, address, balance, points, credit_limit, discount_percent FROM customers ORDER BY name"))
    suppliers = rows_to_dicts(query("SELECT id, name, phone, address, balance FROM suppliers ORDER BY name"))
    employees = rows_to_dicts(query("SELECT id, name, job_title, phone FROM employees WHERE status='نشط' ORDER BY name"))
    return jsonify({
        "ok": True,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "products": products,
        "customers": customers,
        "suppliers": suppliers,
        "employees": employees,
        "settings": all_settings(),
    })


@app.route("/api/sync/status")
@login_required
def api_sync_status():
    pending = query("SELECT COUNT(*) c FROM sync_queue WHERE status='pending'", one=True)["c"]
    failed = query("SELECT COUNT(*) c FROM sync_queue WHERE status='error'", one=True)["c"]
    logs = rows_to_dicts(query("SELECT * FROM sync_log ORDER BY id DESC LIMIT 12"))
    queue = rows_to_dicts(query("SELECT id, op_uuid, op_type, status, error, created_at FROM sync_queue ORDER BY id DESC LIMIT 20"))
    return jsonify({
        "ok": True,
        "online": True,
        "pending": pending,
        "failed": failed,
        "last_sync_at": get_setting("last_sync_at", ""),
        "last_sync_status": get_setting("last_sync_status", ""),
        "device_id": get_setting("device_id", ""),
        "logs": logs,
        "queue": queue,
    })


@app.route("/api/sync/push", methods=["POST"])
@login_required
def api_sync_push():
    body = request.get_json(silent=True) or {}
    ops = body.get("ops") or []
    device_id = body.get("device_id") or get_setting("device_id")
    device_name = body.get("device_name") or get_setting("device_name", "جهاز")
    platform = body.get("platform") or "web"
    execute(
        """INSERT INTO sync_devices (device_id, device_name, platform, last_seen, last_push)
           VALUES (?,?,?,?,?)
           ON CONFLICT(device_id) DO UPDATE SET device_name=excluded.device_name, platform=excluded.platform, last_seen=excluded.last_seen, last_push=excluded.last_push""",
        (device_id, device_name, platform, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    applied, skipped, errors = 0, 0, []
    for op in ops:
        op_uuid = op.get("op_uuid") or str(uuid.uuid4())
        exists = query("SELECT id FROM sync_queue WHERE op_uuid=?", (op_uuid,), one=True)
        if exists:
            skipped += 1
            continue
        created = op.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        payload = op.get("payload") or {}
        execute(
            """INSERT INTO sync_queue (op_uuid, device_id, op_type, payload, status, created_at)
               VALUES (?,?,?,?, 'pending', ?)""",
            (op_uuid, device_id, op.get("op_type") or op.get("type") or "unknown", json.dumps(payload, ensure_ascii=False), created),
        )
        try:
            apply_sync_op(op)
            execute(
                "UPDATE sync_queue SET status='applied', applied_at=? WHERE op_uuid=?",
                (datetime.now().strftime("%Y-%m-%d %H:%M:%S"), op_uuid),
            )
            applied += 1
        except Exception as exc:
            execute("UPDATE sync_queue SET status='error', error=? WHERE op_uuid=?", (str(exc), op_uuid))
            errors.append({"op_uuid": op_uuid, "error": str(exc)})
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    execute("INSERT INTO settings (key, value) VALUES ('last_sync_at', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (stamp,))
    status_txt = "ok" if not errors else "partial"
    execute("INSERT INTO settings (key, value) VALUES ('last_sync_status', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (status_txt,))
    log_sync("push", status_txt, f"تطبيق {applied} وتخطي {skipped}")
    return jsonify({"ok": True, "applied": applied, "skipped": skipped, "errors": errors, "synced_at": stamp})


@app.route("/api/sync/pull")
@login_required
def api_sync_pull():
    snap = snapshot_payload()
    execute(
        "INSERT INTO settings (key, value) VALUES ('last_sync_at', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (snap["generated_at"],),
    )
    log_sync("pull", "ok", "تم سحب النسخة المحلية")
    return jsonify({"ok": True, "snapshot": snap})


@app.route("/sync", methods=["GET", "POST"])
@login_required
def sync_center():
    if request.method == "POST" and request.form.get("action") == "test":
        flash("المزامنة المحلية جاهزة. عند توفر الإنترنت تُرسل العمليات المؤجلة تلقائياً.", "ok")
        return redirect(url_for("sync_center"))
    pending = query("SELECT COUNT(*) c FROM sync_queue WHERE status='pending'", one=True)["c"]
    failed = query("SELECT COUNT(*) c FROM sync_queue WHERE status='error'", one=True)["c"]
    logs = query("SELECT * FROM sync_log ORDER BY id DESC LIMIT 20")
    queue = query("SELECT * FROM sync_queue ORDER BY id DESC LIMIT 30")
    devices = query("SELECT * FROM sync_devices ORDER BY last_seen DESC")
    return render_template("sync.html", pending=pending, failed=failed, logs=logs, queue=queue, devices=devices)


init_db()

if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000)),
        debug=os.environ.get("FLASK_DEBUG", "").lower() in ("1", "true", "yes"),
    )
